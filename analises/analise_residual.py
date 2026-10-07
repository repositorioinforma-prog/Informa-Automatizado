"""
Análise Residual (A.R) — a tela ORIGINAL do Relatória Analítica dentro do
Streamlit.

A interface (HTML/CSS/JS) é a mesma do app desktop:
componentes/relatoria_residual/index.html (cópia de residuals_page.py, só
trocando o QWebChannel do Qt pela ponte). A `PonteResidual` abaixo é a
`ResidualsBridge` original com o MESMO código; só foi trocado o que era do Qt
(trechos marcados com "[Streamlit]"):
- "Carregar dados": o arquivo vem do seletor do navegador (gravado num
  temporário com a mesma extensão) no lugar do QFileDialog;
- "Exportar XLSX": os arquivos (analise_residual.xlsx e, se pedidos, _DEBUG e
  _CASES_DEBUG) são baixados pelo navegador no lugar do "Salvar como";
- os erros (QMessageBox) aparecem como faixas no canto da ferramenta;
- melhorias desta versão (marcadas com "[Streamlit]"): fila na barra lateral
  (cards, pop-ups, substituir repetido), colunas na ordem dos códigos do SPSS,
  arquivos de conferência (_DEBUG) opcionais e estado da tela guardado ao
  trocar de ferramenta;
- .sav é lido pelo pyreadstat ou, se a DLL estiver bloqueada no computador,
  pelo leitor em Python puro (core/sav_puro.py).
Os auxiliares de cálculo estão em core/residuos_math.py (cópia sem alteração).
"""
from __future__ import annotations

import json
import os
import shutil
import tempfile
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analises._pilha_lateral import pilha_lateral
from analises._ponte_relatoria import (
    PonteBase, conferir_versao_tela, desenhar_componente, mostrar_notificacoes, obter_ponte, processar_pedidos,
)
from analises.tutoriais_analitica import botao_tutorial
from core.residuos_math import (  # noqa: F401
    FIXED_BOTTOM_LABELS, _adjusted_residuals, _apply_fixed_bottom_order, _canon, _cell_style_for_z,
    _expected_from_ct, _is_peso_name, _pct_matrix, _round_obs_to_integer, _safe_label_from_map, _safe_str,
    _unique_sheet_name,
)

#: precisa bater com <meta name="versao-tela"> de componentes/relatoria_residual/index.html
VERSAO_TELA = "4"

_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class PonteResidual(PonteBase):
    """Equivalente à `ResidualsBridge` de residuals_page.py (mesmo código;
    só o que era do Qt foi trocado — ver comentários "[Streamlit]")."""

    ARQUIVOS = {"pickData": {"accept": ".csv,.xlsx,.xls,.sav"}}

    def __init__(self):
        super().__init__()
        self.df: Optional[pd.DataFrame] = None
        self.file_label: str = ""
        self.var_labels: Dict[str, str] = {}
        self.value_labels: Dict[str, Dict[Any, Any]] = {}
        self.missing_ranges: Dict[str, List[Any]] = {}
        self.queue: List[Dict[str, Any]] = []
        # [Streamlit]
        self.ordenar_colunas_por_codigo = True   # item 9: categorias das colunas na ordem dos códigos
        self.ui_state_json = ""                  # seleções, nomes, modo, prévia (ao reabrir a página)
        self.display_map: Dict[str, str] = {}    # nomes ajustados na tela (para exportar pela barra lateral)
        self.versao_fila = 0
        self._cache_xlsx = None

    # ---------------- [Streamlit] fila: identidade, pop-ups, exportação pela barra lateral
    @staticmethod
    def assinatura(item: Dict[str, Any]) -> str:
        cols = ",".join(str(b.get("col_var") or "") for b in (item.get("blocks") or []))
        return f"{item.get('row_var') or ''}|{cols}|{item.get('weight_var') or ''}"

    def _mudou_fila(self) -> None:
        self.versao_fila = getattr(self, "versao_fila", 0) + 1

    def removerDaFila(self, sig: str) -> None:
        for i, it in enumerate(self.queue):
            if self.assinatura(it) == sig:
                self.queue.pop(i)
                self._mudou_fila()
                self.notificar(f"Removido da fila: {it.get('title', '')}", "🗑️")
                self._push_queue_state(message="Removido.")
                return

    def saveUiState(self, estado_json: str) -> None:
        self.ui_state_json = estado_json or ""
        try:
            dm = json.loads(estado_json or "{}").get("dispMap") or {}
            if isinstance(dm, dict):
                self.display_map = {str(k): str(v) for k, v in dm.items()}
        except Exception:  # noqa: BLE001
            pass

    def xlsx_da_fila(self) -> bytes:
        chave = (getattr(self, "versao_fila", 0), json.dumps(getattr(self, "display_map", {}), sort_keys=True))
        cache = getattr(self, "_cache_xlsx", None)
        if not cache or cache[0] != chave:
            tmp_dir = tempfile.mkdtemp(prefix="ar_fila_")
            try:
                caminho = os.path.join(tmp_dir, "analise_residual.xlsx")
                self._export_tables(self.queue, dict(getattr(self, "display_map", {})), caminho)
                with open(caminho, "rb") as f:
                    cache = (chave, f.read())
            finally:
                shutil.rmtree(tmp_dir, ignore_errors=True)
            self._cache_xlsx = cache
        return cache[1]

    def _ordenar_colunas(self, base_cols: List[Any], col_map: Dict[Any, str]) -> List[Any]:
        """[Streamlit] item 9: categorias na ordem dos CÓDIGOS (como no SPSS),
        no lugar da ordem em que aparecem no banco."""
        def chave(code):
            try:
                return (0, float(code), "")
            except (TypeError, ValueError):
                return (1, 0.0, _canon(_safe_str(code)))
        ordem: List[Any] = []
        for code in sorted(col_map.keys(), key=chave):
            lab = col_map[code]
            if lab in base_cols and lab not in ordem:
                ordem.append(lab)
        return ordem + [c for c in base_cols if c not in ordem]

    def _push_queue_state(self, message: str = "", reset_row: bool = False) -> None:
        items = []
        for it in self.queue:
            items.append({
                "title": str(it.get("title") or ""),
                "sig": self.assinatura(it),  # [Streamlit] para a tela saber o que já está na fila
            })
        payload = {
            "items": items,
            "message": message,
            "resetRowSelection": bool(reset_row),
        }
        self.run_js(
            f"window.__setQueueState({json.dumps(json.dumps(payload, ensure_ascii=False))});"
        )

    def _apply_value_labels_for_var(self, var: str, s: pd.Series) -> Tuple[pd.Series, Dict[Any, str]]:
        labels = self.value_labels.get(var) or {}
        out = s.copy()
        label_map: Dict[Any, str] = {}

        def _label_of(v: Any) -> str:
            if pd.isna(v):
                return ""
            if v in labels:
                return _safe_str(labels[v])
            try:
                fv = float(v)
                if fv in labels:
                    return _safe_str(labels[fv])
                iv = int(fv)
                if iv in labels:
                    return _safe_str(labels[iv])
            except Exception:
                pass
            return _safe_str(v)

        out = s.map(_label_of)
        for raw in s.dropna().unique().tolist():
            label_map[raw] = _label_of(raw)

        return out, label_map

    def _missing_specs_for_var(self, var: str) -> List[Any]:
        raw = self.missing_ranges.get(str(var), []) or []
        if isinstance(raw, list):
            return raw
        return [raw]

    def _values_match_for_missing(self, value: Any, target: Any) -> bool:
        try:
            if pd.isna(value) or pd.isna(target):
                return False
        except Exception:
            pass

        if isinstance(value, str) or isinstance(target, str):
            v = _safe_str(value)
            t = _safe_str(target)
            return (v == t) or (_canon(v) == _canon(t))

        try:
            return float(value) == float(target)
        except Exception:
            return _safe_str(value) == _safe_str(target)

    def _is_spss_user_missing(self, var: str, value: Any) -> bool:
        try:
            if pd.isna(value):
                return False
        except Exception:
            pass

        specs = self._missing_specs_for_var(var)
        if not specs:
            return False

        for spec in specs:
            if isinstance(spec, dict):
                values = spec.get("values")
                if isinstance(values, (list, tuple, set)):
                    for vv in values:
                        if self._values_match_for_missing(value, vv):
                            return True

                lo = spec.get("lo")
                hi = spec.get("hi")
                if lo is None and hi is None and "value" in spec:
                    lo = spec.get("value")
                    hi = spec.get("value")
            elif isinstance(spec, (list, tuple)) and len(spec) >= 2:
                lo, hi = spec[0], spec[1]
            else:
                lo = spec
                hi = spec

            if lo is None and hi is None:
                continue
            if hi is None:
                hi = lo
            if lo is None:
                lo = hi

            if isinstance(value, str) or isinstance(lo, str) or isinstance(hi, str):
                if self._values_match_for_missing(value, lo) or self._values_match_for_missing(value, hi):
                    return True
                continue

            try:
                fv = float(value)
                flo = float(lo)
                fhi = float(hi)
                if flo <= fv <= fhi:
                    return True
            except Exception:
                if self._values_match_for_missing(value, lo) or self._values_match_for_missing(value, hi):
                    return True

        return False

    def _user_missing_mask(self, s: pd.Series, var: str) -> pd.Series:
        specs = self._missing_specs_for_var(var)
        if not specs:
            return pd.Series(False, index=s.index, dtype=bool)
        return s.map(lambda x: bool(self._is_spss_user_missing(var, x))).astype(bool)

    def _weighted_crosstab(self, df: pd.DataFrame, row_var: str, col_var: str, w: pd.Series) -> pd.DataFrame:
        tmp = pd.DataFrame({
            "_row": df[row_var],
            "_col": df[col_var],
            "_w": w,
        })
        ct = pd.pivot_table(
            tmp,
            index="_row",
            columns="_col",
            values="_w",
            aggfunc="sum",
            fill_value=0.0,
            dropna=False,
        )
        ct = ct.sort_index(axis=0).sort_index(axis=1)
        return ct

    def _unweighted_bases(self, df: pd.DataFrame, row_var: str, col_var: str) -> Tuple[List[Any], List[Any]]:
        base_rows = [x for x in pd.unique(df[row_var].dropna())]
        base_cols = [x for x in pd.unique(df[col_var].dropna())]
        return base_rows, base_cols

    def _sort_rows_by_total_desc(self, ct: pd.DataFrame, row_display_labels: List[str]) -> List[Any]:
        row_totals = ct.sum(axis=1)
        aux = []
        for i, code in enumerate(ct.index.tolist()):
            lab = row_display_labels[i] if i < len(row_display_labels) else _safe_str(code)
            aux.append((code, float(row_totals.loc[code]), lab))

        aux.sort(key=lambda t: (-t[1], _canon(t[2])))

        ordered_labels = [t[2] for t in aux]
        ordered_labels = _apply_fixed_bottom_order(ordered_labels, FIXED_BOTTOM_LABELS)

        final_codes = []
        used = set()
        for lab in ordered_labels:
            for code, _, lb in aux:
                if lb == lab and code not in used:
                    final_codes.append(code)
                    used.add(code)
                    break

        for code, _, _ in aux:
            if code not in used:
                final_codes.append(code)

        return final_codes

    def _prepare_pair(
        self,
        row_var: str,
        col_var: str,
        weight_var: str,
    ) -> Tuple[pd.DataFrame, pd.Series, Dict[Any, str], Dict[Any, str]]:
        assert self.df is not None

        cols = [row_var, col_var, weight_var]
        df = self.df[cols].copy()

        for v in [row_var, col_var, weight_var]:
            if v in df.columns:
                mask = self._user_missing_mask(df[v], v)
                if bool(mask.any()):
                    df.loc[mask, v] = np.nan

        df = df.dropna(subset=[row_var, col_var, weight_var])

        w = pd.to_numeric(df[weight_var], errors="coerce").fillna(0.0).astype(float)
        mask = (w > 0)
        df = df.loc[mask, [row_var, col_var]].copy()
        w = w.loc[mask]

        df[row_var], row_map = self._apply_value_labels_for_var(row_var, df[row_var])
        df[col_var], col_map = self._apply_value_labels_for_var(col_var, df[col_var])

        return df, w, row_map, col_map

    def _compute_block(
        self,
        row_var: str,
        col_var: str,
        weight_var: str,
        pct_mode: str,
    ) -> Dict[str, Any]:
        dfp, w, row_map, col_map = self._prepare_pair(row_var, col_var, weight_var)

        if dfp.empty:
            return {
                "col_var": col_var,
                "row_cats": [],
                "col_cats": [],
                "percent": [],
                "z": [],
                "debug_row_codes": [],
                "debug_col_codes": [],
                "debug_row_display": [],
                "debug_col_display": [],
                "debug_obs": [],
                "debug_exp": [],
            }

        ct = self._weighted_crosstab(dfp, row_var, col_var, w)

        base_rows, base_cols = self._unweighted_bases(dfp, row_var, col_var)
        if getattr(self, "ordenar_colunas_por_codigo", False):  # [Streamlit] item 9
            base_cols = self._ordenar_colunas(base_cols, col_map)
        ct = ct.reindex(index=base_rows, columns=base_cols, fill_value=0.0)

        row_disp = [row_map.get(code, _safe_str(code)) for code in ct.index]
        ordered_codes = self._sort_rows_by_total_desc(ct, row_disp)
        ct = ct.reindex(index=ordered_codes, fill_value=0.0)

        ct = _round_obs_to_integer(ct)

        row_disp = [row_map.get(code, _safe_str(code)) for code in ct.index]
        col_disp = [col_map.get(code, _safe_str(code)) for code in ct.columns]

        pct = _pct_matrix(ct, pct_mode)
        z = _adjusted_residuals(ct.to_numpy(dtype=float))

        return {
            "col_var": col_var,
            "row_cats": row_disp,
            "col_cats": col_disp,
            "percent": pct.values.tolist(),
            "z": np.round(z, 6).tolist(),
            "debug_row_codes": [_safe_str(code) for code in ct.index],
            "debug_col_codes": [_safe_str(code) for code in ct.columns],
            "debug_row_display": row_disp,
            "debug_col_display": col_disp,
            "debug_obs": ct.to_numpy(dtype=float).tolist(),
            "debug_exp": _expected_from_ct(ct).tolist(),
        }

    def _reorder_block_rows_to_match(self, block: Dict[str, Any], ordered_row_labels: List[str]) -> Dict[str, Any]:
        old_rows = list(block.get("row_cats") or [])
        idx_map = {lab: i for i, lab in enumerate(old_rows)}

        col_count = len(block.get("col_cats") or [])

        def blank_num() -> List[float]:
            return [0.0] * col_count

        def blank_nan() -> List[Any]:
            return [None] * col_count

        new_percent = []
        new_z = []
        new_obs = []
        new_exp = []

        for lab in ordered_row_labels:
            idx = idx_map.get(lab)
            if idx is None:
                new_percent.append(blank_num())
                new_z.append(blank_nan())
                new_obs.append(blank_num())
                new_exp.append(blank_nan())
            else:
                new_percent.append((block.get("percent") or [])[idx] if idx < len(block.get("percent") or []) else blank_num())
                new_z.append((block.get("z") or [])[idx] if idx < len(block.get("z") or []) else blank_nan())
                new_obs.append((block.get("debug_obs") or [])[idx] if idx < len(block.get("debug_obs") or []) else blank_num())
                new_exp.append((block.get("debug_exp") or [])[idx] if idx < len(block.get("debug_exp") or []) else blank_nan())

        block["row_cats"] = ordered_row_labels[:]
        block["percent"] = new_percent
        block["z"] = new_z
        block["debug_obs"] = new_obs
        block["debug_exp"] = new_exp
        return block

    def pickData(self, arquivo=None) -> None:
        # [Streamlit] no lugar do QFileDialog: arquivo enviado pelo navegador,
        # gravado num temporário com a mesma extensão (o resto é o original)
        if not arquivo:
            return
        tmp_dir = tempfile.mkdtemp(prefix="ar_")
        path = os.path.join(tmp_dir, os.path.basename(arquivo["name"]) or "banco")
        with open(path, "wb") as f:
            f.write(arquivo["dados"])
        try:

            self.queue = []
            self.ui_state_json = ""        # [Streamlit] banco novo: tela começa do zero
            self.display_map = {}
            self._mudou_fila()
            self.var_labels = {}
            self.value_labels = {}
            self.missing_ranges = {}

            if path.lower().endswith(".csv"):
                df = pd.read_csv(path)
            elif path.lower().endswith(".xlsx"):
                df = pd.read_excel(path)
            elif path.lower().endswith(".xls"):
                df = pd.read_excel(path)
            elif path.lower().endswith(".sav"):
                from core.pyreadstat_seguro import pyreadstat  # [Streamlit] lê .sav mesmo com a DLL bloqueada

                df = None
                meta = None
                last_err = None

                file_enc = None
                try:
                    _, m0 = pyreadstat.read_sav(path, metadataonly=True)  # type: ignore
                    file_enc = (getattr(m0, "file_encoding", None) or "").strip() or None
                except Exception:
                    file_enc = None

                candidates = []
                for enc in [None, file_enc, "utf-8", "cp1252", "latin1", "iso-8859-1", "iso-8859-15", "utf-16le", "ucs-2le"]:
                    if enc and enc not in candidates:
                        candidates.append(enc)

                for enc in candidates:
                    try:
                        if enc is None:
                            df, meta = pyreadstat.read_sav(path, user_missing=True)  # type: ignore
                        else:
                            df, meta = pyreadstat.read_sav(path, encoding=enc, user_missing=True)  # type: ignore
                        break
                    except Exception as e:
                        last_err = e
                        msg = str(e)
                        if (
                            "Unable to convert string" in msg
                            or "invalid byte sequence" in msg
                            or "codec can't decode" in msg
                            or "invalid continuation byte" in msg
                        ):
                            continue
                        raise

                if df is None or meta is None:
                    raise last_err or Exception("Falha ao ler o arquivo .sav.")

                try:
                    m = getattr(meta, "column_names_to_labels", None)
                    if isinstance(m, dict):
                        self.var_labels = {str(k): str(v) for k, v in m.items() if v}
                except Exception:
                    self.var_labels = {}

                try:
                    for var, m in (meta.variable_value_labels or {}).items():
                        self.value_labels[str(var)] = dict(m)
                except Exception:
                    self.value_labels = {}

                try:
                    raw_missing = getattr(meta, "missing_ranges", None) or {}
                    self.missing_ranges = {str(var): list(specs) for var, specs in raw_missing.items()}
                except Exception:
                    self.missing_ranges = {}
            else:
                raise RuntimeError("Formato não suportado. Use CSV, XLSX ou SAV.")

            df = df.dropna(axis=1, how="all")
            self.df = df
            self.file_label = os.path.basename(path)

            cols = [{"name": str(c), "dtype": str(df[c].dtype)} for c in df.columns]
            weights = [str(c) for c in df.columns if _is_peso_name(str(c))]

            payload = {
                "label": self.file_label,
                "columns": cols,
                "labels": self.var_labels,
                "weight_candidates": weights,
            }

            self.run_js(
                f"window.__setData({json.dumps(json.dumps(payload, ensure_ascii=False))});"
            )
            self._push_queue_state(message="Fila resetada.")

        except Exception as e:
            self.aviso(f"Falha ao carregar dados.\n\n{e}", "erro")
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    def computePreview(self, payload_json: str) -> None:
        try:
            if self.df is None:
                raise RuntimeError("Nenhum arquivo carregado.")

            payload = json.loads(payload_json) if payload_json else {}
            mode = str(payload.get("mode") or "batch").strip().lower()
            row_vars = payload.get("rowVars") or []
            col_vars = payload.get("colVars") or []
            weight_var = str(payload.get("weightVar") or "").strip()
            pct_mode = str(payload.get("pctMode") or "col").strip().lower()
            decimals = int(payload.get("decimals") or 1)

            if not isinstance(row_vars, list) or not isinstance(col_vars, list):
                raise RuntimeError("Seleção inválida.")
            row_vars = [str(x) for x in row_vars if str(x) in self.df.columns]
            col_vars = [str(x) for x in col_vars if str(x) in self.df.columns]

            if not weight_var or weight_var not in self.df.columns:
                raise RuntimeError("Selecione um peso válido.")
            if not row_vars or not col_vars:
                raise RuntimeError("Selecione ao menos 1 LINHA e 1 COLUNA.")
            if mode == "queue" and len(row_vars) != 1:
                raise RuntimeError("No modo fila, selecione 1 variável LINHA.")

            tables: List[Dict[str, Any]] = []
            for rv in row_vars:
                valid_cols = [cv for cv in col_vars if cv != rv]
                blocks = [self._compute_block(rv, cv, weight_var, pct_mode) for cv in valid_cols]

                if blocks:
                    ordered_rows = list(blocks[0].get("row_cats") or [])
                    blocks = [self._reorder_block_rows_to_match(b, ordered_rows) for b in blocks]

                title = f"A.R {self.var_labels.get(rv, rv) or rv}"
                tables.append({
                    "row_var": rv,
                    "weight_var": weight_var,
                    "decimals": decimals,
                    "title": title,
                    "blocks": blocks,
                })

            out = {"ok": True, "tables": tables, "decimals": decimals}

        except Exception as e:
            out = {"ok": False, "error": str(e)}

        self.run_js(
            f"window.__setPreview({json.dumps(json.dumps(out, ensure_ascii=False))});"
        )

    def addToQueueFromPreview(self, table_json: str) -> None:
        try:
            t = json.loads(table_json) if table_json else {}
            if not isinstance(t, dict):
                raise RuntimeError("Tabela inválida.")

            title = str(t.get("title") or "").strip()
            if not title:
                title = "A.R"
            if not title.upper().startswith("A.R"):
                title = "A.R " + title

            item = {
                "title": title,
                "row_var": t.get("row_var"),
                "weight_var": t.get("weight_var"),
                "decimals": int(t.get("decimals") or 1),
                "blocks": t.get("blocks") or [],
            }

            # [Streamlit] mesma tabela (linha, colunas e peso) já na fila: substitui se pedido
            sig = self.assinatura(item)
            pos = next((i for i, it in enumerate(self.queue) if self.assinatura(it) == sig), None)
            if pos is not None and not t.get("substituir"):
                self.notificar(f"Essa tabela já está na fila (item {pos + 1}): {self.queue[pos]['title']}.", "⚠️")
                self._push_queue_state(message=f"Já está na fila (item {pos + 1}).")
                return
            if pos is not None:
                self.queue[pos] = item
                self.notificar(f"Tabela substituída na fila (item {pos + 1}): {title}.", "🔁")
                msg = "Substituída na fila."
            else:
                self.queue.append(item)
                self.notificar(f"Adicionado à fila: {title} — {len(self.queue)} tabela(s) na fila.", "✅")
                msg = "Adicionado à fila."
            self._mudou_fila()
            self._push_queue_state(message=msg, reset_row=True)

        except Exception as e:
            self._push_queue_state(message=f"Erro: {e}")

    def removeQueueItem(self, idx: int) -> None:
        try:
            idx = int(idx)
            if 0 <= idx < len(self.queue):
                self.queue.pop(idx)
                self._mudou_fila()
            self._push_queue_state(message="Removido.")
        except Exception as e:
            self._push_queue_state(message=f"Erro: {e}")

    def clearQueue(self) -> None:
        if self.queue:
            self.notificar("Fila esvaziada.", "🗑️")
        self.queue = []
        self._mudou_fila()
        self._push_queue_state(message="Fila limpa.")

    def _export_tables(self, tables: List[Dict[str, Any]], display_map: Dict[str, str], out_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        used_names = set()

        font_title = Font(name="DIN", size=15, bold=True, color="000000")
        font_head = Font(name="DIN", size=11, bold=True, color="000000")
        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin_black = Side(style="thin", color="000000")
        medium_black = Side(style="medium", color="000000")

        fill_head = PatternFill(fill_type="solid", fgColor="F2F2F2")
        fill_sub = PatternFill(fill_type="solid", fgColor="EAF3FF")

        def disp(v: str) -> str:
            v = str(v)
            return (display_map.get(v) or self.var_labels.get(v) or v)

        def apply_border(cell, left=None, right=None, top=None, bottom=None):
            cell.border = Border(
                left=left or thin_black,
                right=right or thin_black,
                top=top or thin_black,
                bottom=bottom or thin_black,
            )

        for t in tables:
            title = str(t.get("title") or "A.R").strip()
            if not title.upper().startswith("A.R"):
                title = "A.R " + title

            sheet_name = _unique_sheet_name(title, used_names)
            ws = wb.create_sheet(sheet_name)

            blocks = t.get("blocks") or []
            decimals = int(t.get("decimals") or 1)
            row_var = str(t.get("row_var") or "")
            row_labels = list((blocks[0].get("row_cats") if blocks else []) or [])

            ws.cell(1, 1, title)
            ws.cell(1, 1).font = font_title
            ws.cell(1, 1).alignment = align_left

            if not blocks:
                continue

            start_row = 3

            ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 1, end_column=1)
            c = ws.cell(start_row, 1, disp(row_var))
            c.font = font_head
            c.alignment = align_center
            c.fill = fill_head
            apply_border(c, left=medium_black, right=medium_black, top=medium_black, bottom=medium_black)

            col_cursor = 2
            total_data_cols = 0
            for b in blocks:
                col_cats = b.get("col_cats") or []
                span = max(1, len(col_cats))
                total_data_cols += span

                ws.merge_cells(start_row=start_row, start_column=col_cursor, end_row=start_row, end_column=col_cursor + span - 1)
                top_cell = ws.cell(start_row, col_cursor, disp(str(b.get("col_var") or "")))
                top_cell.font = font_head
                top_cell.alignment = align_center
                top_cell.fill = fill_sub

                for cc in range(col_cursor, col_cursor + span):
                    apply_border(
                        ws.cell(start_row, cc),
                        left=medium_black if cc == col_cursor else thin_black,
                        right=medium_black if cc == (col_cursor + span - 1) else thin_black,
                        top=medium_black,
                        bottom=medium_black
                    )

                for j, cc_label in enumerate(col_cats, start=col_cursor):
                    cell = ws.cell(start_row + 1, j, str(cc_label))
                    cell.font = font_head
                    cell.alignment = align_center
                    cell.fill = fill_head
                    apply_border(
                        cell,
                        left=medium_black if j == col_cursor else thin_black,
                        right=medium_black if j == (col_cursor + span - 1) else thin_black,
                        top=thin_black,
                        bottom=medium_black
                    )

                col_cursor += span

            row_cursor = start_row + 2

            for i, row_lab in enumerate(row_labels):
                tr_pct = row_cursor
                tr_z = row_cursor + 1

                ws.merge_cells(start_row=tr_pct, start_column=1, end_row=tr_z, end_column=1)
                c0 = ws.cell(tr_pct, 1, str(row_lab))
                c0.font = font_head
                c0.alignment = align_left
                c0.fill = fill_head
                apply_border(
                    c0,
                    left=medium_black,
                    right=medium_black,
                    top=medium_black if i == 0 else thin_black,
                    bottom=medium_black if i == len(row_labels) - 1 else thin_black
                )

                col_cursor = 2
                for b in blocks:
                    pct = b.get("percent") or []
                    z = b.get("z") or []
                    col_cats = b.get("col_cats") or []

                    start_block_col = col_cursor
                    end_block_col = col_cursor + len(col_cats) - 1

                    for j in range(len(col_cats)):
                        cellp = ws.cell(tr_pct, col_cursor)
                        cellz = ws.cell(tr_z, col_cursor)

                        cellp.alignment = align_center
                        cellp.font = font_base
                        cellz.alignment = align_center
                        cellz.font = font_base

                        apply_border(
                            cellp,
                            left=medium_black if col_cursor == start_block_col else thin_black,
                            right=medium_black if col_cursor == end_block_col else thin_black,
                            top=medium_black if i == 0 else thin_black,
                            bottom=thin_black
                        )
                        apply_border(
                            cellz,
                            left=medium_black if col_cursor == start_block_col else thin_black,
                            right=medium_black if col_cursor == end_block_col else thin_black,
                            top=thin_black,
                            bottom=medium_black if i == len(row_labels) - 1 else thin_black
                        )

                        try:
                            pv = float((pct[i] or [])[j])
                            cellp.value = pv
                            cellp.number_format = '0.0"%"'
                        except Exception:
                            cellp.value = None

                        try:
                            zv = float((z[i] or [])[j])
                            cellz.value = round(float(zv), decimals)
                            cellz.number_format = "0" if decimals <= 0 else ("0." + ("0" * decimals))

                            st = _cell_style_for_z(zv)
                            if st is not None:
                                fill_hex, font_hex = st
                                cellz.fill = PatternFill("solid", fgColor=fill_hex)
                                cellz.font = Font(name="DIN", size=10, bold=True, color=font_hex)

                                cellp.fill = PatternFill("solid", fgColor=fill_hex)
                                cellp.font = Font(name="DIN", size=10, bold=True, color=font_hex)
                        except Exception:
                            cellz.value = None

                        col_cursor += 1

                row_cursor += 2

            max_col = 1 + total_data_cols
            for c in range(1, max_col + 1):
                ws.column_dimensions[get_column_letter(c)].width = 16 if c == 1 else 12

            for r in range(1, ws.max_row + 1):
                ws.row_dimensions[r].height = 22

            ws.freeze_panes = "B5"

        wb.save(out_path)

    # ============================================================
    # DEBUG-ONLY START
    # ============================================================
    def _build_case_debug_for_pair(
        self,
        row_var: str,
        col_var: str,
        weight_var: str,
    ) -> pd.DataFrame:
        assert self.df is not None

        cols = [row_var, col_var, weight_var]
        df = self.df[cols].copy()

        df["_row_isna_orig"] = df[row_var].isna()
        df["_col_isna_orig"] = df[col_var].isna()
        df["_w_isna_orig"] = df[weight_var].isna()

        row_user_missing = self._user_missing_mask(df[row_var], row_var)
        col_user_missing = self._user_missing_mask(df[col_var], col_var)
        weight_user_missing = self._user_missing_mask(df[weight_var], weight_var)

        row_series, row_map = self._apply_value_labels_for_var(row_var, df[row_var])
        col_series, col_map = self._apply_value_labels_for_var(col_var, df[col_var])
        _ = row_series, col_series

        weight_clean = df[weight_var].mask(weight_user_missing)
        w_num = pd.to_numeric(weight_clean, errors="coerce")

        out = pd.DataFrame({
            "case_index": df.index,
            "row_code": df[row_var],
            "row_label": [_safe_label_from_map(row_map, v) for v in df[row_var]],
            "col_code": df[col_var],
            "col_label": [_safe_label_from_map(col_map, v) for v in df[col_var]],
            "weight_raw": df[weight_var],
            "weight_num": w_num,
            "row_isna_orig": df["_row_isna_orig"],
            "col_isna_orig": df["_col_isna_orig"],
            "weight_isna_orig": df["_w_isna_orig"],
            "row_is_user_missing": row_user_missing,
            "col_is_user_missing": col_user_missing,
            "weight_is_user_missing": weight_user_missing,
        })

        included = (
            (~out["row_isna_orig"])
            & (~out["col_isna_orig"])
            & (~out["weight_isna_orig"])
            & (~out["row_is_user_missing"])
            & (~out["col_is_user_missing"])
            & (~out["weight_is_user_missing"])
            & (pd.to_numeric(out["weight_num"], errors="coerce").fillna(0.0) > 0)
        )
        out["included"] = included

        reasons = []
        for _, r in out.iterrows():
            if bool(r["included"]):
                reasons.append("IN")
            elif bool(r["row_is_user_missing"]):
                reasons.append("OUT: row user-missing")
            elif bool(r["col_is_user_missing"]):
                reasons.append("OUT: col user-missing")
            elif bool(r["weight_is_user_missing"]):
                reasons.append("OUT: weight user-missing")
            elif bool(r["row_isna_orig"]):
                reasons.append("OUT: row NA")
            elif bool(r["col_isna_orig"]):
                reasons.append("OUT: col NA")
            elif bool(r["weight_isna_orig"]):
                reasons.append("OUT: weight NA")
            else:
                try:
                    wn = float(r["weight_num"])
                    if not np.isfinite(wn):
                        reasons.append("OUT: weight invalid")
                    elif wn <= 0:
                        reasons.append("OUT: weight <= 0")
                    else:
                        reasons.append("OUT: unknown")
                except Exception:
                    reasons.append("OUT: weight invalid")
        out["include_reason"] = reasons

        return out

    def _write_debug_matrix_sheet(self, ws, title: str, row_codes: List[str], col_codes: List[str], matrix: List[List[Any]], decimals: int) -> None:
        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        font_head = Font(name="DIN", size=10, bold=True, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        num_fmt = "0." + ("0" * decimals) if decimals > 0 else "0"

        ws["A1"] = title
        ws["A1"].font = font_head
        ws["A1"].alignment = align_left

        ws.cell(3, 1, "ROW\\COL").font = font_head
        ws.cell(3, 1).alignment = align_center
        ws.cell(3, 1).border = border

        for j, cc in enumerate(col_codes, start=2):
            cell = ws.cell(3, j, cc)
            cell.font = font_head
            cell.alignment = align_center
            cell.border = border

        for i, rc in enumerate(row_codes, start=4):
            cell = ws.cell(i, 1, rc)
            cell.font = font_head
            cell.alignment = align_left
            cell.border = border

            row = matrix[i - 4] if (i - 4) < len(matrix) else []
            for j, v in enumerate(row, start=2):
                cell = ws.cell(i, j)
                cell.border = border
                cell.alignment = align_center
                cell.font = font_base

                try:
                    fv = float(v)
                    if np.isfinite(fv):
                        cell.value = fv
                        cell.number_format = num_fmt
                    else:
                        cell.value = None
                except Exception:
                    cell.value = None

        max_col = max(1, ws.max_column)
        max_row = max(1, ws.max_row)

        for c in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(c)].width = 12
        for r in range(1, max_row + 1):
            ws.row_dimensions[r].height = 22

        ws.freeze_panes = "B4"

    def _export_debug_tables(self, tables: List[Dict[str, Any]], debug_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        sheet_idx = 1

        for t in tables:
            row_var = str(t.get("row_var") or "")
            blocks = t.get("blocks") or []

            for b in blocks:
                col_var = str(b.get("col_var") or "")

                row_codes = b.get("debug_row_codes") or []
                col_codes = b.get("debug_col_codes") or []
                row_disp = b.get("debug_row_display") or []
                col_disp = b.get("debug_col_display") or []
                obs = b.get("debug_obs") or []
                exp = b.get("debug_exp") or []
                z = b.get("z") or []

                tag = f"{sheet_idx:02d}"

                ws_map = wb.create_sheet(f"MAP {tag}")
                ws_map["A1"] = f"Bloco {tag}"
                ws_map["A2"] = f"Linha: {row_var}"
                ws_map["A3"] = f"Coluna: {col_var}"

                ws_map["A6"] = "ROW_CODE"
                ws_map["B6"] = "ROW_DISPLAY"
                for i, (rc, rd) in enumerate(zip(row_codes, row_disp), start=7):
                    ws_map.cell(i, 1, rc)
                    ws_map.cell(i, 2, rd)

                ws_map["D6"] = "COL_CODE"
                ws_map["E6"] = "COL_DISPLAY"
                for i, (cc, cd) in enumerate(zip(col_codes, col_disp), start=7):
                    ws_map.cell(i, 4, cc)
                    ws_map.cell(i, 5, cd)

                for c in range(1, 6):
                    ws_map.column_dimensions[get_column_letter(c)].width = 22

                ws_obs = wb.create_sheet(f"OBS {tag}")
                self._write_debug_matrix_sheet(
                    ws_obs,
                    f"Observado ponderado | {row_var} x {col_var}",
                    row_codes,
                    col_codes,
                    obs,
                    decimals=6
                )

                ws_exp = wb.create_sheet(f"EXP {tag}")
                self._write_debug_matrix_sheet(
                    ws_exp,
                    f"Esperado | {row_var} x {col_var}",
                    row_codes,
                    col_codes,
                    exp,
                    decimals=6
                )

                ws_z = wb.create_sheet(f"Z {tag}")
                self._write_debug_matrix_sheet(
                    ws_z,
                    f"Residual Ajustado | {row_var} x {col_var}",
                    row_codes,
                    col_codes,
                    z,
                    decimals=6
                )

                sheet_idx += 1

        wb.save(debug_path)

    def _write_case_debug_sheet(self, ws, title: str, df_cases: pd.DataFrame) -> None:
        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        font_head = Font(name="DIN", size=10, bold=True, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        ws["A1"] = title
        ws["A1"].font = font_head
        ws["A1"].alignment = align_left

        headers = list(df_cases.columns)
        for j, h in enumerate(headers, start=1):
            cell = ws.cell(3, j, h)
            cell.font = font_head
            cell.alignment = align_center
            cell.border = border

        for i, row in enumerate(df_cases.itertuples(index=False), start=4):
            for j, v in enumerate(row, start=1):
                cell = ws.cell(i, j, v)
                cell.font = font_base
                cell.alignment = align_left if j <= 8 else align_center
                cell.border = border

        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16
        ws.freeze_panes = "A4"

    def _export_case_level_debug(self, tables: List[Dict[str, Any]], debug_cases_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        sheet_idx = 1

        for t in tables:
            row_var = str(t.get("row_var") or "")
            weight_var = str(t.get("weight_var") or "")
            blocks = t.get("blocks") or []

            for b in blocks:
                col_var = str(b.get("col_var") or "")
                tag = f"{sheet_idx:02d}"

                df_cases = self._build_case_debug_for_pair(
                    row_var,
                    col_var,
                    weight_var,
                )

                ws_cases = wb.create_sheet(f"CASES {tag}")
                self._write_case_debug_sheet(
                    ws_cases,
                    f"Casos usados | {row_var} x {col_var} | peso={weight_var}",
                    df_cases
                )

                ws_group = wb.create_sheet(f"GROUP {tag}")
                grouped = (
                    df_cases[df_cases["included"] == True]
                    .groupby(["row_code", "row_label", "col_code", "col_label"], dropna=False)["weight_num"]
                    .sum()
                    .reset_index()
                    .rename(columns={"weight_num": "weighted_sum"})
                )
                self._write_case_debug_sheet(
                    ws_group,
                    f"Resumo agrupado | {row_var} x {col_var} | peso={weight_var}",
                    grouped
                )

                sheet_idx += 1

        wb.save(debug_cases_path)
    # ============================================================
    # DEBUG-ONLY END
    # ============================================================

    def exportXlsx(self, payload_json: str) -> None:
        try:
            if self.df is None:
                raise RuntimeError("Nenhum arquivo carregado.")

            payload = json.loads(payload_json) if payload_json else {}
            mode = str(payload.get("mode") or "batch").strip().lower()
            display_map = payload.get("displayMap") or {}
            preview_tables = payload.get("previewTables") or []

            if not isinstance(display_map, dict):
                display_map = {}

            # [Streamlit] no lugar do "Salvar como": pasta temporária + download
            tmp_dir = tempfile.mkdtemp(prefix="ar_exp_")
            save_path = os.path.join(tmp_dir, "analise_residual.xlsx")

            if mode == "queue":
                if not self.queue:
                    raise RuntimeError("Fila vazia.")
                tables = self.queue
            else:
                if not isinstance(preview_tables, list) or not preview_tables:
                    raise RuntimeError("No modo 1, gere a prévia antes de exportar.")
                tables = preview_tables

            self._export_tables(tables, display_map, save_path)
            arquivos = [save_path]

            # [Streamlit] item 8: arquivos de conferência só quando pedidos
            incluir_debug = bool(payload.get("incluirDebug", True))
            if incluir_debug:
                debug_path = save_path[:-5] + "_DEBUG.xlsx"
                self._export_debug_tables(tables, debug_path)

                debug_cases_path = save_path[:-5] + "_CASES_DEBUG.xlsx"
                self._export_case_level_debug(tables, debug_cases_path)
                arquivos += [debug_path, debug_cases_path]

            for caminho in arquivos:
                with open(caminho, "rb") as f:
                    self.download(os.path.basename(caminho), f.read(), _MIME_XLSX)
            shutil.rmtree(tmp_dir, ignore_errors=True)

            if incluir_debug:
                self._push_queue_state(
                    message=(
                        "XLSX exportado com sucesso. "
                        f"Debugs salvos em: {os.path.basename(debug_path)} e {os.path.basename(debug_cases_path)}"
                    )
                )
            else:
                self._push_queue_state(message="XLSX exportado com sucesso.")

        except Exception as e:
            self._push_queue_state(message=f"Erro ao exportar: {e}")


    # [Streamlit] estado ao (re)abrir a página (no desktop a página não fechava)
    def estado_inicial(self) -> None:
        if self.df is not None:
            cols = [{"name": str(c), "dtype": str(self.df[c].dtype)} for c in self.df.columns]
            weights = [str(c) for c in self.df.columns if _is_peso_name(str(c))]
            payload = {"label": self.file_label, "columns": cols, "labels": self.var_labels,
                       "weight_candidates": weights}
            self.run_js(f"window.__setData({json.dumps(json.dumps(payload, ensure_ascii=False))});")
        self._push_queue_state()
        if self.df is not None and getattr(self, "ui_state_json", ""):
            self.run_js(f"window.__restoreUiState({json.dumps(self.ui_state_json)});")


def _itens_fila(ponte: PonteResidual) -> List[Dict[str, Any]]:
    dm = getattr(ponte, "display_map", {}) or {}

    def disp(v):
        v = str(v or "")
        return dm.get(v) or ponte.var_labels.get(v) or v

    itens = []
    for it in ponte.queue:
        blocos = it.get("blocks") or []
        itens.append({
            "sig": ponte.assinatura(it),
            "title": it.get("title") or "A.R",
            "chips": [{"rotulo": "Blocos", "valor": str(len(blocos))},
                      {"rotulo": "Peso", "valor": disp(it.get("weight_var"))}],
            "secoes": [
                {"rotulo": "Linha", "texto": disp(it.get("row_var")), "cor": "#60a5fa"},
                {"rotulo": "Colunas", "texto": " · ".join(disp(b.get("col_var")) for b in blocos), "cor": "#a78bfa"},
            ],
        })
    return itens


def _fila_na_barra_lateral(ponte: PonteResidual) -> None:
    pilha_lateral(
        ponte, prefixo="ar", itens=lambda: _itens_fila(ponte), remover=ponte.removerDaFila,
        titulo="Fila da A.R",
        vazio_msg="Vazia. No modo 2, gere a prévia e clique em “Adicionar à fila”.",
        exportar=lambda: {"rotulo": f"Exportar XLSX ({len(ponte.queue)})", "dados": ponte.xlsx_da_fila,
                          "arquivo": "analise_residual.xlsx", "mime": _MIME_XLSX},
        limpar=ponte.clearQueue, rotulo_limpar="Limpar fila",
    )


def modulo_analise_residual():
    st.header("Análise Residual (A.R)")
    botao_tutorial("analise_residual")

    ponte: PonteResidual = obter_ponte("ar_ponte", PonteResidual)
    processar_pedidos(ponte, "ar_componente")
    conferir_versao_tela(ponte, "relatoria_residual", VERSAO_TELA)
    _fila_na_barra_lateral(ponte)
    mostrar_notificacoes(ponte)
    # altura "auto": a ferramenta cresce com o conteúdo e a PÁGINA rola (sem tela dentro da tela)
    desenhar_componente(ponte, "relatoria_residual", altura="auto", key="ar_componente")
