"""
Matriz de Correlação — a tela ORIGINAL do Relatória Analítica dentro do
Streamlit.

A interface (HTML/CSS/JS) é a mesma do app desktop:
componentes/relatoria_correlacao/index.html (cópia de correlation_page.py,
só trocando o QWebChannel do Qt pela ponte). Ela conversa com a
`PonteCorrelacao` abaixo, que tem os mesmos métodos (@Slot) da
`CorrelationBridge` original, com o mesmo código. Os auxiliares (Pearson /
Spearman ponderados, Sig., cores do VBA) estão em core/correlacao_math.py,
copiados sem alteração.

O que muda em relação ao desktop:
- "Carregar dados" abre o seletor de arquivo do navegador;
- "Exportar CSV" / "Exportar XLSX" baixam o arquivo pelo navegador;
- os avisos (QMessageBox) aparecem como faixas no canto da ferramenta;
- integrada à página: sem título repetido, sem rolagem interna (a matriz
  acompanha a rolagem), painel na ordem do fluxo (Dados → Peso → Colunas →
  Missing → Nomes → Configurações → Exportar), textos inteiros, Todos/Nenhum
  em verde/vermelho com contador, peso escolhido sozinho quando só existe uma
  coluna com "peso" e o estado da tela guardado ao trocar de ferramenta;
- .sav é lido pelo pyreadstat ou, se ele estiver bloqueado no computador,
  pelo leitor em Python puro (core/sav_puro.py).
"""
from __future__ import annotations

import io
import json
import os
import tempfile
from typing import Dict, List, Optional

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analises._ponte_relatoria import (
    PonteBase, conferir_versao_tela, desenhar_componente, mostrar_notificacoes, obter_ponte, processar_pedidos,
)
from analises.tutoriais_analitica import botao_tutorial
from core.correlacao_math import (
    _corr_style_rgb, _format_token_for_ui, _is_peso_name, _jsonable_matrix, _looks_numeric_token,
    _numeric_candidates, _parse_missing_list, _sig_style_rgb, _token_to_float, _weighted_corr_and_sig_matrices,
    scipy_stats,
)

#: precisa bater com <meta name="versao-tela"> de componentes/relatoria_correlacao/index.html
VERSAO_TELA = "2"

_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _js(funcao: str, payload: dict) -> str:
    return f"{funcao}({json.dumps(json.dumps(payload, ensure_ascii=False))});"


def ler_base(conteudo: bytes, nome: str):
    """Mesma leitura do pickData original (CSV utf-8/latin1, XLSX, SAV com as
    tentativas de codificação), a partir do arquivo enviado pelo navegador.
    Devolve (df, rótulos das variáveis)."""
    ext = os.path.splitext(nome)[1].lower()
    var_labels: Dict[str, str] = {}

    if ext == ".csv":
        try:
            df = pd.read_csv(io.BytesIO(conteudo), encoding="utf-8")
        except UnicodeDecodeError:
            df = pd.read_csv(io.BytesIO(conteudo), encoding="latin1")

    elif ext == ".xlsx":
        df = pd.read_excel(io.BytesIO(conteudo))

    elif ext == ".sav":
        from core.pyreadstat_seguro import pyreadstat

        # o pyreadstat precisa de um caminho no disco
        tmp = tempfile.NamedTemporaryFile(suffix=".sav", delete=False)
        try:
            tmp.write(conteudo)
            tmp.close()
            path = tmp.name

            df = None
            meta = None
            last_err = None

            file_enc = None
            try:
                _, m0 = pyreadstat.read_sav(path, metadataonly=True)  # type: ignore
                file_enc = (getattr(m0, "file_encoding", None) or "").strip() or None
            except Exception:
                file_enc = None

            candidates = []
            for enc in [None, file_enc, "utf-8", "cp1252", "latin1", "iso-8859-1", "iso-8859-15", "utf-16le", "ucs-2le"]:
                if enc and enc not in candidates:
                    candidates.append(enc)

            for enc in candidates:
                try:
                    if enc is None:
                        df, meta = pyreadstat.read_sav(path)  # type: ignore
                    else:
                        df, meta = pyreadstat.read_sav(path, encoding=enc)  # type: ignore
                    break
                except Exception as e:  # noqa: BLE001
                    last_err = e
                    msg = str(e)
                    if (
                        "Unable to convert string" in msg
                        or "invalid byte sequence" in msg
                        or "codec can't decode" in msg
                        or "invalid continuation byte" in msg
                    ):
                        continue
                    raise

            if df is None or meta is None:
                raise last_err or Exception("Falha ao ler o arquivo .sav.")
        finally:
            try:
                os.unlink(tmp.name)
            except OSError:
                pass

        labels_map: Dict[str, str] = {}
        try:
            m = getattr(meta, "column_names_to_labels", None)
            if isinstance(m, dict):
                labels_map = {str(k): str(v) for k, v in m.items() if v}
            else:
                cn = getattr(meta, "column_names", None)
                cl = getattr(meta, "column_labels", None)
                if cn and cl and len(cn) == len(cl):
                    labels_map = {str(k): str(v) for k, v in zip(cn, cl) if v}
        except Exception:  # noqa: BLE001
            labels_map = {}

        var_labels = labels_map

    else:
        raise RuntimeError("Formato não suportado. Use CSV, XLSX ou SAV.")

    df = df.dropna(axis=1, how="all")
    return df, var_labels


class PonteCorrelacao(PonteBase):
    """Equivalente à `CorrelationBridge` de correlation_page.py."""

    ARQUIVOS = {"pickData": {"accept": ".csv,.xlsx,.sav"}}

    def __init__(self):
        super().__init__()
        self.df: Optional[pd.DataFrame] = None
        self.data_label: str = ""
        self.var_labels: Dict[str, str] = {}

        # último resultado
        self.last_cols_vars: List[str] = []
        self.last_matrix: Optional[np.ndarray] = None          # matriz filtrada (exibida/exportada)
        self.last_raw_matrix: Optional[np.ndarray] = None      # matriz bruta
        self.last_p_matrix: Optional[np.ndarray] = None        # Sig. (2-tailed)
        self.last_method: str = "pearson"
        self.last_decimals: int = 3
        self.last_weight_var: str = ""
        self.last_result_payload: Optional[dict] = None       # para reabrir a página
        self.ui_state_json: str = ""                           # colunas, missing, nomes, método...

    # ---------------- estado ao (re)abrir a página
    def estado_inicial(self) -> None:
        if self.df is not None:
            self._push_columns()
            if getattr(self, "ui_state_json", ""):
                self.run_js(f"window.__restoreUiState({json.dumps(self.ui_state_json)});")
            if self.last_result_payload:
                self.run_js(_js("window.__setCorrResult", self.last_result_payload))

    def saveUiState(self, estado_json: str) -> None:
        """Estado da tela (colunas, missing, nomes, método...) para não perder ao
        trocar de ferramenta no menu e voltar."""
        self.ui_state_json = estado_json or ""

    def _push_columns(self) -> None:
        if self.df is None:
            payload = {"columns": [], "label": "", "labels": {}, "weight_candidates": []}
        else:
            cols = _numeric_candidates(self.df)
            labels = {c: (self.var_labels.get(c) or "") for c in cols}
            weights = [str(c) for c in self.df.columns if _is_peso_name(str(c))]
            payload = {
                "columns": cols,
                "label": self.data_label,
                "labels": labels,
                "weight_candidates": weights,
            }
        self.run_js(_js("window.__setColumns", payload))

    # ---------------- @Slot pickData
    def pickData(self, arquivo=None) -> None:
        if not arquivo:
            return
        try:
            df, self.var_labels = ler_base(arquivo["dados"], arquivo["name"])
            self.df = df
            self.data_label = os.path.basename(arquivo["name"])

            # reset último resultado
            self.last_cols_vars = []
            self.last_matrix = None
            self.last_raw_matrix = None
            self.last_p_matrix = None
            self.last_method = "pearson"
            self.last_decimals = 3
            self.last_weight_var = ""
            self.last_result_payload = None
            self.ui_state_json = ""

            self._push_columns()
        except Exception as e:  # noqa: BLE001
            self.var_labels = {}
            self.aviso(f"Falha ao carregar dados.\n\n{e}", "erro")

    # ---------------- @Slot getUniqueValues
    def getUniqueValues(self, req_json: str) -> None:
        try:
            if self.df is None:
                raise RuntimeError("Nenhum arquivo carregado.")

            req = json.loads(req_json) if req_json else {}
            col = str(req.get("col") or "").strip()
            max_items = int(req.get("max") or 200)
            max_items = max(20, min(500, max_items))

            if not col or col not in self.df.columns:
                raise RuntimeError("Coluna inválida.")

            s = self.df[col]

            if pd.api.types.is_numeric_dtype(s):
                vc = s.dropna().value_counts().head(max_items)
                items = []
                for val, cnt in vc.items():
                    token, disp = _format_token_for_ui(val)
                    if token == "__NA__":
                        continue
                    items.append({"token": token, "display": disp, "count": int(cnt)})
            else:
                ss = s.copy()
                mask = ss.notna()
                ss2 = ss[mask].astype(str).map(lambda x: x.strip())
                ss2 = ss2.replace("", "__EMPTY__")
                vc = ss2.value_counts().head(max_items)
                items = []
                for val, cnt in vc.items():
                    if val == "__EMPTY__":
                        items.append({"token": "__EMPTY__", "display": "<vazio>", "count": int(cnt)})
                    else:
                        items.append({"token": str(val), "display": str(val), "count": int(cnt)})

            payload = {
                "ok": True,
                "col": col,
                "n_unique": int(s.dropna().nunique()) if int(s.notna().sum()) else 0,
                "values": items,
            }
        except Exception as e:  # noqa: BLE001
            col = ""
            try:
                col = str(json.loads(req_json).get("col", "")) if req_json else ""
            except Exception:  # noqa: BLE001
                col = ""
            payload = {"ok": False, "col": col, "error": str(e), "values": [], "n_unique": 0}

        self.run_js(_js("window.__setUniqueValues", payload))

    # ---------------- @Slot computeCorr
    def computeCorr(self, payload_json: str) -> None:
        try:
            if self.df is None:
                raise RuntimeError("Nenhum arquivo carregado.")

            payload = json.loads(payload_json) if payload_json else {}
            if not isinstance(payload, dict):
                raise RuntimeError("Payload inválido.")

            cols = payload.get("cols") or []
            weight_var = (payload.get("weightVar") or "").strip()

            missing_vals_global = payload.get("missingVals") or ""
            missing_rules = payload.get("missingRules") or {}

            use_spss_labels = bool(payload.get("useSpssLabels", True))
            method = (payload.get("method") or "pearson").lower().strip()
            decimals = int(payload.get("decimals") or 3)
            convert_text = bool(payload.get("convertText", True))

            if not isinstance(cols, list) or len(cols) < 2:
                raise RuntimeError("Selecione pelo menos 2 colunas.")
            cols = [str(c) for c in cols]

            for c in cols:
                if c not in self.df.columns:
                    raise RuntimeError(f"Coluna inválida: {c}")

            if not weight_var:
                raise RuntimeError("Selecione o peso (obrigatório).")
            if weight_var not in self.df.columns:
                raise RuntimeError("Variável de peso inválida.")

            if method not in ("pearson", "spearman", "kendall"):
                method = "pearson"

            if method in ("pearson", "spearman") and scipy_stats is None:
                raise RuntimeError("Para calcular Sig. (2-tailed), instale: pip install scipy")

            df2 = self.df[cols + [weight_var]].copy()

            # 1) Missing guiado por variável
            if isinstance(missing_rules, dict):
                for c, tokens in missing_rules.items():
                    c = str(c)
                    if c not in df2.columns:
                        continue
                    if not isinstance(tokens, list):
                        continue

                    ser = df2[c]

                    if "__EMPTY__" in tokens:
                        try:
                            if pd.api.types.is_string_dtype(ser) or pd.api.types.is_object_dtype(ser):
                                empty_mask = ser.notna() & ser.astype(str).map(lambda x: x.strip() == "")
                                df2.loc[empty_mask, c] = np.nan
                        except Exception:  # noqa: BLE001
                            pass

                    num_tokens = [t for t in tokens if _looks_numeric_token(t)]
                    str_tokens = [str(t).strip().lower() for t in tokens
                                  if not _looks_numeric_token(t) and t not in ("__EMPTY__", "__NA__")]

                    if num_tokens:
                        tmp_num = pd.to_numeric(
                            ser.astype(str).str.replace(",", ".", regex=False),
                            errors="coerce"
                        )
                        for t in num_tokens:
                            fv = _token_to_float(t)
                            if fv is None:
                                continue
                            df2.loc[tmp_num == fv, c] = np.nan

                    if str_tokens:
                        tmp_txt = ser.astype(str).str.strip().str.lower()
                        df2.loc[tmp_txt.isin(str_tokens), c] = np.nan

            # 2) Missing global manual
            missing_str, missing_num = _parse_missing_list(str(missing_vals_global))
            if missing_str:
                for c in cols:
                    tmp_txt = df2[c].astype(str).str.strip().str.lower()
                    df2.loc[tmp_txt.isin(missing_str), c] = np.nan

            # 3) conversão + missing numérico global
            for c in cols:
                if convert_text:
                    df2[c] = pd.to_numeric(
                        df2[c].astype(str).str.replace(",", ".", regex=False),
                        errors="coerce"
                    )
                else:
                    df2[c] = pd.to_numeric(df2[c], errors="coerce")

                if missing_num:
                    for mv in missing_num:
                        df2.loc[df2[c] == mv, c] = np.nan

            # peso
            w = pd.to_numeric(df2[weight_var], errors="coerce").fillna(0.0).astype(float)

            valid_w = w > 0
            dfv = df2.loc[valid_w, cols].copy()
            wv = w.loc[valid_w].copy()

            # labels iniciais
            cols_display: List[str] = []
            for c in cols:
                if use_spss_labels:
                    lab = (self.var_labels.get(c) or "").strip()
                    cols_display.append(lab if lab else c)
                else:
                    cols_display.append(c)

            # cálculo
            if method == "kendall":
                corr_df = dfv.corr(method="kendall")
                raw_mat = corr_df.to_numpy(dtype=float)
                p_mat = np.full_like(raw_mat, np.nan, dtype=float)
                filtered_mat = raw_mat.copy()
                method_used = "kendall"
                subtitle_extra = f"· peso: {weight_var} (Kendall sem ponderação)"
            else:
                corr_df, sig_df, filtered_df = _weighted_corr_and_sig_matrices(dfv, cols, wv, method=method)
                raw_mat = corr_df.to_numpy(dtype=float)
                p_mat = sig_df.to_numpy(dtype=float)
                filtered_mat = filtered_df.to_numpy(dtype=float)
                method_used = method
                subtitle_extra = f"· peso: {weight_var} · filtro Sig. 0,000 a 0,005"

            # guarda para export
            self.last_cols_vars = cols
            self.last_raw_matrix = raw_mat.copy()
            self.last_p_matrix = p_mat.copy()
            self.last_matrix = filtered_mat.copy()   # matriz exibida/exportada
            self.last_method = method_used
            self.last_decimals = int(decimals)
            self.last_weight_var = weight_var

            payload_out = {
                "ok": True,
                "title": "Matriz de correlação",
                "subtitle": f"{self.data_label} · {len(cols)} variáveis {subtitle_extra}",
                "method": method_used,
                "decimals": int(decimals),
                "vars": cols,
                "cols_display": cols_display,
                "matrix": _jsonable_matrix(filtered_mat),
            }
            self.last_result_payload = payload_out
            self.run_js(_js("window.__setCorrResult", payload_out))

        except Exception as e:  # noqa: BLE001
            payload_out = {"ok": False, "error": str(e)}
            self.run_js(_js("window.__setCorrResult", payload_out))

    # ---------------- exportações
    def _display(self, payload_json: str) -> List[str]:
        payload = json.loads(payload_json) if payload_json else {}
        display = payload.get("display") or []
        if not isinstance(display, list) or len(display) != len(self.last_cols_vars):
            display = self.last_cols_vars[:]  # fallback
        return display

    def csv_bytes(self, display: List[str]) -> bytes:
        df = pd.DataFrame(self.last_matrix, index=display, columns=display)
        for i in range(df.shape[0]):
            for j in range(df.shape[1]):
                if i > j:
                    df.iat[i, j] = np.nan
        df = df.round(int(self.last_decimals))
        buf = io.BytesIO()
        df.to_csv(buf, index=True, sep=";", encoding="utf-8-sig")
        return buf.getvalue()

    def xlsx_bytes(self, display: List[str]) -> bytes:
        cols = display
        mat = self.last_matrix
        pmat = self.last_p_matrix
        dec = int(self.last_decimals)

        wb = Workbook()
        ws = wb.active
        ws.title = "Correlações"

        font_base = Font(name="DIN", size=10, bold=False, color="000000")
        font_head = Font(name="DIN", size=10, bold=True, color="000000")
        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        def write_matrix_sheet(ws_obj, data_matrix: np.ndarray, *, blank_diag: bool, number_format: str, style_func):
            n = len(cols)

            for c in range(1, n + 2):
                ws_obj.column_dimensions[get_column_letter(c)].width = 10
            for r in range(1, n + 2):
                ws_obj.row_dimensions[r].height = 40

            ws_obj.cell(1, 1, "").font = font_head
            ws_obj.cell(1, 1).alignment = align_left
            ws_obj.cell(1, 1).border = border

            for j, name in enumerate(cols, start=2):
                cell = ws_obj.cell(1, j, name)
                cell.font = font_head
                cell.alignment = align_center
                cell.border = border

            for i, row_name in enumerate(cols, start=2):
                c0 = ws_obj.cell(i, 1, row_name)
                c0.font = font_head
                c0.alignment = align_left
                c0.border = border

                for j in range(2, n + 2):
                    ii = i - 2
                    jj = j - 2

                    cell = ws_obj.cell(i, j)
                    cell.border = border
                    cell.alignment = align_center
                    cell.font = font_base

                    if ii > jj:
                        cell.value = None
                        continue

                    if blank_diag and ii == jj:
                        cell.value = None
                        continue

                    v = float(data_matrix[ii, jj]) if np.isfinite(data_matrix[ii, jj]) else np.nan
                    if not np.isfinite(v):
                        cell.value = None
                        continue

                    cell.value = v
                    cell.number_format = number_format

                    stl = style_func(v) if style_func is not None else None
                    if stl is not None:
                        fill_hex, font_hex = stl
                        cell.fill = PatternFill("solid", fgColor=fill_hex)
                        cell.font = Font(name="DIN", size=10, bold=True, color=font_hex)

            ws_obj.freeze_panes = "B2"

        corr_num_fmt = "0" if dec <= 0 else ("0." + ("0" * dec))
        write_matrix_sheet(ws, mat, blank_diag=False, number_format=corr_num_fmt, style_func=_corr_style_rgb)

        ws_sig = wb.create_sheet("Sig. (2-tailed)")
        if pmat is None:
            pmat = np.full((len(cols), len(cols)), np.nan, dtype=float)

        write_matrix_sheet(ws_sig, pmat, blank_diag=True, number_format="0.000", style_func=_sig_style_rgb)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ---------------- @Slot exportCorrCsv
    def exportCorrCsv(self, payload_json: str) -> None:
        try:
            if self.last_matrix is None or not self.last_cols_vars:
                self.aviso("Exportar\n\nCalcule a correlação antes de exportar.", "info")
                return
            self.download("correlacao.csv", self.csv_bytes(self._display(payload_json)), "text/csv")
            self.aviso("Exportar\n\nCSV gerado com sucesso.", "ok")
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao exportar CSV.\n\n{e}", "erro")

    # ---------------- @Slot exportCorrXlsx
    def exportCorrXlsx(self, payload_json: str) -> None:
        try:
            if self.last_matrix is None or not self.last_cols_vars:
                self.aviso("Exportar\n\nCalcule a correlação antes de exportar.", "info")
                return
            self.download("correlacao.xlsx", self.xlsx_bytes(self._display(payload_json)), _MIME_XLSX)
            self.aviso("Exportar\n\nXLSX gerado com sucesso.", "ok")
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao exportar XLSX.\n\n{e}", "erro")


def modulo_correlacao():
    st.header("Matriz de Correlação")
    botao_tutorial("correlacao")

    ponte: PonteCorrelacao = obter_ponte("corr_ponte", PonteCorrelacao)
    processar_pedidos(ponte, "corr_componente")
    conferir_versao_tela(ponte, "relatoria_correlacao", VERSAO_TELA)
    mostrar_notificacoes(ponte)
    # altura "auto": a ferramenta cresce com o conteúdo e a PÁGINA rola (sem tela dentro da tela)
    desenhar_componente(ponte, "relatoria_correlacao", altura="auto", key="corr_componente")
