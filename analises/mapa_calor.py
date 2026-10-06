"""
Mapa de Calor — a tela ORIGINAL do Relatória Analítica dentro do Streamlit.

A interface (HTML/CSS/JS + Leaflet) é a mesma do app desktop:
componentes/relatoria_mapa_calor/index.html (cópia de heatmap_page.py, só
trocando o QWebChannel do Qt pela ponte). Ela conversa com a
`PonteMapaCalor` abaixo, que tem os mesmos métodos (@Slot) da `QtBridge`
original. Leitura do Excel, contas e PPTX: core/mapa_calor_math.py (código
original sem alterações).

O que muda em relação ao desktop:
- a fila do PPTX fica na barra lateral (cards com miniatura, Mín/Média/Máx
  e paleta; Exportar PPTX e Limpar fila); adicionar mostra um pop-up e o
  mesmo caso + segmento pergunta se deve SUBSTITUIR o item da fila;
- os KMLs, a pintura, a numeração e os ajustes visuais ficam guardados no
  Python: trocar de ferramenta no menu e voltar não perde o mapa;
- conferências visíveis: KMLs × regiões do caso, regiões com base < 30 e o
  resultado da pintura (verde = tudo pintado, vermelho = faltou algo);
- "Importar Excel" e "Carregar logo" abrem o seletor de arquivo do navegador;
- "Exportar PNG" e "Exportar PPTX" baixam o arquivo pelo navegador;
- os avisos (QMessageBox) aparecem como faixas no canto da ferramenta;
- o arquivo de debug (.processamentos_relatoria_debug.json) não é gravado
  ao lado do Excel (não há pasta do usuário no servidor).
"""
from __future__ import annotations

import base64
import io
import json
import math
from typing import Any, Dict, List, Optional

import streamlit as st

from analises._pilha_lateral import pilha_lateral
from analises._ponte_relatoria import (
    PonteBase, conferir_versao_tela, desenhar_componente, mostrar_notificacoes, obter_ponte, processar_pedidos,
)
from analises.tutoriais_analitica import botao_tutorial
from core.mapa_calor_math import (
    base_usada, fmt_1_half_up, gerar_pptx, ler_casos_excel, parse_number, round_half_up,
)

#: precisa bater com <meta name="versao-tela"> de componentes/relatoria_mapa_calor/index.html
VERSAO_TELA = "2"

_MIME_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _js(funcao: str, *args) -> str:
    return f"{funcao}({', '.join(json.dumps(a, ensure_ascii=False) for a in args)});"


class PonteMapaCalor(PonteBase):
    """Equivalente à `QtBridge` de heatmap_page.py."""

    ARQUIVOS = {
        "pickExcel": {"accept": ".xlsx"},
        "pickLogo": {"accept": ".png,.jpg,.jpeg"},
    }

    def __init__(self):
        super().__init__()
        self.cases_by_key: Dict[str, Dict[str, Any]] = {}
        self.nome_excel = ""
        self.logo_bytes: Optional[bytes] = None
        self.ppt_queue: List[Dict[str, Any]] = []
        self.bases_regioes: Dict[str, Dict[str, float]] = {}
        self.kmls_json: str = ""        # [{name, text}] dos KMLs importados
        self.estado_mapa_json: str = ""  # pintura, numeração, ajustes visuais...
        self.versao_fila = 0
        self._cache_pptx = None

    # ---------------- estado ao (re)abrir a página
    def estado_inicial(self) -> None:
        if self.cases_by_key:
            self._push_cases()
        self._push_ppt_state()
        if self.kmls_json or self.estado_mapa_json:
            self.run_js(_js("window.__restoreMapState", self.kmls_json or "[]", self.estado_mapa_json or "{}"))

    # ---------------- fila: identidade, miniatura, PPTX
    @staticmethod
    def assinatura(it: Dict[str, Any]) -> str:
        return f"{it.get('case_key', '')}|{it.get('segment', '')}"

    def _mudou_fila(self) -> None:
        self.versao_fila = getattr(self, "versao_fila", 0) + 1
        self._push_ppt_state()

    def pptx_da_fila(self) -> bytes:
        cache = getattr(self, "_cache_pptx", None)
        if not cache or cache[0] != self.versao_fila:
            cache = (self.versao_fila, gerar_pptx(self.ppt_queue, self.logo_bytes))
            self._cache_pptx = cache
        return cache[1]

    def removerDaFila(self, sig: str) -> None:
        for i, it in enumerate(self.ppt_queue):
            if self.assinatura(it) == sig:
                del self.ppt_queue[i]
                self.notificar(f"Removido da fila: {it.get('case_title', '')} · {it.get('segment', '')}", "🗑️")
                self._mudou_fila()
                return

    # ---------------- estado do mapa (guardado para não perder ao trocar de ferramenta)
    def saveKmls(self, kmls_json: str) -> None:
        self.kmls_json = kmls_json or ""

    def saveMapState(self, estado_json: str) -> None:
        self.estado_mapa_json = estado_json or ""

    def _push_ppt_state(self) -> None:
        items = []
        for i, it in enumerate(self.ppt_queue):
            label = f"{it.get('case_title','')} - {it.get('segment','')}".strip(" -")
            items.append({"index": i, "label": label, "sig": self.assinatura(it)})
        payload = {"count": len(self.ppt_queue), "items": items, "logoLoaded": bool(self.logo_bytes)}
        self.run_js(_js("window.__setPptQueueState", json.dumps(payload, ensure_ascii=False)))

    def _push_cases(self) -> None:
        segments_by_case: Dict[str, List[str]] = {}
        base_by_case: Dict[str, str] = {}
        for case_key, case_data in self.cases_by_key.items():
            segments_by_case[case_key] = [k for k in case_data.keys() if not k.startswith("__")]
            meta = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}
            base_total = meta.get("base_total", None)
            base_by_case[case_key] = (f"{float(base_total):.0f}" if isinstance(base_total, (int, float)) else "")
        regions_by_case = {
            k: list(((v.get("__meta") or {}).get("region_names") or [])) for k, v in self.cases_by_key.items()
        }
        payload = {
            "cases": list(self.cases_by_key.keys()),
            "segmentsByCase": segments_by_case,
            "baseByCase": base_by_case,
            "debugPath": "",
            # extras (para as conferências da tela): regiões e base de cada região
            "regionsByCase": regions_by_case,
            "baseRegionsByCase": getattr(self, "bases_regioes", {}) or {},
        }
        self.run_js(_js("window.__setCases", json.dumps(payload, ensure_ascii=False)))

    # ---------------- @Slot removePptItem
    def removePptItem(self, index: int) -> None:
        try:
            index = int(index)
            if 0 <= index < len(self.ppt_queue):
                del self.ppt_queue[index]
            self._mudou_fila()
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao remover item da fila.\n\n{e}", "erro")

    # ---------------- @Slot pickExcel
    def pickExcel(self, arquivo=None) -> None:
        if not arquivo:
            return
        try:
            self.cases_by_key, _debug = ler_casos_excel(arquivo["dados"])
            self.nome_excel = arquivo["name"]
            self.bases_regioes = _bases_por_regiao(arquivo["dados"], self.cases_by_key)
            self._push_cases()
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Erro ao ler Excel\n\nNão consegui analisar o Excel.\n\nDetalhes: {e}", "erro")

    # ---------------- @Slot computeStats
    def computeStats(self, case_key: str, segment_label: str, base_override_text: str) -> None:
        try:
            case_data = self.cases_by_key.get(case_key, {})
            meta = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}
            base_total_used, base_source = base_usada(meta, base_override_text)

            region_map = case_data.get(segment_label, {}) if isinstance(case_data, dict) else {}
            items = [(name, val) for name, val in region_map.items() if isinstance(val, (int, float))]
            if not items:
                self.run_js("window.__setStats('—','—','—','','','',null,null,null);")
                return

            values = [v for (_, v) in items]
            count = len(values)
            total_sum = sum(values)
            mean = total_sum / count

            non_occ = 100.0 - mean
            margin_error: Optional[float] = None
            if base_total_used is not None and base_total_used > 0:
                margin_error = 1.96 * math.sqrt((mean * non_occ) / float(base_total_used))

            min_value = (mean - margin_error) if margin_error is not None else None
            max_value = (mean + margin_error) if margin_error is not None else None

            def fmt_stat(x: Optional[float]) -> str:
                if x is None:
                    return "—"
                return fmt_1_half_up(float(x))

            def fmt_dbg(x: Optional[float]) -> str:
                if x is None:
                    return "—"
                return f"{x:.2f}".replace(".", ",")

            segment_sources = (meta.get("segment_sources", {}) or {}).get(segment_label, {}) or {}
            used_lines: List[str] = []
            for region_name, value in items:
                src = segment_sources.get(region_name, {}) or {}
                parsed = src.get("parsed")
                used_lines.append(
                    f"{region_name}={fmt_dbg(value)} | cell={src.get('addr')} | raw={src.get('raw')} | "
                    f"fmt={src.get('num_format')} | "
                    f"parsed={fmt_dbg(parsed) if isinstance(parsed, (int, float)) else parsed}"
                )
            values_str = "\n".join(used_lines[:300])

            location = (
                f"Aba: {meta.get('sheet')}\n"
                f"Marcador: ({meta.get('marker_row')},{meta.get('marker_col')}) '{meta.get('marker_text')}'\n"
                f"Header: {meta.get('header_row')} | 1º segmento: {meta.get('first_segment_row')} | "
                f"Base: {meta.get('base_row')} | Total col: {meta.get('total_col')}\n"
            )

            tip_mean = (
                f"{location}\n"
                f"Segmento: {segment_label}\n"
                f"Obs.: Total NÃO entra na média.\n\n"
                f"Soma = {fmt_dbg(total_sum)}\n"
                f"n = {count}\n"
                f"Média = Soma/n = {fmt_dbg(total_sum)} / {count} = {fmt_dbg(mean)}\n\n"
                f"Células/valores usados:\n{values_str}"
            )

            if margin_error is None:
                tip_min = tip_mean + "\n\nBase inválida/não definida -> Erro/Mín/Máx não calculados."
                tip_max = tip_min
            else:
                tip_common = (
                    f"{tip_mean}\n\n"
                    f"Base usada ({base_source}) = {fmt_dbg(float(base_total_used))}\n"
                    f"Não ocorrência = 100 - Média = 100 - {fmt_dbg(mean)} = {fmt_dbg(non_occ)}\n"
                    f"Erro = 1,96 * √((Média * Não ocorrência) / Base)\n"
                    f"Erro = 1,96 * √(({fmt_dbg(mean)} * {fmt_dbg(non_occ)}) / {fmt_dbg(float(base_total_used))})"
                    f" = {fmt_dbg(margin_error)}\n"
                )
                tip_min = tip_common + f"\nMínima = Média - Erro = {fmt_dbg(mean)} - {fmt_dbg(margin_error)} = {fmt_dbg(min_value)}"
                tip_max = tip_common + f"\nMáxima = Média + Erro = {fmt_dbg(mean)} + {fmt_dbg(margin_error)} = {fmt_dbg(max_value)}"

            raw_mean = round_half_up(mean, 1)
            raw_min = round_half_up(min_value, 1) if isinstance(min_value, (int, float)) else None
            raw_max = round_half_up(max_value, 1) if isinstance(max_value, (int, float)) else None

            self.run_js(_js("window.__setStats", fmt_stat(min_value), fmt_stat(mean), fmt_stat(max_value),
                            tip_min, tip_mean, tip_max, raw_min, raw_mean, raw_max))
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao calcular estatísticas.\n\n{e}", "aviso")

    # ---------------- @Slot applyAutoPaint
    def applyAutoPaint(self, case_key: str, segment_label: str, base_override_text: str,
                       no_overwrite_manual: bool) -> None:
        try:
            case_data = self.cases_by_key.get(case_key, {})
            meta = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}
            base_total_used, _src = base_usada(meta, base_override_text)
            if base_total_used is None:
                self.aviso("Base inválida\n\nDefina uma Base (Total) válida para aplicar ao mapa.", "aviso")
                return

            region_map = case_data.get(segment_label, {}) if isinstance(case_data, dict) else {}
            items = [(name, val) for name, val in region_map.items() if isinstance(val, (int, float))]
            if not items:
                self.aviso("Sem dados\n\nNão há valores numéricos para este segmento.", "info")
                return

            values = [v for (_, v) in items]
            mean = sum(values) / len(values)
            non_occ = 100.0 - mean
            margin_error = 1.96 * math.sqrt((mean * non_occ) / float(base_total_used))
            min_r = round_half_up(mean - margin_error, 1)
            max_r = round_half_up(mean + margin_error, 1)

            out_items = []
            for name, v in items:
                v_r = round_half_up(float(v), 1)
                if v_r < min_r:
                    cls = "low"
                elif v_r > max_r:
                    cls = "high"
                else:
                    cls = "mid"
                out_items.append({"name": name, "value": v_r, "cls": cls})

            payload = {"segment": segment_label, "min": min_r, "max": max_r,
                       "noOverwriteManual": bool(no_overwrite_manual), "items": out_items}
            self.run_js(_js("window.__applyAutoPaint", json.dumps(payload, ensure_ascii=False)))
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao aplicar pintura automática.\n\n{e}", "erro")

    # ---------------- @Slot pickLogo
    def pickLogo(self, arquivo=None) -> None:
        if not arquivo:
            return
        try:
            self.logo_bytes = arquivo["dados"]
            self._mudou_fila()  # a logo entra em todos os slides
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Não consegui carregar a logo.\n\n{e}", "erro")

    # ---------------- @Slot addToPptQueue
    def addToPptQueue(self, meta_json: str, data_url: str) -> None:
        try:
            meta = json.loads(meta_json)
            if not str(data_url).startswith("data:image/png"):
                self.aviso("PPTX\n\nImagem inválida (esperado PNG).", "aviso")
                return
            _, b64 = data_url.split(",", 1)
            png_bytes = base64.b64decode(b64)

            case_key = meta.get("caseKey", "")
            segment = meta.get("segment", "")
            palette = meta.get("palette", {})
            stats = meta.get("stats", {})

            legend_in = meta.get("legend", [])
            legend_rows: List[Dict[str, Any]] = []
            if isinstance(legend_in, list):
                for row in legend_in:
                    if not isinstance(row, dict):
                        continue
                    name = str(row.get("name", "")).strip()
                    color = str(row.get("color", "")).strip() or "#ffffff"
                    value = row.get("value", None)
                    if name and isinstance(value, (int, float)):
                        legend_rows.append({"name": name, "color": color, "value": float(value)})

            if not case_key or not segment:
                self.aviso("PPTX\n\nMeta incompleta (case/segment).", "aviso")
                return
            if not all(k in palette for k in ("low", "mid", "high")):
                self.aviso("PPTX\n\nPaleta incompleta.", "aviso")
                return
            if not all(k in stats for k in ("min", "mean", "max")):
                self.aviso("PPTX\n\nEstatísticas incompletas (min/mean/max).", "aviso")
                return

            case_data = self.cases_by_key.get(case_key, {})
            meta_case = case_data.get("__meta", {}) if isinstance(case_data, dict) else {}
            case_title = meta_case.get("title", "") or case_key.split(" — ")[0]

            novo = {
                "case_key": case_key,
                "case_title": case_title,
                "segment": segment,
                "palette": palette,
                "stats": {"min": float(stats["min"]), "mean": float(stats["mean"]), "max": float(stats["max"])},
                "png_bytes": png_bytes,
                "legend_rows": legend_rows,
                "miniatura": _miniatura(png_bytes),
            }
            sig = self.assinatura(novo)
            pos = next((i for i, it in enumerate(self.ppt_queue) if self.assinatura(it) == sig), None)
            nome = f"{case_title} · {segment}"
            if pos is None:
                self.ppt_queue.append(novo)
                self.notificar(f"Adicionado à fila: {nome} — {len(self.ppt_queue)} mapa(s) na fila.", "✅")
            elif meta.get("substituir"):
                self.ppt_queue[pos] = novo
                self.notificar(f"Mapa substituído na fila (item {pos + 1}): {nome}.", "🔁")
            else:
                self.notificar(f"Esse mapa já está na fila (item {pos + 1}): {nome}.", "⚠️")
                return
            self._mudou_fila()
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao adicionar à fila.\n\n{e}", "erro")

    # ---------------- @Slot clearPptQueue
    def clearPptQueue(self) -> None:
        if self.ppt_queue:
            self.notificar("Fila esvaziada.", "🗑️")
        self.ppt_queue.clear()
        self._mudou_fila()

    # ---------------- @Slot exportPptx
    def exportPptx(self) -> None:
        if not self.ppt_queue:
            self.aviso("PPTX\n\nA fila está vazia. Adicione mapas antes de exportar.", "info")
            return
        try:
            dados = gerar_pptx(self.ppt_queue, self.logo_bytes)
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Falha ao gerar PPTX.\n\n{e}", "erro")
            return
        self.download("mapas_calor.pptx", dados, _MIME_PPTX)
        self.aviso(f"PPTX\n\nApresentação gerada com {len(self.ppt_queue)} slide(s).", "ok")

    # ---------------- @Slot savePng
    def savePng(self, data_url: str) -> None:
        try:
            if not str(data_url).startswith("data:image/png"):
                self.aviso("Exportação\n\nFormato inesperado ao exportar PNG.", "aviso")
                return
            _, b64 = data_url.split(",", 1)
            self.download("mapa_calor.png", base64.b64decode(b64), "image/png")
        except Exception as e:  # noqa: BLE001
            self.aviso(f"Erro ao salvar PNG\n\n{e}", "erro")


def _miniatura(png_bytes: bytes, largura: int = 420) -> str:
    """JPEG pequeno do mapa capturado (para o card da fila)."""
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(png_bytes)).convert("RGB")
        if im.width > largura:
            im = im.resize((largura, max(1, round(im.height * largura / im.width))))
        out = io.BytesIO()
        im.save(out, "JPEG", quality=80)
        return "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode("ascii")
    except Exception:  # noqa: BLE001
        return ""


def _bases_por_regiao(conteudo: bytes, casos: Dict[str, Dict[str, Any]]) -> Dict[str, Dict[str, float]]:
    """Base (N) de cada região de cada caso, lida na linha "Base" do relatório."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=False)
    except Exception:  # noqa: BLE001
        return {}
    out: Dict[str, Dict[str, float]] = {}
    for chave, caso in casos.items():
        meta = caso.get("__meta") or {}
        linha, aba = meta.get("base_row"), meta.get("sheet")
        if not linha or aba not in wb.sheetnames:
            continue
        ws = wb[aba]
        bases = {}
        for nome, col in zip(meta.get("region_names") or [], meta.get("region_cols") or []):
            v = parse_number(ws.cell(linha, col))
            if isinstance(v, (int, float)):
                bases[nome] = float(v)
        out[chave] = bases
    return out


def _fmt(x: Any) -> str:
    try:
        return fmt_1_half_up(float(x))
    except (TypeError, ValueError):
        return "—"


def _itens_fila(ponte: PonteMapaCalor) -> List[Dict[str, Any]]:
    itens = []
    for it in ponte.ppt_queue:
        pal, stt = it.get("palette") or {}, it.get("stats") or {}
        itens.append({
            "sig": ponte.assinatura(it),
            "title": f"{it.get('case_title', '')} · {it.get('segment', '')}",
            "imagem": it.get("miniatura") or "",
            "chips": [
                {"rotulo": "Mín", "cor": pal.get("low"), "valor": _fmt(stt.get("min"))},
                {"rotulo": "Média", "cor": pal.get("mid"), "valor": _fmt(stt.get("mean"))},
                {"rotulo": "Máx", "cor": pal.get("high"), "valor": _fmt(stt.get("max"))},
            ],
            "secoes": [
                {"rotulo": "Caso", "texto": it.get("case_key", ""), "cor": "#60a5fa"},
                {"rotulo": "Segmento", "texto": it.get("segment", ""), "cor": "#a78bfa"},
            ],
        })
    return itens


def _fila_na_barra_lateral(ponte: PonteMapaCalor) -> None:
    pilha_lateral(
        ponte, prefixo="mc", itens=lambda: _itens_fila(ponte), remover=ponte.removerDaFila,
        titulo="Fila do PPTX",
        subtitulo=lambda: ("✓ Logo carregada" if ponte.logo_bytes
                           else "Sem logo — use “Carregar logo” na ferramenta"),
        subtitulo_tipo="ok" if ponte.logo_bytes else "aviso",
        vazio_msg="Vazia. Monte o mapa e clique em “Adicionar à fila”.",
        exportar=lambda: {"rotulo": f"Exportar PPTX ({len(ponte.ppt_queue)})", "dados": ponte.pptx_da_fila,
                          "arquivo": "mapas_calor.pptx", "mime": _MIME_PPTX},
        limpar=ponte.clearPptQueue, rotulo_limpar="Limpar fila",
    )


def modulo_mapa_calor():
    st.header("Mapa de Calor")
    botao_tutorial("mapa_calor")

    ponte: PonteMapaCalor = obter_ponte("mc_ponte", PonteMapaCalor)
    processar_pedidos(ponte, "mc_componente")
    conferir_versao_tela(ponte, "relatoria_mapa_calor", VERSAO_TELA)
    _fila_na_barra_lateral(ponte)
    mostrar_notificacoes(ponte)
    # altura "janela": a ferramenta ocupa a altura do navegador (o mapa precisa de altura fixa)
    desenhar_componente(ponte, "relatoria_mapa_calor", altura="janela", key="mc_componente")
