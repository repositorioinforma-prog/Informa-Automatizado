"""
Codificação automática de pergunta aberta sobre candidato (Presidente,
Governador, Senador, Deputado Estadual/Federal, Prefeito, Vereador),
usando uma base de candidatos (ex.: TSE) como gabarito de nomes.

Casa cada resposta de texto livre contra os candidatos do cargo (nome
de urna + nome completo + nome social, quando disponíveis), por essa
ordem de confiança:
    1. match exato contra qualquer forma de nome do candidato;
    2. toda palavra da resposta aparece dentro de alguma forma de nome
       do candidato (cobre nome de batismo/nome completo parcial, tipo
       "Luiz Inácio" -> "Luiz Inácio Lula da Silva", sem precisar de
       tabela de apelido manual);
    3. token único (sobrenome, apelido de uma palavra) que só bate num
       candidato — vira ambiguidade de "Família X" se bater em mais de
       um;
    4. fuzzy matching genérico (fallback pra variação leve de escrita).

Categorias especiais (por regra, não por candidato): Não sabe (99),
Branco/Nulo (97), Não irá votar (98), Outros (96), Partidos citados,
Candidatos de esquerda/direita.

Sem dependência do Streamlit — pode ser testado isoladamente (exceto a
leitura de `.sav`, que precisa do pyreadstat).
"""
from __future__ import annotations

import re
import unicodedata
from collections import Counter

try:
    from rapidfuzz import fuzz

    def _similaridade(a, b):
        return fuzz.token_sort_ratio(a, b) / 100.0
except ImportError:
    import difflib

    def _similaridade(a, b):
        # token_sort: ordena as palavras antes de comparar, pra "SILVA
        # JOAO" bater com "JOAO SILVA" igual o rapidfuzz faz — sem isso,
        # ordem diferente de nome/sobrenome derrubava a similaridade à
        # toa. Recomendado adicionar `rapidfuzz` no requirements.txt pra
        # produção — é mais robusto com erro de digitação pesado e mais
        # rápido em listas de candidatos grandes; o difflib aqui é só um
        # fallback que sempre funciona sem depender de instalar nada.
        ta = " ".join(sorted(a.split()))
        tb = " ".join(sorted(b.split()))
        return difflib.SequenceMatcher(None, ta, tb).ratio()


def normalizar(texto):
    if texto is None:
        return ""
    s = str(texto)
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.upper()
    s = re.sub(r"[^A-Z0-9\s]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


# ---------------------------------------------------------------------
# Regras de categorias especiais (rodam ANTES do casamento com candidato)
# ---------------------------------------------------------------------
# ---------------------------------------------------------------------
# Padrões de "não sabe"/"branco-nulo"/"não vai votar" — ampliados a
# partir de uma lista de ~150 frases levantada em cima de respostas
# reais de pesquisa (curadoria de outro projeto do próprio usuário),
# roteadas pra cada categoria de acordo com o esquema deste módulo
# (algumas frases da lista original, tipo "não votar", pertencem
# semanticamente à categoria separada "Não irá votar" aqui, e
# "nenhum"/"ninguém" entram em Branco/Nulo, não em Não sabe). Continuam
# usadas com MATCH EXATO (`re.fullmatch` via `_bate_algum_padrao`), não
# substring — isso é proposital: comparar substring (tipo checar se
# "ns" aparece em QUALQUER LUGAR do texto) classificaria errado
# respostas legítimas que só por acaso contêm essas letras juntas (ex.:
# "ns" aparece dentro de "transporte" e "construção").
PADROES_NAO_SABE = [
    r"^NAO SEI$", r"^NAO SEI AINDA$", r"^N SEI$", r"^NAO SABE$", r"^NS$",
    r"^NAO TENHO CERTEZA$", r"^AINDA NAO SEI$", r"^NAO DECIDI$", r"^INDECISO A?$",
    r"^NAO PENSEI$", r"^NAO LEMBRO$", r"^AINDA EM DUVIDA$", r"^AINDA INDECISO$",
    r"^AINDA N SEI$", r"^AINDA NAO DECIDI$", r"^AINDA NAO DEFINI CANDIDATO$",
    r"^AINDA NAO DEFINI MEU VOTO$", r"^AINDA NAO SEI AINDA$",
    r"^AINDA NAO SEI OS CANDIDATOS$", r"^AINDA NAO SEI QUEM SAO$",
    r"^AINDA NAO SEI QUEM SAO OS CANDIDATOS$", r"^AINDA NAO TEM CONHECIMENTO DE NENHUM NOME$",
    r"^AINDA NAO TENHO CANDIDATO$", r"^AINDA NAO TENHO CANDIDATO ESPECIFICO$",
    r"^AINDA NAO TENHO CONHECIMENTO$", r"^ATE O MOMENTO NAO SAIU A LISTA DOS CANDIDATOS$",
    r"^DESCONHECO$", r"^DESCONHECO OS CANDIDATOS$", r"^EM DUVIDA$", r"^ESTOU INDECISA$",
    r"^ESTOU INDECISA SOBRE$", r"^ESTOU INDECISO$", r"^EU NAO SEI$", r"^INDECISA$",
    r"^INDECISAO DUVIDA$", r"^INDECISO$", r"^INDEFINIDO$",
    r"^INDEFINIDO NAO CONHECO OS CANDIDATOS AINDA$", r"^N S$", r"^N SABE$", r"^N TEM$",
    r"^NA DUVIDA$", r"^NADA$", r"^NAO CONHECO$", r"^NAO CONHECO ISSO$",
    r"^NAO CONHECO NENHUM$", r"^NAO CONHECO NENHUM CANDIDATO$",
    r"^NAO CONHECO NENHUM CANDIDATO A GOVERNO$", r"^NAO CONHECO NENHUM CANDIDATO A SENADO$",
    r"^NAO CONHECO NINGUEM$", r"^NAO DECIDI AINDA$", r"^NAO DECIDIDO$", r"^NAO DEFINIDO$",
    r"^NAO ESTOU A PAR DOS PRE CANDIDATOS$", r"^NAO LEMBRA$", r"^NAO ME LEMBRO$",
    r"^NAO ME RECORDO$", r"^NAO OUVI$", r"^NAO POSSO DIZER$", r"^NAO RECORDO$",
    r"^NAO RESOLVI$", r"^NAO SEI AINDA NAO SEI AINDA$", r"^NAO SEI DEFINIR AINDA$",
    r"^NAO SEI DIZER$", r"^NAO SEI INFORMAR$", r"^NAO SEI OPINAR$",
    r"^NAO SEI OS CANDIDATOS$", r"^NAO SEI QUAIS SAO OS CANDIDATOS$",
    r"^NAO SEI QUAL E OS CANDIDATOS$", r"^NAO SEI QUE SAO OS CANDIDATOS$",
    r"^NAO SEI QUEM E CANDIDATO$", r"^NAO SEI QUEM E OS CANDIDATOS$", r"^NAO SEI QUEM SAO$",
    r"^NAO SEI QUEM SAO AS OPCOES AINDA$", r"^NAO SEI QUEM SAO OS CANDIDATOS$", r"^NAO TEM$",
    r"^NAO TENHO$", r"^NAO TENHO AINDA$", r"^NAO TENHO CADIDATO$", r"^NAO TENHO CANDIDATO$",
    r"^NAO TENHO CANDIDATO AINDA$", r"^NAO TENHO CANDIDATO ESPECIFICO$",
    r"^NAO TENHO CANDIDATOS$", r"^NAO TENHO CANDITADO$", r"^NAO TENHO CANDITATO$",
    r"^NAO TENHO CONHECIMENTO DE NENHUM$", r"^NAO TENHO IDEIA$", r"^NAO TENHO NEM UMA OPCAO$",
    r"^NAO TENHO NENHUM CANDIDATO EM MENTE$", r"^NAO TENHO NENHUM NOME$",
    r"^NAO TENHO OPCAO$", r"^NAO TENHO OPCAO NO MOMENTO$", r"^NAO TENHO UM CANDIDATO$",
    r"^NAO TENHO UM CANDIDATO AINDA$", r"^NAO TENHO UMA RESPOSTA$", r"^NAO TENHO VOTO$",
    r"^NSEI$", r"^NTEM$", r"^NUNCA OUVI$", r"^NUNCA OUVI FALAR$", r"^POR EQUANTO NAO SEI$",
    r"^PREFIRO NAO DIZER$",
    r"^QUAIS CANDIDATOS A SENADOR NAO SEI E NEM TENHO NOMES DE CANDIDATOS PRA ELEICAO$",
    r"^SEI$", r"^SEM$", r"^SEM CANDIDATO$", r"^SEM CANDIDATO AINDA$",
    r"^SEM CANDIDATO ATE O MOMENTO$", r"^SEM CANDITATO$", r"^SEM OPCAO$", r"^SEM OPINIAO$",
    r"^SEM RESPOSTA$", r"^SEM VOTO$", r"^TAMBEM AINDA NAO SEI$", r"^TAMBEM NAO SEI$",
    r"^TAMBEM NAO SEI AINDA$", r"^TAMBEM NAO SEI NAO$", r"^TB NAO SEI$", r"^TMBN EM DUVIDA$",
    r"^TO EM DUVIDA$", r"^N SE86$", r"^NT$",
    r"^NAO TENHO CANDIDATO AINDA ESTOU PESQUISANDO AINDA$",
    r"^NAO LEMBRA O NOME EX PREFEITO$", r"^MS$", r"^TBM NAO SEI$", r"^ESQUECEU O NOME$",
    r"^NAO TENHO NENHUM AINDA$", r"^ESQUECI$", r"^NAO VOU FALAR$", r"^AINDA ESTOU DEFININDO$",
    r"^NO SEU$", r"^NAO ESTOU ATUALIZADA$", r"^NAO LEMBRA O NOME$", r"^NAO SE$",
    r"^NAO SEIB$", r"^NAO RESPONDO$", r"^NAO QUERO FALAR$", r"^TO ANALISANDO AINDA$",
    r"^NAO DECIDIU$", r"^NAO SRI$", r"^NA SEI$", r"^NAO SUS$", r"^NAO ZEIB$", r"^AVALIANDO$",
    r"^NAO PROCUROU$", r"^TO POR FORA$",
]
PADROES_BRANCO_NULO = [
    r"^BRANCO$", r"^NULO$", r"^VOU ANULAR$", r"^VOTO NULO$", r"^VOTO BRANCO$", r"^NENHUM$",
    r"^NENHUMA$", r"^NAO VOU ESCOLHER NINGUEM$", r"^ANULAR$", r"^EM NINGUEM$", r"^NEM UM$",
    r"^NENHUM CANDIDATO BOM AINDA$", r"^NENHUM DESSES TRAIRAS$", r"^NENHUM DOS CANDIDATOS$",
    r"^NENHUM DOS CANDIDATOS A CANDIDATO$", r"^NENHUM NENHUM$", r"^NENHUM TAMBEM$",
    r"^NINGUEM$", r"^NINGUEM NO MOMENTO$", r"^PARA NINGUEM$", r"^PRA NENHUM$", r"^ANULADO$",
    r"^ANULOU$", r"^BRANCO NULO$", r"^BRANCO OU NULO$", r"^NWNHUM$", r"^NGM$", r"^NG$",
    r"^NIN$", r"^NE$", r"^NL$",
]
PADROES_NAO_VAI_VOTAR = [
    r"^NAO VOU VOTAR$", r"^NAO IREI VOTAR$", r"^NAO VOTO$", r"^NAO PRETENDO VOTAR$",
    r"^VOTO OBRIGATORIO$", r"^NAO QUERO VOTAR$", r"^NAO VOTAR$", r"^NAO VAI VOTAR$",
    r"^NAO IRA VOTAR$",
]
PADROES_ESQUERDA = [r"^ESQUERDA$", r"^ALGUEM DE ESQUERDA$", r"^CANDIDATO DE ESQUERDA$"]
PADROES_DIREITA = [r"^DIREITA$", r"^ALGUEM DE DIREITA$", r"^CANDIDATO DE DIREITA$"]

PARTIDOS_CONHECIDOS = {
    "PT", "PL", "PSD", "PSDB", "MDB", "UNIAO", "UNIAO BRASIL", "PP", "REPUBLICANOS",
    "PDT", "PSB", "PODEMOS", "NOVO", "REDE", "PSOL", "AVANTE", "SOLIDARIEDADE",
    "PRTB", "DEMOCRATA", "PCO", "PCB", "PSTU", "UP", "DC", "MISSAO", "AGIR",
}

CATEGORIAS_FIXAS = {"Não sabe", "Branco/Nulo", "Não irá votar", "Outros"}
CODIGOS_RESERVADOS = {"Outros": 96, "Branco/Nulo": 97, "Não irá votar": 98, "Não sabe": 99}


def _bate_algum_padrao(texto_norm, padroes):
    return any(re.fullmatch(p, texto_norm) for p in padroes)


def _tokens(texto_norm):
    return set(texto_norm.split())


# ---------------------------------------------------------------------
# Candidatos
# ---------------------------------------------------------------------
def montar_candidato(nome_urna, nome_completo=None, nome_social=None, partido=None, uf=None):
    """Monta o dict de candidato usado pelo classificador — reúne todas
    as formas de nome conhecidas (urna, completo, social) num único
    lugar, já que o casamento tenta contra QUALQUER uma delas."""
    variantes = []
    for v in (nome_urna, nome_completo, nome_social):
        if v and str(v).strip() and str(v).strip() not in variantes:
            variantes.append(str(v).strip())
    return {"nome": nome_urna, "variantes": variantes, "partido": partido, "uf": uf}


def classificar_resposta(resposta, candidatos, limiar_confianca=0.72):
    """
    candidatos: lista de dicts no formato de `montar_candidato()`.

    Devolve dict: {"categoria": str, "candidato": str|None,
                    "confianca": float, "motivo": str}
    (o código numérico é atribuído depois, em `classificar_banco`, que
    conhece a distribuição inteira pra decidir a ordem/base mínima).
    """
    texto_norm = normalizar(resposta)

    if not texto_norm:
        return {"categoria": "Branco/Nulo", "candidato": None,
                "confianca": 1.0, "motivo": "resposta vazia"}

    if _bate_algum_padrao(texto_norm, PADROES_NAO_SABE):
        return {"categoria": "Não sabe", "candidato": None,
                "confianca": 1.0, "motivo": "padrão não-sabe"}

    if _bate_algum_padrao(texto_norm, PADROES_BRANCO_NULO):
        return {"categoria": "Branco/Nulo", "candidato": None,
                "confianca": 1.0, "motivo": "padrão branco/nulo"}

    if _bate_algum_padrao(texto_norm, PADROES_NAO_VAI_VOTAR):
        return {"categoria": "Não irá votar", "candidato": None,
                "confianca": 1.0, "motivo": "padrão não-vai-votar"}

    if _bate_algum_padrao(texto_norm, PADROES_ESQUERDA):
        return {"categoria": "Candidatos de esquerda", "candidato": None,
                "confianca": 1.0, "motivo": "menção de espectro (esquerda) sem candidato"}

    if _bate_algum_padrao(texto_norm, PADROES_DIREITA):
        return {"categoria": "Candidatos de direita", "candidato": None,
                "confianca": 1.0, "motivo": "menção de espectro (direita) sem candidato"}

    if texto_norm in PARTIDOS_CONHECIDOS:
        return {"categoria": "Partidos citados", "candidato": None,
                "confianca": 1.0, "motivo": f"partido citado sozinho ({texto_norm})"}

    texto_tokens = _tokens(texto_norm)

    # 1) Match exato contra QUALQUER variante de nome do candidato.
    for c in candidatos:
        if any(texto_norm == normalizar(v) for v in c["variantes"]):
            return {"categoria": c["nome"], "candidato": c["nome"],
                    "confianca": 1.0, "motivo": "match exato"}

    # 2) Subconjunto de tokens: toda palavra da resposta aparece dentro de
    # alguma variante de nome do candidato (ex.: "Luiz Inácio" ⊂ "Luiz
    # Inácio Lula da Silva"). Exige 2+ tokens pra não disparar à toa.
    if len(texto_tokens) >= 2:
        candidatos_subset = []
        for c in candidatos:
            for v in c["variantes"]:
                if texto_tokens <= _tokens(normalizar(v)):
                    candidatos_subset.append(c)
                    break
        candidatos_subset_unicos = list({c["nome"]: c for c in candidatos_subset}.values())
        if len(candidatos_subset_unicos) == 1:
            return {"categoria": candidatos_subset_unicos[0]["nome"],
                    "candidato": candidatos_subset_unicos[0]["nome"], "confianca": 0.98,
                    "motivo": "todas as palavras da resposta aparecem no nome do candidato"}
        if len(candidatos_subset_unicos) >= 2:
            return {"categoria": f"Família {resposta.strip().title()}", "candidato": None,
                    "confianca": 0.9,
                    "motivo": f"tokens batem em {len(candidatos_subset_unicos)} candidatos diferentes"}

    # 3) Token único (sobrenome, apelido de 1 palavra) — precisa ser
    # exclusivo de UM candidato entre TODAS as variantes de nome de
    # TODOS os candidatos, senão vira ambiguidade de família.
    if len(texto_tokens) == 1:
        token = next(iter(texto_tokens))
        candidatos_com_token = []
        for c in candidatos:
            if any(token in _tokens(normalizar(v)) for v in c["variantes"]):
                candidatos_com_token.append(c)
        candidatos_com_token_unicos = list({c["nome"]: c for c in candidatos_com_token}.values())
        if len(candidatos_com_token_unicos) == 1:
            return {"categoria": candidatos_com_token_unicos[0]["nome"],
                    "candidato": candidatos_com_token_unicos[0]["nome"], "confianca": 0.95,
                    "motivo": "sobrenome/token único, sem ambiguidade"}
        if len(candidatos_com_token_unicos) >= 2:
            return {"categoria": f"Família {resposta.strip().title()}", "candidato": None,
                    "confianca": 1.0,
                    "motivo": f"token ambíguo entre {len(candidatos_com_token_unicos)} candidatos"}

    # 4) Fuzzy genérico: compara contra CADA variante de nome de cada
    # candidato, fica com o melhor score entre todas.
    melhor_candidato, melhor_score = None, 0.0
    for c in candidatos:
        score_candidato = max((_similaridade(texto_norm, normalizar(v)) for v in c["variantes"]), default=0.0)
        if score_candidato > melhor_score:
            melhor_score, melhor_candidato = score_candidato, c

    if melhor_candidato and melhor_score >= limiar_confianca:
        return {"categoria": melhor_candidato["nome"], "candidato": melhor_candidato["nome"],
                "confianca": round(melhor_score, 3), "motivo": "fuzzy match"}

    return {"categoria": "Outros", "candidato": None,
            "confianca": round(melhor_score, 3) if melhor_candidato else 0.0,
            "motivo": f"abaixo do limiar (melhor: {melhor_candidato['nome'] if melhor_candidato else '-'} @ {melhor_score:.2f})"}


# ---------------------------------------------------------------------
# Leitura da base de candidatos (TSE ou similar)
# ---------------------------------------------------------------------
_COLUNAS_RECONHECIDAS = {
    "uf": ("UF",),
    "nome_completo": ("NOME COMPLETO", "NM CANDIDATO"),
    "nome_urna": ("NOME NA URNA", "NOME DO CANDIDATO", "NM URNA CANDIDATO"),
    "nome_social": ("NOME SOCIAL APELIDO", "NOME SOCIAL", "APELIDO"),
    "cargo": ("CARGO",),
    "partido": ("PARTIDO", "SIGLA PARTIDO"),
}


def _mapear_colunas(cabecalho):
    idx = {}
    for i, nome_col in enumerate(cabecalho, start=1):
        n = normalizar(nome_col or "")
        for chave, variantes in _COLUNAS_RECONHECIDAS.items():
            if chave in idx:
                continue
            if n in variantes:
                idx[chave] = i
    return idx


def ler_base_candidatos(caminho_ou_arquivo, aba=None):
    """
    Lê um arquivo de candidatos (ex.: exportação do TSE) e devolve
    (candidatos_por_cargo, cargos_disponiveis, ufs_disponiveis) —
    `candidatos_por_cargo` é um dict CARGO -> lista de dicts brutos
    {uf, nome_completo, nome_urna, nome_social, partido}, sem filtrar
    ainda (o filtro por cargo/UF fica em `filtrar_candidatos`).

    Reconhece as colunas por nome (tolerante a variação de escrita:
    "Nome na urna"/"Nome do candidato", "Nome completo"/"NM_CANDIDATO"
    etc.) — exige pelo menos uma coluna de nome e uma de cargo.
    """
    import openpyxl

    wb = openpyxl.load_workbook(caminho_ou_arquivo, rich_text=True)
    ws = wb[aba] if aba else wb[wb.sheetnames[0]]
    cabecalho = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    idx = _mapear_colunas(cabecalho)

    if "cargo" not in idx or ("nome_urna" not in idx and "nome_completo" not in idx):
        raise ValueError(
            "Não encontrei as colunas esperadas (preciso pelo menos de "
            "'Cargo' e 'Nome na urna' ou 'Nome completo'). Colunas "
            f"encontradas: {', '.join(str(c) for c in cabecalho if c)}."
        )

    candidatos_por_cargo = {}
    ufs = set()
    for r in range(2, ws.max_row + 1):
        cargo = ws.cell(row=r, column=idx["cargo"]).value
        if not cargo or not str(cargo).strip():
            continue
        cargo = str(cargo).strip()

        registro = {
            "uf": ws.cell(row=r, column=idx["uf"]).value if "uf" in idx else None,
            "nome_completo": ws.cell(row=r, column=idx["nome_completo"]).value if "nome_completo" in idx else None,
            "nome_urna": ws.cell(row=r, column=idx["nome_urna"]).value if "nome_urna" in idx else None,
            "nome_social": ws.cell(row=r, column=idx["nome_social"]).value if "nome_social" in idx else None,
            "partido": ws.cell(row=r, column=idx["partido"]).value if "partido" in idx else None,
        }
        if not registro["nome_urna"] and not registro["nome_completo"]:
            continue

        candidatos_por_cargo.setdefault(cargo, []).append(registro)
        if registro["uf"]:
            ufs.add(str(registro["uf"]).strip())

    cargos_disponiveis = sorted(candidatos_por_cargo.keys())
    return candidatos_por_cargo, cargos_disponiveis, sorted(ufs)


def filtrar_candidatos(candidatos_por_cargo, cargo, uf=None):
    """Monta a lista final de candidatos (via `montar_candidato`) pro
    cargo escolhido, opcionalmente filtrada por UF (obrigatório na
    prática pra cargos estaduais/municipais — Senador, Governador,
    Deputado Estadual/Federal, Prefeito, Vereador — já que o mesmo
    sobrenome/nome pode aparecer em estados diferentes sem relação
    nenhuma entre si; Presidente é nacional, não precisa de UF)."""
    registros = candidatos_por_cargo.get(cargo, [])
    if uf:
        uf_norm = normalizar(uf)
        registros = [r for r in registros if normalizar(r.get("uf") or "") == uf_norm]

    candidatos = [
        montar_candidato(
            nome_urna=r["nome_urna"] or r["nome_completo"],
            nome_completo=r["nome_completo"],
            nome_social=r["nome_social"],
            partido=r["partido"],
            uf=r["uf"],
        )
        for r in registros
    ]
    # remove duplicatas exatas de nome de urna (candidato repetido na
    # base, ex.: mesmo candidato listado 2x por algum motivo de origem)
    vistos = set()
    candidatos_unicos = []
    for c in candidatos:
        if c["nome"] not in vistos:
            vistos.add(c["nome"])
            candidatos_unicos.append(c)
    return candidatos_unicos


# ---------------------------------------------------------------------
# Classificação do banco inteiro + frequência + código numérico
# ---------------------------------------------------------------------
def classificar_banco(valores, candidatos, limiar_confianca=0.72, base_minima=3):
    """
    valores: lista de respostas textuais brutas (uma por caso — vem
    direto da coluna do banco, pode ter repetição literal).

    Classifica só os valores DISTINTOS (mais rápido quando há muita
    repetição) e expande o resultado de volta pra cada caso. Aplica a
    base mínima: categorias de candidato/família/etc. com menos de
    `base_minima` casos no total viram "Outros" — mesma convenção do
    resto do projeto pras demais abertas.

    Devolve dict:
        "resultado_por_caso": lista (mesma ordem/tamanho de `valores`)
            de {"valor_original", "categoria", "codigo", "candidato",
                "confianca", "motivo"}
        "frequencia": lista de {"categoria", "codigo", "n", "percentual"},
            candidatos por N decrescente primeiro, reservados
            (Outros/Branco-Nulo/Não irá votar/Não sabe) por último
        "mapa_valor_categoria": dict valor_original -> resultado
            completo (chave "" representa vazio/ausente)
        "codigos": dict categoria -> código numérico
    """
    def _chave(v):
        return "" if v is None or str(v).strip() == "" else v

    valores_distintos = sorted({_chave(v) for v in valores}, key=str)
    mapa = {v: classificar_resposta(v, candidatos, limiar_confianca) for v in valores_distintos}

    contagem_valores = Counter(mapa[_chave(v)]["categoria"] for v in valores)

    for r in mapa.values():
        if r["categoria"] not in CATEGORIAS_FIXAS and contagem_valores[r["categoria"]] < base_minima:
            r["categoria_antes_base_minima"] = r["categoria"]
            r["categoria"] = "Outros"
            r["motivo"] = r["motivo"] + f" (menos de {base_minima} casos, agrupado em Outros)"

    contagem_final = Counter(mapa[_chave(v)]["categoria"] for v in valores)
    total = len(valores)

    nao_reservadas = sorted(
        (c for c in contagem_final if c not in CODIGOS_RESERVADOS),
        key=lambda c: -contagem_final[c],
    )
    codigos = {c: i + 1 for i, c in enumerate(nao_reservadas)}
    codigos.update(CODIGOS_RESERVADOS)

    for r in mapa.values():
        r["codigo"] = codigos[r["categoria"]]

    resultado_por_caso = [
        {"valor_original": v, **mapa[_chave(v)]} for v in valores
    ]

    frequencia = [
        {
            "categoria": c, "codigo": codigos[c], "n": n,
            "percentual": round(100 * n / total, 1) if total else 0.0,
        }
        for c, n in contagem_final.items()
    ]
    frequencia.sort(key=lambda x: (
        x["categoria"] in CODIGOS_RESERVADOS,
        x["codigo"] if x["categoria"] in CODIGOS_RESERVADOS else 0,
        -x["n"],
    ))

    return {
        "resultado_por_caso": resultado_por_caso,
        "frequencia": frequencia,
        "mapa_valor_categoria": mapa,
        "codigos": codigos,
    }


# ---------------------------------------------------------------------
# Geração da Syntax SPSS
# ---------------------------------------------------------------------
def gerar_syntax(nome_variavel_origem, nome_variavel_destino, mapa_valor_categoria, codigos, rotulo_cargo=""):
    """
    Monta o texto pronto de RECODE + VARIABLE LABELS + VALUE LABELS pra
    colar no fluxo de programação SPSS do projeto (Syntax.sps).

    Todo comando termina com ponto, como o padrão do projeto exige.
    """
    linhas = []
    if rotulo_cargo:
        linhas.append(f"* Codificação automática de {nome_variavel_origem} — {rotulo_cargo}.")
    linhas.append(f"RECODE {nome_variavel_origem}")
    for valor, r in sorted(mapa_valor_categoria.items(), key=lambda kv: (kv[1]["codigo"], str(kv[0]))):
        if valor == "":
            continue  # vazio/ausente não é um valor literal comparável no RECODE
        valor_escapado = str(valor).replace("'", "''")
        linhas.append(f"    ('{valor_escapado}'={r['codigo']})")
    linhas.append(f"    (ELSE={CODIGOS_RESERVADOS['Outros']})")
    linhas.append(f"    INTO {nome_variavel_destino}.")
    linhas.append("EXECUTE.")
    linhas.append("")
    rotulo_var = f"{nome_variavel_origem} codificada" + (f" — {rotulo_cargo}" if rotulo_cargo else "")
    linhas.append(f"VARIABLE LABELS {nome_variavel_destino} '{rotulo_var}'.")
    linhas.append(f"VALUE LABELS {nome_variavel_destino}")
    itens_ordenados = sorted(codigos.items(), key=lambda kv: kv[1])
    for i, (categoria, codigo) in enumerate(itens_ordenados):
        categoria_escapada = categoria.replace("'", "''")
        fim = "." if i == len(itens_ordenados) - 1 else ""
        linhas.append(f"    {codigo} '{categoria_escapada}'{fim}")
    linhas.append("EXECUTE.")
    return "\n".join(linhas)
