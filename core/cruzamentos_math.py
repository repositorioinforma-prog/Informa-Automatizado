"""
Cruzamentos (Tabelas) — lógica de cálculo e exportação Excel.

Porte fiel do `crosstabs_page.py` do app desktop "Relatória Analítica"
(classe QtBridge: _prepare_df, _weighted_crosstab, _pct_matrix,
_sort_rows_by_total_desc, computeCrosstabs e a exportação Excel). As
listas de ordenação (FIXED_BOTTOM_LABELS, PRIORITY_ROW_SEQUENCES,
VBA_EXCLUDE_LABELS) foram copiadas sem alteração.

Diferenças em relação ao original, só de "encanamento":
- recebe/devolve objetos Python em vez de JSON via QWebChannel;
- a exportação devolve os bytes do .xlsx (o Streamlit faz o download)
  em vez de gravar num caminho escolhido no QFileDialog;
- a reordenação/remoção de linhas e colunas (que no original era feita
  em JavaScript, nos modais) virou `reordenar_linhas` / `reordenar_colunas`.

Não depende de Streamlit.
"""
from __future__ import annotations

import io
import math
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from core.analitica_dados import canon as _canon, eh_nome_de_peso as _is_peso_name


# ============================================================
# CONFIG DE ORDENAÇÃO AUTOMÁTICA (copiado do original)
# ============================================================
FIXED_BOTTOM_LABELS = [
    "Não sei",
    "Branco/Nulo",
    "Branco / Nulo",
    "Nulo / Branco",
    "Não votarei",
    "Não irá votar",
    "Não sabe",
    "NS",
]

# Ordem fixa para escalas específicas.
# Se a tabela contiver pelo menos 2 itens dessa sequência,
# a ordem abaixo prevalece sobre a ordenação por Total.
PRIORITY_ROW_SEQUENCES = [
    [
        "Certamente voto",
        "Provavelmente voto",
        "Provavelmente não voto",
        "Certamente não voto",
        "Não conheço suficiente para votar",
        "Não sabe",
        "NS",
    ],
]

# Itens herdados do VBA que não entram na ordenação por Total.
# Eles ficam preservados e vão para o fim, mantendo a ordem original entre si.
VBA_EXCLUDE_LABELS = [
    "Base",
    "Base reduzida",
    "Outros",
    "Direita",
    "Centro",
    "Esquerda",
    "Branco/Nulo",
    "Votaria em todos",
    "Nenhum",
    "Aumenta",
    "Diminui",
    "Não sabe",
    "Concordo totalmente",
    "Nem concordo nem discordo",
    "Discordo totalmente",
    "Um candidato opositor a Família Reis ",
    "Um candidato da Família Reis ",
    "Um candidato aliado aos Reis, mas que não seja da família; ",
    "Concordo plenamente",
    "Concordo",
    "Discordo",
    "Discordo plenamente",
    "Nenhuma  palavra",
    "É conhecedor de Brasília e das prioridades das pessoas ",
    "Não conhece a cidade e o que as pessoas precisam ",
    "Tem vontade de trazer melhorias para Brasília",
    "Não tem vontade de trazer melhorias para Brasília",
    "É conhecedor de Brasília e das prioridades das pessoas",
    "Não conhece a cidade e o que as pessoas precisam",
    "Está preparado para exercer o cargo",
    "Não está preparado para exercer o cargo",
    "Alta associação",
    "Associação moderada",
    "Baixa associação",
    "Diariamente",
    "Algumas vezes por semana",
    "Algumas vezes no mês",
    "Raramente",
    "Nunca",
    "Grande interesse",
    "Interesse moderado",
    "Baixo interesse",
    "Certamente voto",
    "Certamente não voto",
    "Não conheço suficiente para votar",
    "Provavelmente voto",
    "Provavelmente não voto",
    "Todas as alternativas",
    "Nenhum desses",
    "Branco / Nulo",
    "Ajuda muito na decisão",
    "Ajuda na decisão",
    "Atrapalha na decisão",
    "Atrapalha muito na decisão",
    "Preparado",
    "Despreparado",
    "Experiente",
    "Inexperiente",
    "Tem influência política",
    "Não tem influência política",
    "Conhece a cidade",
    "Não sei",
    "Não votarei",
    "Regular",
    "Aprovação",
    "Reprovação",
    "Nenhum deles",
    "Melhorou",
    "Piorou",
    "Continua da mesma forma",
    "Nada a melhorar",
    "Vai melhorar",
    "Ficar como está",
    "Vai piorar",
    "Confia muito",
    "Confia",
    "Confia mais ou menos",
    "Confia pouco",
    "Não confia",
    "Trocaria",
    "Não trocaria",
    "Outras cidades",
    "Não faz nada",
    "Nada contra a gestão atual",
    "Votou branco/nulo",
    "Não irá votar",
    "Está preparado",
    "Não está preparado",
    "Tem capacidade",
    "Não tem capacidade",
    "Tem força política",
    "Não tem força política",
    "Tem autonomia",
    "Não tem autonomia",
    "Ajudaria",
    "Não ajudaria",
    "Demais, deveria diminuir",
    "Na medida certa",
    "Pouca, deveria aumentar",
    "Lembra",
    "Não lembra",
    "Desde que nasceu",
    "Mora há mais de 20 anos",
    "Mora entre 20 e 10 anos",
    "Mora entre 10 e 5 anos",
    "Mora há menos de 5 anos",
    "Irá cumprir todas as promessas que faz",
    "Irá cumprir somente algumas",
    "Não irá cumprir nenhuma promessa que faz",
    "Promete muito mais do que conseguirá fazer em Maricá",
    "Promete na medida certa",
    "Promete pouco para Maricá, pois conseguiria fazer mais",
    "Sim",
    "Não",
    "Indiferente",
    "Outros aspectos",
    "Nada de bom",
    "Não vê nada de ruim",
    "Não sente falta de nada",
    "Muito interesse",
    "Interesse regular",
    "Nenhum interesse",
    "Excelente",
    "Boa",
    "Ruim",
    "Péssima",
    "Votou Branco ou Nulo",
    "Repro- vação",
    "Média",
    "Otimista",
    "Pessimista",
    "Nada importante",
    "Pouco importante",
    "Importante",
    "Muito importante",
    "NS",
    "Evoluindo",
    "Mesma coisa / parada",
    "Regredindo",
    "Da esquerda",
    "Da direita",
    "Do centro",
    "Conheço bem",
    "Conheço de ouvir falar",
    "Conheço das redes sociais",
    "Não conheço",
    "Indeciso",
    "Aumenta minha chance de votar nesse candidato",
    "É indiferente o apoio",
    "Diminui a minha chance de votar nesse candidato",
    "Não lembra/ Não tem candidato",
    "Não sabe em quem votar",
    "Outros candidatos",
    "Nulo / Branco",
    "Outras áreas",
    "Nenhuma área",
    "Todas as áreas",
    "Outras necessidades",
    "Nenhuma necessidade",
    "Conhece um pouco",
    "Nunca ouvi falar",
    "Conhece muito",
    "Influencia meu voto no candidato que apoiar",
    "Posso considerar o voto no candidato apoiado",
    "Posso considerar não votar no candidato apoiado",
    "Influencia a não votar no candidato apoiado",
    "Conheço muito",
    "Conheço um pouco",
    "Conheço somente de foto",
    "Conheço de foto, mas, não sabia que era Rodrigo Bacellar",
    "Nunca ouvi falar e nunca vi",
    "Tenho informações positivas",
    "Tenho informações negativas",
    "Não tenho informações",
    "Já viu",
    "Nunca viu",
    "Nenhum dos dois governos está bem",
    "Concorda totalmente",
    "Concorda em parte",
    "Não concorda e nem discorda",
    "Discorda em parte",
    "Discorda totalmente",
    "Melhor",
    "Igual",
    "Pior",
    "Todos devem permanecer presos e ser condenados",
    "A prisão deve ser apenas para quem danificou prédios",
    "Todos devem receber anistia",
    "Maricá oferece vagas de emprego até demais para os moradores",
    "Maricá oferece emprego na medida certa, todos poderiam estar empregados",
    "Maricá oferece poucas vagas de emprego, e vai oferecer cada vez menos",
    "Maricá oferece poucas vagas de emprego, mas, começará a oferecer mais",
    "O turismo está crescendo, mas falta gerar emprego e renda aos moradores",
    "O turismo está crescendo e tem ajudado a gerar emprego e renda aos moradores",
    "O turismo está estagnado",
    "O turismo está diminuindo",
    "Maricá está crescendo, mas sem planejamento e ficará pior",
    "Maricá está crescendo, sem planejamento, mas haverá uma organização",
    "Maricá já está crescendo de forma organizada",
    "Maricá está estagnada",
    "Maricá está regredindo",
]


# ==========================
# Helpers
# ==========================
def _round_half_up(x: float, ndigits: int = 1) -> float:
    q = Decimal("1").scaleb(-ndigits)
    return float(Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP))


def _safe_str(x: Any) -> str:
    if x is None:
        return ""
    try:
        if isinstance(x, float) and math.isnan(x):
            return ""
    except Exception:
        pass
    return str(x)


def _dedupe_keep_order(values: List[str]) -> List[str]:
    seen = set()
    out: List[str] = []
    for v in values:
        key = _canon(v)
        if key in seen:
            continue
        seen.add(key)
        out.append(v)
    return out


def _apply_priority_sequences(labels: List[str], sequences: List[List[str]], min_present: int = 2) -> Tuple[List[str], List[str]]:
    remaining = labels[:]
    prefix: List[str] = []

    for seq in sequences:
        present: List[str] = []
        for wanted in seq:
            wanted_c = _canon(wanted)
            match = next((lab for lab in remaining if _canon(lab) == wanted_c), None)
            if match is not None:
                present.append(match)

        present = _dedupe_keep_order(present)
        if len(present) < min_present:
            continue

        for lab in present:
            remaining = [x for x in remaining if _canon(x) != _canon(lab)]
        prefix.extend(present)

    return prefix, remaining


def _apply_fixed_bottom_order(labels: List[str], fixed_bottom: List[str]) -> List[str]:
    fixed_canon = [_canon(x) for x in fixed_bottom]
    fixed_canon_set = set(fixed_canon)
    present = labels[:]

    fixed_present: List[str] = []
    for f in fixed_canon:
        for lab in present:
            if _canon(lab) == f and lab not in fixed_present:
                fixed_present.append(lab)
                break

    rest = [lab for lab in present if _canon(lab) not in fixed_canon_set]
    return rest + fixed_present


# ==========================
# Cálculo
# ==========================
def _prepare_df(
    df_full: pd.DataFrame,
    value_labels: Dict[str, Dict[Any, str]],
    row_var: str,
    col_var: str,
    weight_var: str,
    include_na: bool,
) -> Tuple[pd.DataFrame, Optional[pd.Series]]:
    needed = [row_var, col_var]
    if weight_var:
        needed.append(weight_var)

    src = df_full.loc[:, needed]
    if include_na:
        df = src.copy()
    else:
        mask = src[row_var].notna() & src[col_var].notna()
        df = src.loc[mask].copy()

    for c in (row_var, col_var):
        m = value_labels.get(c)
        if m:
            df[c] = df[c].map(lambda v, labels=m: labels.get(v, v))

    if include_na:
        df[row_var] = df[row_var].where(df[row_var].notna(), "(Sem resposta)")
        df[col_var] = df[col_var].where(df[col_var].notna(), "(Sem resposta)")

    df[row_var] = df[row_var].map(_safe_str)
    df[col_var] = df[col_var].map(_safe_str)

    w = None
    if weight_var:
        w = pd.to_numeric(df[weight_var], errors="coerce").fillna(0.0)
    return df, w


def _weighted_crosstab(df: pd.DataFrame, row_var: str, col_var: str, w: Optional[pd.Series]) -> pd.DataFrame:
    if w is None:
        return pd.crosstab(df[row_var], df[col_var], dropna=False).astype(float)

    ct = pd.crosstab(
        df[row_var],
        df[col_var],
        values=w,
        aggfunc="sum",
        dropna=False,
    ).fillna(0.0)
    return ct.astype(float)


def _unweighted_bases(df: pd.DataFrame, row_var: str, col_var: str) -> pd.DataFrame:
    return pd.crosstab(df[row_var], df[col_var], dropna=False)


def _pct_matrix(ct: pd.DataFrame, mode: str) -> pd.DataFrame:
    if mode == "col":
        denom = ct.sum(axis=0).replace(0, float("nan"))
        p = ct.divide(denom, axis=1) * 100.0
    elif mode == "row":
        denom = ct.sum(axis=1).replace(0, float("nan"))
        p = ct.divide(denom, axis=0) * 100.0
    elif mode == "total":
        denom = float(ct.values.sum())
        p = (ct / denom) * 100.0 if denom else ct * 0.0
    else:
        p = ct * 0.0

    p = p.fillna(0.0)
    return p.apply(lambda s: s.map(lambda v: _round_half_up(float(v), 1)))


def _pct_total_col(ct: pd.DataFrame) -> List[float]:
    row_sums = ct.sum(axis=1)
    grand = float(ct.values.sum()) if ct.values.size else 0.0
    if grand <= 0:
        return [0.0 for _ in row_sums]
    out = (row_sums / grand) * 100.0
    return [_round_half_up(float(v), 1) for v in out.tolist()]


def _sort_rows_by_total_desc(
    ct: pd.DataFrame,
    fixed_bottom: List[str],
    exclude_labels: Optional[List[str]] = None,
    priority_sequences: Optional[List[List[str]]] = None,
) -> List[str]:
    original_labels = [str(x) for x in list(ct.index.astype(str))]
    exclude_canon = {_canon(x) for x in (exclude_labels or [])}
    priority_sequences = priority_sequences or []

    priority_prefix, remaining_labels = _apply_priority_sequences(original_labels, priority_sequences)

    grand = float(ct.values.sum()) if ct.values.size else 0.0
    if grand > 0:
        row_sums = ct.sum(axis=1)
        pct = (row_sums / grand) * 100.0
    else:
        pct = pd.Series(0.0, index=ct.index)

    pinned = [lab for lab in remaining_labels if _canon(lab) in exclude_canon]
    sortable = [lab for lab in remaining_labels if _canon(lab) not in exclude_canon]

    if grand > 0:
        sortable.sort(key=lambda lab: float(pct.get(lab, 0.0)), reverse=True)

    ordered_tail = _apply_fixed_bottom_order(sortable + pinned, fixed_bottom)
    return priority_prefix + ordered_tail


def calcular_cruzamentos(
    df_full: pd.DataFrame,
    var_labels: Dict[str, str],
    value_labels: Dict[str, Dict[Any, str]],
    row_var: str,
    col_vars: List[str],
    weight_var: str,
    include_na: bool,
    pct_mode: str,
) -> List[Dict[str, Any]]:
    """Equivalente ao `computeCrosstabs` original. Devolve a lista de
    tabelas (dicts com o mesmo formato que o JS recebia)."""
    if not isinstance(col_vars, list) or not col_vars:
        raise ValueError("Lista de colunas inválida.")

    row_var = (row_var or "").strip()
    weight_var = (weight_var or "").strip()

    if weight_var and not _is_peso_name(weight_var):
        weight_var = ""
    if weight_var and weight_var not in df_full.columns:
        weight_var = ""

    tables: List[Dict[str, Any]] = []
    row_var_display = (var_labels.get(row_var) or "").strip() or row_var

    for col_var in col_vars:
        col_var = (col_var or "").strip()
        if not col_var:
            continue

        df_prep, w = _prepare_df(df_full, value_labels, row_var, col_var, weight_var, include_na)
        base_ct = _unweighted_bases(df_prep, row_var, col_var)
        if base_ct.empty:
            continue

        base_row_labels = [str(x) for x in base_ct.index.tolist()]
        base_col_labels = [str(x) for x in base_ct.columns.tolist()]
        base_total_n = int(base_ct.to_numpy().sum())

        ct = _weighted_crosstab(df_prep, row_var, col_var, w)
        ct = ct.reindex(index=base_row_labels, columns=base_col_labels, fill_value=0.0)

        ordered_rows = _sort_rows_by_total_desc(
            ct,
            FIXED_BOTTOM_LABELS,
            exclude_labels=VBA_EXCLUDE_LABELS,
            priority_sequences=PRIORITY_ROW_SEQUENCES,
        )
        ct = ct.reindex(index=ordered_rows, fill_value=0.0)
        base_ct = base_ct.reindex(index=ordered_rows, columns=base_col_labels, fill_value=0)

        pct_m = _pct_matrix(ct, pct_mode)
        pct_total = _pct_total_col(ct)

        col_labels = base_col_labels
        row_labels = [str(x) for x in ct.index.tolist()]
        base_col_n_aligned = base_ct.sum(axis=0).astype(int).tolist()

        weighted_total = float(ct.values.sum())
        base_display = f"Base (N): {base_total_n}"
        if weight_var:
            base_display += f" | Ponderada: {weighted_total:.1f}"
        else:
            base_display += " | Sem ponderação"

        col_var_display = (var_labels.get(col_var) or "").strip() or col_var
        title = f"{row_var_display} × {col_var_display}"

        tables.append({
            "key": f"{row_var}__{col_var}",
            "title": title,
            "row_var": row_var,
            "col_var": col_var,
            "row_var_display": row_var_display,
            "col_var_display": col_var_display,
            "pct_mode": pct_mode,
            "weight_var": weight_var,
            "include_na": bool(include_na),
            "fixed_bottom_labels": FIXED_BOTTOM_LABELS,
            "priority_row_sequences": PRIORITY_ROW_SEQUENCES,
            "row_labels": row_labels,
            "col_labels": col_labels,
            "pct_total": pct_total,
            "pct_matrix": pct_m.values.tolist(),
            "base_total_n": int(base_total_n),
            "base_col_n": base_col_n_aligned,
            "base_display": base_display,
        })

    return tables


# ==========================
# Ajustes manuais (antes feitos no JS, nos modais)
# ==========================
def reordenar_linhas(t: Dict[str, Any], novos_rotulos: List[str]) -> Dict[str, Any]:
    """`applyOrderToTableRows` do JS original: reordena as linhas (rótulos
    que não existirem na tabela são ignorados)."""
    old = list(t["row_labels"])
    idx_map = {lab: i for i, lab in enumerate(old)}
    new_total, new_matrix, new_labels = [], [], []
    for lab in novos_rotulos:
        i = idx_map.get(lab)
        if i is None:
            continue
        new_labels.append(lab)
        new_total.append(t["pct_total"][i])
        new_matrix.append(t["pct_matrix"][i])
    t["row_labels"] = new_labels
    t["pct_total"] = new_total
    t["pct_matrix"] = new_matrix
    return t


def reordenar_colunas(t: Dict[str, Any], indices_ativos: List[int]) -> Dict[str, Any]:
    """`applyOrderToTableCols` do JS original: `indices_ativos` são os
    índices das colunas que ficam, na nova ordem (as removidas somem).
    O Total (coluna B) não muda."""
    t["col_labels"] = [t["col_labels"][i] for i in indices_ativos]
    t["base_col_n"] = [t["base_col_n"][i] for i in indices_ativos]
    t["pct_matrix"] = [[row[i] for i in indices_ativos] for row in t["pct_matrix"]]
    return t


def fmt_pct(x: Any) -> str:
    if x is None:
        return ""
    try:
        if math.isnan(float(x)):
            return ""
    except Exception:
        return ""
    return f"{float(x):.1f}".replace(".", ",")


# ==========================
# Excel Export (bordas mínimas)
# ==========================
def _excel_styles() -> Dict[str, Any]:
    return {
        "font_title": Font(name="DIN", size=12, bold=False),
        "font_header": Font(name="DIN", size=10, bold=False),
        "font_rowlab": Font(name="DIN", size=10, bold=False),
        "font_num": Font(name="DIN", size=10, bold=False),
        "align_center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "align_left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "thin": Side(style="thin", color="000000"),
        "thick": Side(style="thick", color="000000"),
        "none": Side(style=None),
    }


def _compute_colA_width(tables: List[Dict[str, Any]]) -> float:
    max_len = len("Categoria")
    for t in tables:
        for lab in (t.get("row_labels") or []):
            max_len = max(max_len, len(str(lab)))
        max_len = max(max_len, len("Base"))
    return min(max(12.0, max_len * 1.05), 70.0)


def _apply_table_borders_minimal(ws, top: int, bottom: int, left: int, right: int,
                                 none: Side, thin: Side, thick: Side) -> None:
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            l = none
            t = none
            b = none
            rr = thin

            if c == left:
                l = thick
            if c == right:
                rr = thick
            if r == top:
                t = thick
            if r == bottom:
                b = thick

            if r == top:
                l = thick
                rr = thick
                t = thick
                b = thick

            ws.cell(r, c).border = Border(left=l, right=rr, top=t, bottom=b)


def _write_table_block(ws, start_row: int, t: Dict[str, Any], styles: Dict[str, Any]) -> int:
    title = (t.get("title") or "").strip()
    row_labels: List[str] = t["row_labels"]
    col_labels: List[str] = t["col_labels"]
    pct_total: List[float] = t["pct_total"]
    pct_matrix: List[List[float]] = t["pct_matrix"]
    base_total_n: int = int(t["base_total_n"])
    base_col_n: List[int] = [int(x) for x in t["base_col_n"]]

    last_col = 2 + len(col_labels)

    r_title = start_row
    ctitle = ws.cell(r_title, 1, title)
    ctitle.font = styles["font_title"]
    ctitle.alignment = styles["align_center"]
    ws.merge_cells(start_row=r_title, start_column=1, end_row=r_title, end_column=last_col)

    r_blank = r_title + 1
    r_head = r_blank + 1
    headers = ["Categoria", "Total"] + col_labels
    for j, h in enumerate(headers, start=1):
        c = ws.cell(r_head, j, h)
        c.font = styles["font_header"]
        c.alignment = styles["align_left"] if j == 1 else styles["align_center"]

    r0 = r_head + 1
    for i, lab in enumerate(row_labels):
        rr = r0 + i
        c0 = ws.cell(rr, 1, lab)
        c0.font = styles["font_rowlab"]
        c0.alignment = styles["align_left"]

        cT = ws.cell(rr, 2, float(pct_total[i]))
        cT.font = styles["font_num"]
        cT.alignment = styles["align_center"]
        cT.number_format = "0.0"

        for j in range(len(col_labels)):
            v = float(pct_matrix[i][j])
            cc = ws.cell(rr, 3 + j, v)
            cc.font = styles["font_num"]
            cc.alignment = styles["align_center"]
            cc.number_format = "0.0"

    r_base = r0 + len(row_labels)
    cB0 = ws.cell(r_base, 1, "Base")
    cB0.font = styles["font_rowlab"]
    cB0.alignment = styles["align_left"]

    cBT = ws.cell(r_base, 2, int(base_total_n))
    cBT.font = styles["font_rowlab"]
    cBT.alignment = styles["align_center"]
    cBT.number_format = "0"

    for j in range(len(col_labels)):
        cc = ws.cell(r_base, 3 + j, int(base_col_n[j]))
        cc.font = styles["font_rowlab"]
        cc.alignment = styles["align_center"]
        cc.number_format = "0"

    _apply_table_borders_minimal(
        ws, top=r_head, bottom=r_base, left=1, right=last_col,
        none=styles["none"], thin=styles["thin"], thick=styles["thick"],
    )
    return r_base


def exportar_excel(tabelas: List[Dict[str, Any]], nome_aba: str = "Cruzamentos") -> bytes:
    """Equivalente a `exportQueueExcel` (pilha) / `exportSingleExcel`
    (1 tabela — chame com `nome_aba="Cruzamento"`). Devolve os bytes."""
    if not tabelas:
        raise ValueError("A pilha está vazia.")
    wb = Workbook()
    ws = wb.active
    ws.title = nome_aba

    styles = _excel_styles()
    ws.column_dimensions["A"].width = _compute_colA_width(tabelas)

    max_cols = 0
    for t in tabelas:
        max_cols = max(max_cols, 2 + len(t.get("col_labels") or []))
    for c in range(2, max_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 8.89

    r = 1
    for t in tabelas:
        last_used = _write_table_block(ws, r, t, styles)
        r = last_used + 3

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
