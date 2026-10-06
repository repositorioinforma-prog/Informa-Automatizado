"""
Mapa de Calor — cálculos do app desktop "Relatória Analítica".

O trecho entre as marcas "INÍCIO/FIM DO CÓDIGO ORIGINAL" é cópia SEM
ALTERAÇÕES de app/pages/heatmap_page.py (linhas 1112–2100): leitura dos
casos no Relatório Descritivo (Excel), arredondamento e geração do PPTX.
Abaixo dele ficam só os embrulhos para o Streamlit (arquivos em memória no
lugar de caminhos no disco) e as contas dos @Slot computeStats /
applyAutoPaint da QtBridge original, também sem mudar as regras.
"""
from __future__ import annotations

import io
import json
import math
import re
import unicodedata
from decimal import Decimal, ROUND_HALF_UP
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

# ===================== INÍCIO DO CÓDIGO ORIGINAL (heatmap_page.py) =====================



# =========================
# Excel parsing (com debug)
# =========================

SEGMENT_LABELS = [
    "Certamente voto",
    "Certamente não voto",
    "Aprovação",
    "Reprovação",
]

REQUIRED_SEGMENT_PAIRS = [
    {"Certamente voto", "Certamente não voto"},
    {"Aprovação", "Reprovação"},
]

MARKER_LABELS = {
    "municipio", "municipios",
    "bairro", "bairros",
    "distrito", "distritos",
    "subprefeitura", "subprefeituras",
    "zona", "zonas",
    "mesorregiao", "mesorregioes",
    "microrregiao", "microrregioes",
    "regiao administrativa", "regioes administrativas",
    "regiao", "regioes",
    "area de planejamento", "areas de planejamento",
    "ap", "ap areas de planejamento",
    "rpa", "za",
    "zona administrativa", "zonas administrativas",
    "zona administrativa da capital", "zonas administrativas da capital",
    "uf", "estado", "estados", "capital", "capitais",
}


IGNORE_HEADER_WORDS = {"total", "base", "geral"}


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.lower()
    text = " ".join(text.split())
    return text


def is_marker_text(value: Any) -> bool:
    """
    Reconhece apenas rótulos territoriais usados como marcador de tabela.
    Evita confundir palavras comuns dentro de títulos e perguntas.
    """
    normalized = normalize_text(value)
    if not normalized:
        return False

    cleaned = re.sub(r"[^a-z0-9 ]+", " ", normalized)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    if cleaned in MARKER_LABELS:
        return True

    for label in MARKER_LABELS:
        if cleaned.startswith(label + " "):
            remainder = cleaned[len(label):].strip()
            if remainder in {"", "total", "percentual", "porcentagem", "segmento"}:
                return True

    return False


def parse_number(cell) -> Optional[float]:
    raw = cell.value
    if raw is None:
        return None

    num_format = (cell.number_format or "").lower()

    if isinstance(raw, (int, float)):
        value = float(raw)
        if "%" in num_format and 0 <= value <= 1.0:
            return value * 100.0
        return value

    text = str(raw).strip()
    if not text:
        return None

    has_percent = "%" in text or "%" in num_format
    text = text.replace("%", "").strip()

    if "," not in text and re.match(r"^\d{1,3}(\.\d{3})+$", text):
        text = text.replace(".", "")

    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")

    try:
        value = float(text)
    except ValueError:
        return None

    if has_percent and 0 <= value <= 1.0:
        value *= 100.0

    return value


def get_row_label_text(sheet, row_index: int, max_label_cols: int = 3) -> Optional[str]:
    for col_index in range(1, max_label_cols + 1):
        v = sheet.cell(row_index, col_index).value
        if isinstance(v, str) and v.strip():
            return v.strip()
    return None


def row_has_segment_label(sheet, row_index: int) -> Optional[str]:
    label = get_row_label_text(sheet, row_index, max_label_cols=3)
    if not label:
        return None
    label_norm = normalize_text(label)
    for seg in SEGMENT_LABELS:
        if normalize_text(seg) in label_norm:
            return seg
    return None


def row_is_base(sheet, row_index: int) -> bool:
    label = get_row_label_text(sheet, row_index, max_label_cols=3)
    if not label:
        return False
    return normalize_text(label).startswith("base")


def row_contains_marker(sheet, row_index: int, max_check_cols: int = 10) -> bool:
    for col_index in range(1, max_check_cols + 1):
        v = sheet.cell(row_index, col_index).value
        if isinstance(v, str) and v.strip() and is_marker_text(v):
            return True
    return False


def find_segment_rows_below_marker(sheet, marker_row: int, max_row: int, scan_limit: int = 500) -> List[int]:
    segment_rows: List[int] = []
    end_row = min(max_row, marker_row + scan_limit)
    for row_index in range(marker_row, end_row + 1):
        if row_has_segment_label(sheet, row_index):
            segment_rows.append(row_index)
    return segment_rows


def pick_header_row_from_segments(sheet, first_segment_row: int, max_col: int, lookback: int = 10) -> Optional[int]:
    if first_segment_row <= 1:
        return None

    candidates: List[Tuple[int, int]] = []
    start_row = max(1, first_segment_row - lookback)
    end_row = first_segment_row - 1

    for row_index in range(start_row, end_row + 1):
        if row_has_segment_label(sheet, row_index) or row_is_base(sheet, row_index):
            continue

        text_cells = 0
        numeric_cells = 0
        nonempty_cells = 0

        for col_index in range(2, max_col + 1):
            v = sheet.cell(row_index, col_index).value
            if v in (None, ""):
                continue
            nonempty_cells += 1
            if isinstance(v, str) and v.strip():
                text_cells += 1
            elif isinstance(v, (int, float)):
                numeric_cells += 1

        if text_cells >= 2:
            score = text_cells * 10 - numeric_cells * 20 + nonempty_cells
            candidates.append((score, row_index))

    if not candidates:
        return end_row if end_row >= 1 else None

    candidates.sort(reverse=True)
    return candidates[0][1]


def compute_last_table_col(sheet, header_row: int, first_segment_row: int, max_col: int) -> int:
    last_col = 1
    for col_index in range(1, max_col + 1):
        if sheet.cell(header_row, col_index).value not in (None, ""):
            last_col = col_index
        if sheet.cell(first_segment_row, col_index).value not in (None, ""):
            last_col = max(last_col, col_index)
    return last_col


def find_total_column(sheet, header_row: int, last_table_col: int) -> Tuple[int, bool]:
    """
    Procura a coluna Total em várias linhas próximas ao cabeçalho.
    Não assume automaticamente que a coluna B é o Total.
    """
    scan_start = max(1, header_row - 8)
    scan_end = min(sheet.max_row or header_row, header_row + 2)

    for scan_row in range(scan_end, scan_start - 1, -1):
        for col_index in range(1, last_table_col + 1):
            value = sheet.cell(scan_row, col_index).value
            if value is None:
                continue
            if normalize_text(value) in {"total", "geral", "total geral"}:
                return col_index, True

    for merged in sheet.merged_cells.ranges:
        if merged.max_row < scan_start or merged.min_row > scan_end:
            continue
        value = sheet.cell(merged.min_row, merged.min_col).value
        if value is not None and normalize_text(value) in {"total", "geral", "total geral"}:
            return merged.min_col, True

    return 0, False



def extract_base_total_from_row(
    sheet,
    base_row: int,
    total_col: int,
    last_table_col: int,
) -> Tuple[Optional[float], Optional[int]]:
    """
    Lê a base total. Primeiro usa a coluna Total detectada. Se ela não existir,
    usa o maior valor numérico da linha Base, que normalmente é a base geral.
    """
    if total_col and total_col > 0:
        value = parse_number(sheet.cell(base_row, total_col))
        if isinstance(value, (int, float)) and value > 0:
            return float(value), total_col

    candidates: List[Tuple[float, int]] = []
    for col_index in range(2, last_table_col + 1):
        value = parse_number(sheet.cell(base_row, col_index))
        if isinstance(value, (int, float)) and value > 0:
            candidates.append((float(value), col_index))

    if not candidates:
        return None, None

    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0]

def extract_region_columns(sheet, header_row: int, last_table_col: int) -> Tuple[List[str], List[int], List[str]]:
    region_names: List[str] = []
    region_cols: List[int] = []
    raw_header_labels: List[str] = []

    for col_index in range(2, last_table_col + 1):
        header_value = sheet.cell(header_row, col_index).value
        if header_value is None:
            continue
        header_text = str(header_value).strip()
        if not header_text:
            continue

        raw_header_labels.append(header_text)
        header_norm = normalize_text(header_text)

        if header_norm in IGNORE_HEADER_WORDS:
            continue
        if header_norm == "total":
            continue
        if is_marker_text(header_text) and not header_norm.startswith("ap "):
            continue

        region_names.append(header_text)
        region_cols.append(col_index)

    return region_names, region_cols, raw_header_labels


def find_title_above(sheet, header_row: int, first_col: int, last_col: int) -> Optional[str]:
    merged_ranges = list(sheet.merged_cells.ranges)

    for scan_row in range(max(1, header_row - 30), header_row)[::-1]:
        for rng in merged_ranges:
            if rng.min_row == scan_row and rng.max_row == scan_row and rng.min_col <= first_col and rng.max_col >= last_col:
                v = sheet.cell(rng.min_row, rng.min_col).value
                if v is not None and str(v).strip():
                    return str(v).strip()

        a_val = sheet.cell(scan_row, 1).value
        if a_val is not None and str(a_val).strip():
            empty_count = 0
            width = last_col - first_col + 1
            for col_index in range(first_col, last_col + 1):
                if sheet.cell(scan_row, col_index).value in (None, ""):
                    empty_count += 1
            if empty_count >= int(0.85 * width):
                return str(a_val).strip()

    return None


def detect_table_kind(marker_text: str, header_labels: List[str]) -> str:
    """
    Identifica automaticamente o tipo de recorte territorial.
    """
    nm = normalize_text(marker_text)
    header_norms = [normalize_text(h) for h in header_labels]

    if (
        "ap (areas de planejamento)" in nm
        or "area de planejamento" in nm
        or "areas de planejamento" in nm
        or nm == "ap"
        or any(h == "ap" or h.startswith("ap ") for h in header_norms)
    ):
        return "Áreas de Planejamento"

    if nm == "rpa" or nm.startswith("rpa ") or any(h == "rpa" or h.startswith("rpa ") for h in header_norms):
        return "RPA"

    if (
        nm == "za"
        or "zona administrativa" in nm
        or "zonas administrativas" in nm
        or any(h == "za" or h.startswith("za ") for h in header_norms)
    ):
        return "ZA"

    if "regiao administrativa" in nm or "regioes administrativas" in nm:
        return "Regiões Administrativas"

    if "mesorregiao" in nm or "mesorregioes" in nm:
        return "Mesorregiões"

    if "microrregiao" in nm or "microrregioes" in nm:
        return "Microrregiões"

    if "municipio" in nm or "municipios" in nm:
        return "Municípios"

    if "bairro" in nm or "bairros" in nm:
        return "Bairros"

    if "distrito" in nm or "distritos" in nm:
        return "Distritos"

    if "subprefeitura" in nm or "subprefeituras" in nm:
        return "Subprefeituras"

    if nm == "zona" or "zonas" in nm:
        return "Zonas"

    if nm == "uf" or nm.startswith("uf ") or any(h == "uf" for h in header_norms):
        return "UF"

    if "estado" in nm or "estados" in nm:
        return "Estados"

    if "capital" in nm or "capitais" in nm:
        return "Capitais"

    if "regiao" in nm or "regioes" in nm:
        return "Regiões"

    return "Tabela"


def parse_excel_cases_with_debug(excel_path: str) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Any]]:
    workbook = openpyxl.load_workbook(excel_path, data_only=True)

    cases_by_key: Dict[str, Dict[str, Any]] = {}
    debug_candidates: List[Dict[str, Any]] = []
    seen_tables: Set[Tuple[str, int, int]] = set()

    for sheet in workbook.worksheets:
        max_row = min(sheet.max_row or 0, 20000)
        max_col = min(sheet.max_column or 0, 300)
        if max_row == 0 or max_col == 0:
            continue

        for marker_row in range(1, max_row + 1):
            # Marcadores territoriais costumam ficar nas primeiras colunas.
            for marker_col in range(1, min(max_col, 12) + 1):
                marker_value = sheet.cell(marker_row, marker_col).value
                if not isinstance(marker_value, str):
                    continue
                if not is_marker_text(marker_value):
                    continue

                marker_text = str(marker_value).strip()

                segment_rows = find_segment_rows_below_marker(sheet, marker_row, max_row, scan_limit=500)
                if not segment_rows:
                    continue

                first_segment_row = min(segment_rows)

                header_row = pick_header_row_from_segments(sheet, first_segment_row, max_col, lookback=10)
                if header_row is None:
                    continue

                last_table_col = compute_last_table_col(sheet, header_row, first_segment_row, max_col)
                table_dedupe_key = (sheet.title, header_row, last_table_col)
                if table_dedupe_key in seen_tables:
                    continue
                seen_tables.add(table_dedupe_key)

                total_col, _found_total = find_total_column(sheet, header_row, last_table_col)
                region_names, region_cols, raw_header_labels = extract_region_columns(sheet, header_row, last_table_col)
                if len(region_names) < 2:
                    continue

                table_kind = detect_table_kind(marker_text, raw_header_labels)

                base_total: Optional[float] = None
                base_row: Optional[int] = None

                segment_values_by_label: Dict[str, Dict[str, float]] = {}
                segment_sources_by_label: Dict[str, Dict[str, Dict[str, Any]]] = {}
                segment_row_by_label: Dict[str, int] = {}

                blank_run = 0
                scan_start = header_row + 1
                scan_end = min(max_row, header_row + 1200)

                for data_row in range(scan_start, scan_end + 1):
                    row_is_empty = True
                    for col_index in range(1, min(last_table_col, 20) + 1):
                        if sheet.cell(data_row, col_index).value not in (None, ""):
                            row_is_empty = False
                            break
                    if row_is_empty:
                        blank_run += 1
                    else:
                        blank_run = 0
                    if blank_run >= 2 and data_row > first_segment_row:
                        break

                    if data_row > first_segment_row and row_contains_marker(sheet, data_row, max_check_cols=12):
                        break

                    label_text = get_row_label_text(sheet, data_row, max_label_cols=3)
                    if not label_text:
                        continue
                    label_norm = normalize_text(label_text)

                    if label_norm.startswith("base"):
                        if base_total is None:
                            detected_base, detected_col = extract_base_total_from_row(
                                sheet,
                                data_row,
                                total_col,
                                last_table_col,
                            )
                            base_total = detected_base
                            base_row = data_row
                            if detected_col:
                                total_col = detected_col
                        continue

                    matched_segment = row_has_segment_label(sheet, data_row)
                    if not matched_segment:
                        continue

                    if matched_segment in segment_row_by_label:
                        continue

                    segment_row_by_label[matched_segment] = data_row
                    segment_values_by_label.setdefault(matched_segment, {})
                    segment_sources_by_label.setdefault(matched_segment, {})

                    for region_name, region_col in zip(region_names, region_cols):
                        cell = sheet.cell(data_row, region_col)
                        parsed = parse_number(cell)
                        if parsed is None:
                            continue
                        addr = f"{get_column_letter(region_col)}{data_row}"
                        segment_values_by_label[matched_segment][region_name] = float(parsed)
                        segment_sources_by_label[matched_segment][region_name] = {
                            "addr": addr,
                            "row": data_row,
                            "col": region_col,
                            "raw": cell.value,
                            "num_format": cell.number_format,
                            "parsed": float(parsed),
                        }

                if base_total is None or not isinstance(base_total, (int, float)) or base_total <= 0:
                    continue

                segment_keys = set(segment_values_by_label.keys())
                ok = any(pair.issubset(segment_keys) for pair in REQUIRED_SEGMENT_PAIRS)
                if not ok:
                    continue

                min_required_fill = max(3, int(0.70 * len(region_names)))
                if any(len(segment_values_by_label.get(seg, {})) < min_required_fill for seg in segment_keys):
                    pass

                title = find_title_above(sheet, header_row, 1, last_table_col) or "(Sem título)"
                case_key = (
                    f"{title} — {table_kind} — Aba: {sheet.title} — "
                    f"Header: {header_row} — Base: {base_row if base_row else '—'}"
                )
                if case_key in cases_by_key:
                    case_key = f"{case_key} — marcador({marker_row},{marker_col})"

                cases_by_key[case_key] = {
                    "__meta": {
                        "title": title,
                        "kind": table_kind,
                        "sheet": sheet.title,
                        "marker_row": marker_row,
                        "marker_col": marker_col,
                        "marker_text": marker_text,
                        "first_segment_row": first_segment_row,
                        "header_row": header_row,
                        "last_table_col": last_table_col,
                        "total_col": total_col,
                        "region_names": region_names,
                        "region_cols": region_cols,
                        "base_row": base_row,
                        "base_total": float(base_total),
                        "segment_rows": segment_row_by_label,
                        "segment_sources": segment_sources_by_label,
                    }
                }

                for seg_label, region_map in segment_values_by_label.items():
                    cases_by_key[case_key][seg_label] = region_map

    debug_report = {"excel_path": excel_path, "cases_count": len(cases_by_key)}
    return cases_by_key, debug_report


# =========================
# PPT helpers
# =========================

def dec_round_half_up(value: float, decimals: int = 1) -> Decimal:
    """
    Arredonda em Decimal com ROUND_HALF_UP (0,55 -> 0,6).
    """
    q = Decimal("1").scaleb(-decimals)  # decimals=1 -> 0.1
    return Decimal(str(value)).quantize(q, rounding=ROUND_HALF_UP)

def round_half_up(value: float, decimals: int = 1) -> float:
    return float(dec_round_half_up(value, decimals))

def fmt_1_half_up(value: float) -> str:
    """
    Retorna string com 1 casa decimal e vírgula, usando HALF_UP.
    """
    d = dec_round_half_up(value, 1)
    return str(d).replace(".", ",")


def _hex_to_rgb(hex_color: str) -> RGBColor:
    s = (hex_color or "").strip()
    if s.startswith("#"):
        s = s[1:]
    if len(s) != 6:
        return RGBColor(0, 0, 0)
    r = int(s[0:2], 16)
    g = int(s[2:4], 16)
    b = int(s[4:6], 16)
    return RGBColor(r, g, b)


def _fmt_1(x: float) -> str:
    return fmt_1_half_up(float(x))


def _is_good_segment(seg: str) -> bool:
    t = (seg or "").strip()
    t = unicodedata.normalize("NFD", t)
    t = "".join(ch for ch in t if unicodedata.category(ch) != "Mn")
    t = t.lower()
    return ("certamente voto" in t) or ("aprovacao" in t) or ("aprovação" in t)


def build_pptx(
    out_path: str,
    slides: List[Dict[str, Any]],
    logo_bytes: Optional[bytes]
) -> None:
    """
    Gera os slides em 16:9 mantendo o mapa e o cabeçalho como estavam.

    Alterações visuais desta versão:
    - tabela da direita com apenas 2 colunas: Regiões | %;
    - a própria célula do percentual recebe a cor da região;
    - a antiga legenda horizontal abaixo do mapa foi removida;
    - Mín / Entre / Máx agora aparecem em uma pequena tabela abaixo
      da tabela de regiões, com as cores aplicadas nos rótulos.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    def text_color_for_fill(hex_color: str) -> RGBColor:
        """
        Retorna preto ou branco conforme a luminosidade da cor de fundo.
        """
        s = (hex_color or "").strip().lstrip("#")
        if len(s) != 6:
            return RGBColor(0, 0, 0)

        try:
            r = int(s[0:2], 16)
            g = int(s[2:4], 16)
            b = int(s[4:6], 16)
        except ValueError:
            return RGBColor(0, 0, 0)

        luminance = (0.299 * r + 0.587 * g + 0.114 * b) / 255.0
        return RGBColor(0, 0, 0) if luminance >= 0.56 else RGBColor(255, 255, 255)

    def set_cell_border(cell, color: str = "BFBFBF", width: int = 12700) -> None:
        """
        Aplica borda visível nos quatro lados da célula.
        width=12700 equivale aproximadamente a 1 pt.
        """
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.oxml.ns import qn

        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            ln = tcPr.find(qn(edge))
            if ln is None:
                ln = OxmlElement(edge)
                tcPr.append(ln)

            ln.set("w", str(width))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            ln.set("algn", "ctr")

            solidFill = ln.find(qn("a:solidFill"))
            if solidFill is None:
                solidFill = OxmlElement("a:solidFill")
                ln.append(solidFill)

            srgbClr = solidFill.find(qn("a:srgbClr"))
            if srgbClr is None:
                srgbClr = OxmlElement("a:srgbClr")
                solidFill.append(srgbClr)
            srgbClr.set("val", color)

            prstDash = ln.find(qn("a:prstDash"))
            if prstDash is None:
                prstDash = OxmlElement("a:prstDash")
                ln.append(prstDash)
            prstDash.set("val", "solid")

    for item in slides:
        seg = item["segment"]
        case_title = (item.get("case_title") or "").strip()
        title_text = f"{case_title} - {seg}".strip(" -") if case_title else seg

        pal = item["palette"]
        stats = item["stats"]
        png_bytes = item["png_bytes"]

        region_rows = item.get("legend_rows") or []
        has_table = bool(region_rows)

        slide = prs.slides.add_slide(blank)

        # =====================================================
        # CABEÇALHO: mantido
        # =====================================================
        slide_left = Inches(0.38)
        slide_right = Inches(0.38)

        logo_reserved_w = Inches(2.05)
        logo_gap = Inches(0.18)

        title_left = slide_left
        title_top = Inches(0.08)
        title_w = prs.slide_width - slide_left - slide_right - logo_reserved_w - logo_gap
        title_h = Inches(0.62)

        title_box = slide.shapes.add_textbox(
            title_left,
            title_top,
            title_w,
            title_h
        )
        tf = title_box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title_text
        p.font.bold = True
        p.font.name = "DIN"

        if len(title_text) <= 75:
            p.font.size = Pt(20)
        elif len(title_text) <= 115:
            p.font.size = Pt(18)
        else:
            p.font.size = Pt(16)

        if logo_bytes:
            stream = io.BytesIO(logo_bytes)
            logo_pic = slide.shapes.add_picture(
                stream,
                left=prs.slide_width - slide_right - logo_reserved_w,
                top=Inches(0.10),
                height=Inches(0.42)
            )
            logo_pic.left = prs.slide_width - slide_right - logo_pic.width

        # =====================================================
        # ÁREA PRINCIPAL: dimensões gerais mantidas
        # =====================================================
        content_left = Inches(0.72)
        content_right = Inches(0.62)
        content_top = Inches(0.76)

        map_h = Inches(5.28)

        table_gap = Inches(0.32)

        # Mantemos a mesma largura total da tabela antiga:
        # 1.55 + 0.34 + 0.66 = 2.55 pol.
        # Agora ela é dividida em apenas 2 colunas.
        name_w = Inches(1.55)
        pct_w = Inches(1.00)
        table_w = name_w + pct_w if has_table else Inches(0)

        total_content_w = prs.slide_width - content_left - content_right
        map_box_w = (
            total_content_w - table_gap - table_w
            if has_table
            else total_content_w
        )

        table_left = content_left + map_box_w + table_gap

        # =====================================================
        # MAPA: mantido exatamente com lógica "contain"
        # =====================================================
        img = Image.open(io.BytesIO(png_bytes))
        iw, ih = img.size
        img_ratio = iw / ih if ih else 1.0
        box_ratio = float(map_box_w) / float(map_h)

        if img_ratio >= box_ratio:
            pic_w = int(map_box_w)
            pic_h = int(float(map_box_w) / img_ratio)
        else:
            pic_h = int(map_h)
            pic_w = int(float(map_h) * img_ratio)

        pic_left = int(content_left) + int((float(map_box_w) - pic_w) / 2)
        pic_top = int(content_top) + int((float(map_h) - pic_h) / 2)

        map_stream = io.BytesIO(png_bytes)
        slide.shapes.add_picture(
            map_stream,
            pic_left,
            pic_top,
            width=pic_w,
            height=pic_h
        )

        # =====================================================
        # TABELA À DIREITA: Regiões | %
        # A cor agora fica na MESMA célula do percentual.
        # =====================================================
        if has_table:
            def natural_key(value: str):
                parts = re.split(r"(\d+)", str(value).lower())
                return [
                    int(part) if part.isdigit() else part
                    for part in parts
                ]

            rows_sorted = sorted(
                region_rows,
                key=lambda row: natural_key(row.get("name", ""))
            )

            max_rows = 18
            trimmed = rows_sorted[:max_rows]

            rows = len(trimmed) + 1
            cols = 2

            # Reserva espaço inferior para a tabela Mín / Entre / Máx.
            limits_gap = Inches(0.16)
            limits_h = Inches(1.42)
            max_region_table_h = map_h - limits_gap - limits_h

            if len(trimmed) <= 8:
                desired_row_h = Inches(0.52)
            elif len(trimmed) <= 12:
                desired_row_h = Inches(0.42)
            elif len(trimmed) <= 16:
                desired_row_h = Inches(0.34)
            else:
                desired_row_h = Inches(0.29)

            desired_table_h = desired_row_h * rows
            region_table_h = min(desired_table_h, max_region_table_h)

            tbl_shape = slide.shapes.add_table(
                rows,
                cols,
                table_left,
                content_top,
                table_w,
                region_table_h
            )
            tbl = tbl_shape.table

            tbl.columns[0].width = name_w
            tbl.columns[1].width = pct_w

            row_h = int(region_table_h / rows) if rows else int(region_table_h)
            for row_index in range(rows):
                tbl.rows[row_index].height = row_h

            headers = ["Regiões", "%"]
            for col_index, header in enumerate(headers):
                cell = tbl.cell(0, col_index)
                cell.text = header
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(245, 245, 245)

                paragraph = cell.text_frame.paragraphs[0]
                paragraph.font.bold = True
                paragraph.font.size = Pt(12)
                paragraph.font.name = "DIN"
                paragraph.font.color.rgb = RGBColor(0, 0, 0)
                paragraph.alignment = PP_ALIGN.CENTER
                set_cell_border(cell)

            font_size = 12

            for row_index, row in enumerate(trimmed, start=1):
                name = str(row.get("name", "")).strip()
                value = float(row.get("value", 0.0))
                color = str(row.get("color", "#ffffff")).strip() or "#ffffff"

                # Coluna da região
                cell_name = tbl.cell(row_index, 0)
                cell_name.text = name
                cell_name.fill.solid()
                cell_name.fill.fore_color.rgb = RGBColor(255, 255, 255)

                p_name = cell_name.text_frame.paragraphs[0]
                p_name.font.size = Pt(font_size)
                p_name.font.bold = True
                p_name.font.name = "DIN"
                p_name.font.color.rgb = RGBColor(0, 0, 0)
                p_name.alignment = PP_ALIGN.CENTER
                set_cell_border(cell_name)

                # Coluna do percentual + cor
                cell_pct = tbl.cell(row_index, 1)
                cell_pct.text = f"{_fmt_1(value)}%"
                cell_pct.fill.solid()
                cell_pct.fill.fore_color.rgb = _hex_to_rgb(color)

                p_pct = cell_pct.text_frame.paragraphs[0]
                p_pct.font.size = Pt(font_size)
                p_pct.font.bold = True
                p_pct.font.name = "DIN"
                p_pct.font.color.rgb = text_color_for_fill(color)
                p_pct.alignment = PP_ALIGN.CENTER
                set_cell_border(cell_pct)

            # =================================================
            # TABELA MÍN / ENTRE / MÁX
            # Fica abaixo da tabela de regiões, como na referência.
            # =================================================
            limits_top = content_top + region_table_h + limits_gap

            limits_rows = 4
            limits_cols = 2

            label_w = Inches(0.92)
            value_w = table_w - label_w

            limits_shape = slide.shapes.add_table(
                limits_rows,
                limits_cols,
                table_left,
                limits_top,
                table_w,
                limits_h
            )
            limits_tbl = limits_shape.table

            limits_tbl.columns[0].width = label_w
            limits_tbl.columns[1].width = value_w

            header_h = Inches(0.38)
            data_h = int((limits_h - header_h) / 3)

            limits_tbl.rows[0].height = int(header_h)
            for i in range(1, 4):
                limits_tbl.rows[i].height = data_h

            # Cabeçalho simples, sem inventar período do relatório.
            c00 = limits_tbl.cell(0, 0)
            c00.text = ""
            c00.fill.solid()
            c00.fill.fore_color.rgb = RGBColor(255, 255, 255)
            set_cell_border(c00)

            c01 = limits_tbl.cell(0, 1)
            c01.text = "Limites"
            c01.fill.solid()
            c01.fill.fore_color.rgb = RGBColor(255, 255, 255)

            p_header = c01.text_frame.paragraphs[0]
            p_header.font.name = "DIN"
            p_header.font.size = Pt(12)
            p_header.font.bold = True
            p_header.font.color.rgb = RGBColor(0, 0, 0)
            p_header.alignment = PP_ALIGN.CENTER
            set_cell_border(c01)

            min_v = float(stats["min"])
            max_v = float(stats["max"])

            limits_data = [
                ("Mín.:", f"{_fmt_1(min_v)}%", pal["low"]),
                (
                    "Entre:",
                    f"{_fmt_1(min_v)}% a {_fmt_1(max_v)}%",
                    pal["mid"]
                ),
                ("Máx.:", f"{_fmt_1(max_v)}%", pal["high"]),
            ]

            for idx, (label, value_text, color_hex) in enumerate(limits_data, start=1):
                label_cell = limits_tbl.cell(idx, 0)
                label_cell.text = label
                label_cell.fill.solid()
                label_cell.fill.fore_color.rgb = _hex_to_rgb(color_hex)

                p_label = label_cell.text_frame.paragraphs[0]
                p_label.font.name = "DIN"
                p_label.font.size = Pt(12)
                p_label.font.bold = True
                p_label.font.color.rgb = text_color_for_fill(color_hex)
                p_label.alignment = PP_ALIGN.CENTER
                set_cell_border(label_cell)

                value_cell = limits_tbl.cell(idx, 1)
                value_cell.text = value_text
                value_cell.fill.solid()
                value_cell.fill.fore_color.rgb = RGBColor(255, 255, 255)

                p_value = value_cell.text_frame.paragraphs[0]
                p_value.font.name = "DIN"
                p_value.font.size = Pt(12)
                p_value.font.bold = True
                p_value.font.color.rgb = RGBColor(0, 0, 0)
                p_value.alignment = PP_ALIGN.CENTER
                set_cell_border(value_cell)

        # A antiga legenda horizontal abaixo do mapa foi removida de propósito.

    prs.save(out_path)

# ======================= FIM DO CÓDIGO ORIGINAL (heatmap_page.py) =======================


def ler_casos_excel(conteudo: bytes) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Any]]:
    """parse_excel_cases_with_debug com o arquivo em memória."""
    return parse_excel_cases_with_debug(io.BytesIO(conteudo))


def gerar_pptx(slides: List[Dict[str, Any]], logo_bytes: Optional[bytes]) -> bytes:
    """build_pptx gravando em memória (para download)."""
    saida = io.BytesIO()
    build_pptx(saida, slides, logo_bytes)
    return saida.getvalue()


def base_usada(meta: Dict[str, Any], base_override_text: str) -> Tuple[Optional[float], str]:
    """Mesma regra da QtBridge: base manual (> 0) senão a do Excel."""
    excel_base_total = meta.get("base_total", None)
    manual_base: Optional[float] = None
    txt = (base_override_text or "").strip()
    if txt:
        try:
            manual_base = float(txt.replace(".", "").replace(",", "."))
        except ValueError:
            manual_base = None
    base_total_used = manual_base if (manual_base is not None and manual_base > 0) else (
        float(excel_base_total) if isinstance(excel_base_total, (int, float)) and excel_base_total > 0 else None
    )
    base_source = "manual" if (manual_base is not None and manual_base > 0) else "excel"
    return base_total_used, base_source
