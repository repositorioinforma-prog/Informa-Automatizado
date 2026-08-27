import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

from core.amostra_camadas import distribuir_amostra_em_camadas, normalizar_camadas
from core.amostra_math import calcular_amostra, COLS_RENDA, FAIXAS_ETARIAS
from core.amostra_exportador import exportar_amostra


def _linha(uf, nome_uf, ri, nome_ri, rim, nome_rim, cod_mun, municipio, pop):
    masc = pop // 2
    fem = pop - masc
    row = {
        "uf": uf,
        "nome_uf": nome_uf,
        "cod_regiao_intermediaria": ri,
        "regiao_intermediaria": nome_ri,
        "cod_regiao_imediata": rim,
        "regiao_imediata": nome_rim,
        "cod_municipio": cod_mun,
        "municipio": municipio,
        "populacao_total": pop,
        "masc_total": masc,
        "fem_total": fem,
    }
    # A soma 16+ fecha em 80% da população, distribuída entre as cinco faixas.
    pop_16 = int(pop * 0.8)
    bases = [0.10, 0.22, 0.22, 0.20, 0.26]
    acumulado_m = 0
    acumulado_f = 0
    for i, faixa in enumerate(FAIXAS_ETARIAS):
        total_faixa = int(pop_16 * bases[i]) if i < 4 else pop_16 - (acumulado_m + acumulado_f)
        m = total_faixa // 2
        f = total_faixa - m
        row[f"masc_{faixa}"] = m
        row[f"fem_{faixa}"] = f
        acumulado_m += m
        acumulado_f += f
    renda = [0.45, 0.30, 0.17, 0.08]
    for coluna, frac in zip(COLS_RENDA, renda):
        row[coluna] = int(pop * frac)
    return row


def base_sintetica():
    return pd.DataFrame([
        _linha("AA", "Estado A", "A1", "RI A1", "A11", "RIM A11", "001", "Mun A", 1000),
        _linha("AA", "Estado A", "A1", "RI A1", "A11", "RIM A11", "002", "Mun B", 500),
        _linha("AA", "Estado A", "A2", "RI A2", "A21", "RIM A21", "003", "Mun C", 500),
        _linha("BB", "Estado B", "B1", "RI B1", "B11", "RIM B11", "004", "Mun D", 2000),
    ])


def base_distritos_sintetica():
    linhas = [
        (_linha("AA", "Estado A", "A1", "RI A1", "A11", "RIM A11", "001", "Mun A", 600), "00101", "Distrito A1"),
        (_linha("AA", "Estado A", "A1", "RI A1", "A11", "RIM A11", "001", "Mun A", 400), "00102", "Distrito A2"),
        (_linha("AA", "Estado A", "A2", "RI A2", "A21", "RIM A21", "003", "Mun C", 500), "00301", "Distrito C1"),
    ]
    for linha, codigo, nome in linhas:
        linha["cod_distrito"] = codigo
        linha["distrito"] = nome
    return pd.DataFrame([linha for linha, _, _ in linhas])


class TestAmostraCamadas(unittest.TestCase):
    def test_normalizacao_rejeita_ordem_invertida(self):
        with self.assertRaises(ValueError):
            normalizar_camadas(["municipio", "uf"])

    def test_todos_os_niveis_fecham_no_total(self):
        r = distribuir_amostra_em_camadas(
            base_sintetica(),
            amostra_total=101,
            camadas=["uf", "regiao_intermediaria", "municipio"],
            base_populacional="total",
        )
        self.assertEqual(int(r.resultado["amostra"].sum()), 101)
        for resumo in r.resumos.values():
            self.assertEqual(int(resumo["amostra"].sum()), 101)
        self.assertTrue(all(v["ok"] for v in r.validacoes))

    def test_filhos_fecham_exatamente_com_cada_pai(self):
        r = distribuir_amostra_em_camadas(
            base_sintetica(),
            amostra_total=103,
            camadas=["uf", "regiao_intermediaria", "municipio"],
            base_populacional="total",
        )
        ufs = r.resumos["uf"]
        ris = r.resumos["regiao_intermediaria"]
        por_uf = ris.groupby(["uf", "nome_uf"])["amostra"].sum().reset_index()
        cmp = ufs.merge(por_uf, on=["uf", "nome_uf"], suffixes=("_uf", "_ri"))
        self.assertTrue((cmp["amostra_uf"] == cmp["amostra_ri"]).all())

    def test_pular_nivel_intermediario_e_permitido(self):
        r = distribuir_amostra_em_camadas(
            base_sintetica(),
            amostra_total=99,
            camadas=["regiao_intermediaria", "municipio"],
            base_populacional="total",
        )
        self.assertEqual(int(r.resultado["amostra"].sum()), 99)
        self.assertIn("percentual_no_pai", r.resultado.columns)

    def test_recorte_por_municipios_especificos(self):
        r = calcular_amostra(
            codigos_municipios=["001", "003"],
            amostra_total=31,
            nivel_quebra="municipio",
            base_populacional="total",
            df_base=base_sintetica(),
        )
        self.assertEqual(set(r["cod_municipio"]), {"001", "003"})
        self.assertEqual(int(r["amostra"].sum()), 31)

    def test_recorte_por_distritos_especificos(self):
        r = calcular_amostra(
            codigos_distritos=["00102", "00301"],
            amostra_total=23,
            nivel_quebra="distrito",
            base_populacional="total",
            incluir_distritos=True,
            df_base=base_distritos_sintetica(),
        )
        self.assertEqual(set(r["cod_distrito"]), {"00102", "00301"})
        self.assertEqual(int(r["amostra"].sum()), 23)

    def test_camadas_respeitam_recorte_distrital(self):
        r = distribuir_amostra_em_camadas(
            base_distritos_sintetica(),
            amostra_total=19,
            camadas=["municipio", "distrito"],
            base_populacional="total",
            codigos_distritos=["00102", "00301"],
        )
        self.assertEqual(set(r.resultado["cod_distrito"]), {"00102", "00301"})
        self.assertEqual(int(r.resultado["amostra"].sum()), 19)
        self.assertTrue(all(v["ok"] for v in r.validacoes))

    def test_motor_classico_permanece_fechando(self):
        r = calcular_amostra(
            uf="SP",
            amostra_total=100,
            nivel_quebra="regiao_intermediaria",
            base_populacional="16_mais",
        )
        self.assertEqual(int(r["amostra"].sum()), 100)
        self.assertIn("populacao_base", r.columns)

    def test_exportador_cria_abas_de_controle(self):
        r = distribuir_amostra_em_camadas(
            base_sintetica(),
            amostra_total=50,
            camadas=["uf", "regiao_intermediaria", "municipio"],
            base_populacional="total",
        )
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            exportar_amostra(
                r.resultado,
                caminho_saida=tmp.name,
                resumos_camadas=r.resumos,
            )
            wb = load_workbook(tmp.name, read_only=True)
            self.assertIn("Amostra Completa", wb.sheetnames)
            self.assertIn("Controle - UF", wb.sheetnames)
            self.assertIn("Controle - Reg Interm", wb.sheetnames)
            self.assertIn("Controle - Município", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
