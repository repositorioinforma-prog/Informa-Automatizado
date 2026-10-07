"""Análise Residual (portada do "Relatória Analítica").

A paridade com a ResidualsBridge original foi conferida na migração (SAV com
missing do SPSS, CSV e XLSX: prévias nos 3 tipos de percentual, fila e os 3
arquivos exportados célula a célula); estes testes fixam as regras principais.
"""
import base64
import io
import json
import os
import sys
import types
import unittest

import numpy as np
import openpyxl
import pandas as pd

if "streamlit" not in sys.modules:
    _st = types.ModuleType("streamlit")
    _st.session_state = {}
    _comp = types.ModuleType("streamlit.components")
    _v1 = types.ModuleType("streamlit.components.v1")
    _v1.declare_component = lambda *a, **k: (lambda **kw: None)
    _comp.v1 = _v1
    _st.components = _comp
    sys.modules.update({"streamlit": _st, "streamlit.components": _comp, "streamlit.components.v1": _v1})

from analises.analise_residual import PonteResidual, VERSAO_TELA  # noqa: E402
from core import residuos_math as rm  # noqa: E402


def _csv():
    rng = np.random.default_rng(2)
    n = 600
    reg = rng.integers(1, 4, n)
    voto = np.where(reg == 1, rng.choice([1, 2], n, p=[.7, .3]), rng.choice([1, 2], n, p=[.3, .7]))
    return pd.DataFrame({"voto": voto, "regiao": reg, "Peso": rng.uniform(.5, 1.5, n)}).to_csv(index=False).encode()


def _args(js):
    return json.loads(json.loads(js[js.index("(") + 1:js.rindex(")")]))


class TestAnaliseResidual(unittest.TestCase):
    def setUp(self):
        self.p = PonteResidual()
        self.p.pickData(arquivo={"name": "base.csv", "dados": _csv()})

    def _ultimo(self, nome):
        return _args([e["codigo"] for e in self.p.eventos if e["tipo"] == "js" and e["codigo"].startswith(nome)][-1])

    def test_residuo_ajustado_spss(self):
        obs = np.array([[30.0, 10.0], [20.0, 40.0]])
        z = rm._adjusted_residuals(obs)
        t, r, c = obs.sum(), obs.sum(1), obs.sum(0)
        e = np.outer(r, c) / t
        esperado = (obs - e) / np.sqrt(e * (1 - r[:, None] / t) * (1 - c[None, :] / t))
        np.testing.assert_allclose(z, esperado)
        self.assertEqual(rm._cell_style_for_z(1.85), (rm.FILL_BLUE, rm.FONT_BLUE))
        self.assertEqual(rm._cell_style_for_z(-1.85), (rm.FILL_RED, rm.FONT_RED))
        self.assertIsNone(rm._cell_style_for_z(1.84))

    def test_previa_e_validacoes(self):
        self.p.computePreview(json.dumps({"rowVars": ["voto"], "colVars": ["regiao"], "weightVar": "Peso"}))
        r = self._ultimo("window.__setPreview")
        self.assertTrue(r["ok"])
        bloco = r["tables"][0]["blocks"][0]
        self.assertEqual(r["tables"][0]["title"], "A.R voto")
        self.assertTrue(any(abs(z) >= 1.85 for linha in bloco["z"] for z in linha))
        self.p.computePreview(json.dumps({"rowVars": ["voto"], "colVars": ["regiao"], "weightVar": ""}))
        self.assertEqual(self._ultimo("window.__setPreview")["error"], "Selecione um peso válido.")

    def test_fila_e_exportacao_3_arquivos(self):
        self.p.computePreview(json.dumps({"rowVars": ["voto", "regiao"], "colVars": ["regiao", "voto"],
                                          "weightVar": "Peso"}))
        t1, t2 = self._ultimo("window.__setPreview")["tables"]
        self.p.addToQueueFromPreview(json.dumps(t1))
        self.p.addToQueueFromPreview(json.dumps({**t2, "title": "Sem prefixo"}))
        self.assertEqual([i["title"] for i in self.p.queue], ["A.R voto", "A.R Sem prefixo"])
        self.p.removeQueueItem(1)
        self.p.eventos = []
        self.p.exportXlsx(json.dumps({"mode": "queue", "displayMap": {"voto": "Voto"}, "incluirDebug": True}))
        dls = [e for e in self.p.eventos if e["tipo"] == "download"]
        self.assertEqual([d["nome"] for d in dls],
                         ["analise_residual.xlsx", "analise_residual_DEBUG.xlsx", "analise_residual_CASES_DEBUG.xlsx"])
        ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(dls[0]["b64"]))).active
        self.assertEqual(ws["A1"].value, "A.R voto")
        self.assertEqual(ws["A3"].value, "Voto")

    def test_conferencia_opcional(self):
        self.p.computePreview(json.dumps({"rowVars": ["voto"], "colVars": ["regiao"], "weightVar": "Peso"}))
        t = self._ultimo("window.__setPreview")["tables"]
        self.p.eventos = []
        self.p.exportXlsx(json.dumps({"mode": "batch", "previewTables": t, "incluirDebug": False}))
        self.assertEqual([e["nome"] for e in self.p.eventos if e["tipo"] == "download"], ["analise_residual.xlsx"])

    def test_tabela_repetida_na_fila(self):
        self.p.computePreview(json.dumps({"mode": "queue", "rowVars": ["voto"], "colVars": ["regiao"],
                                          "weightVar": "Peso"}))
        t = self._ultimo("window.__setPreview")["tables"][0]
        self.p.addToQueueFromPreview(json.dumps(t))
        self.p.addToQueueFromPreview(json.dumps({**t, "title": "Outro nome"}))   # repetida: não entra
        self.assertEqual([i["title"] for i in self.p.queue], ["A.R voto"])
        self.assertEqual(self._ultimo("window.__setQueueState")["message"], "Já está na fila (item 1).")
        self.p.addToQueueFromPreview(json.dumps({**t, "title": "Outro nome", "substituir": True}))
        self.assertEqual([i["title"] for i in self.p.queue], ["A.R Outro nome"])
        estado = self._ultimo("window.__setQueueState")
        self.assertEqual(estado["items"][0]["sig"], "voto|regiao|Peso")
        # barra lateral: xlsx da fila e remoção pela assinatura
        self.assertTrue(self.p.xlsx_da_fila().startswith(b"PK"))
        self.p.removerDaFila("voto|regiao|Peso")
        self.assertEqual(self.p.queue, [])

    def test_colunas_na_ordem_dos_codigos(self):
        df = pd.DataFrame({"voto": [1, 2, 1, 2, 1, 2] * 20, "regiao": [3, 3, 2, 1, 3, 2] * 20, "Peso": 1.0})
        p = PonteResidual()
        p.pickData(arquivo={"name": "b.csv", "dados": df.to_csv(index=False).encode()})
        blk = p._compute_block("voto", "regiao", "Peso", "col")
        self.assertEqual(blk["col_cats"], ["1", "2", "3"])
        p.ordenar_colunas_por_codigo = False   # como no desktop: ordem em que aparecem no banco
        blk = p._compute_block("voto", "regiao", "Peso", "col")
        self.assertEqual(blk["col_cats"], ["3", "2", "1"])

    def test_versao_da_tela(self):
        pasta = os.path.join(os.path.dirname(os.path.dirname(__file__)), "componentes", "relatoria_residual")
        html = open(os.path.join(pasta, "index.html"), encoding="utf-8").read()
        self.assertIn(f'<meta name="versao-tela" content="{VERSAO_TELA}" />', html)
        self.assertEqual(open(os.path.join(pasta, "ponte_qt.js"), "rb").read(),
                         open(os.path.join(os.path.dirname(pasta), "ponte_qt.js"), "rb").read())


if __name__ == "__main__":
    unittest.main()
