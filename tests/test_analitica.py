"""Ferramentas analíticas portadas do "Relatória Analítica" (liberadas uma
por vez — por enquanto, Cruzamentos).

A paridade com o código original foi conferida na migração (mesmas
entradas -> mesmos resultados e arquivos idênticos célula a célula);
estes testes fixam as regras principais para não regredirem.
"""
import io
import unittest

import numpy as np
import openpyxl
import pandas as pd

from core import cruzamentos_math as cz
from core.analitica_dados import carregar_banco


def _df(n=300, seed=1):
    rng = np.random.default_rng(seed)
    df = pd.DataFrame({
        "voto": rng.choice(["A", "B", "Não sabe", "Branco/Nulo", "C"], n),
        "sexo": rng.choice(["M", "F"], n),
        "x": rng.normal(10, 2, n),
        "Peso": rng.uniform(0.5, 1.5, n),
    })
    df["y"] = df["x"] * 2 + rng.normal(0, 1, n)
    return df


class TestCarregamento(unittest.TestCase):
    def test_csv(self):
        b = carregar_banco(_df().to_csv(index=False).encode(), "banco.csv")
        self.assertEqual(b.nome_arquivo, "banco.csv")
        self.assertIn("Peso", b.df.columns)

    def test_formato_invalido(self):
        with self.assertRaises(ValueError):
            carregar_banco(b"x", "banco.txt")


class TestCruzamentos(unittest.TestCase):
    def test_ordem_e_percentuais(self):
        t = cz.calcular_cruzamentos(_df(), {}, {}, "voto", ["sexo"], "Peso", False, "col")[0]
        self.assertEqual(t["row_labels"][-2:], ["Branco/Nulo", "Não sabe"])  # fixos no fim
        totais = t["pct_total"][:3]
        self.assertEqual(totais, sorted(totais, reverse=True))  # demais por Total desc
        for j in range(len(t["col_labels"])):
            self.assertAlmostEqual(sum(r[j] for r in t["pct_matrix"]), 100.0, delta=0.3)

    def test_peso_sem_nome_peso_e_ignorado(self):
        df = _df().rename(columns={"Peso": "fator"})
        t = cz.calcular_cruzamentos(df, {}, {}, "voto", ["sexo"], "fator", False, "col")[0]
        self.assertIn("Sem ponderação", t["base_display"])

    def test_reordenar_e_exportar(self):
        t = cz.calcular_cruzamentos(_df(), {}, {}, "voto", ["sexo"], "Peso", False, "col")[0]
        cz.reordenar_linhas(t, list(reversed(t["row_labels"])))
        cz.reordenar_colunas(t, [1])
        self.assertEqual(len(t["col_labels"]), 1)
        ws = openpyxl.load_workbook(io.BytesIO(cz.exportar_excel([t]))).active
        self.assertEqual(ws["A3"].value, "Categoria")
        self.assertEqual(ws["B3"].value, "Total")
        self.assertEqual(ws.cell(4, 1).value, t["row_labels"][0])


if __name__ == "__main__":
    unittest.main()
