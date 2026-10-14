import openpyxl

from core.avisos_consistencia_math import gerar_avisos_consistencia, textos_avisos


def _tabela(ws, linha, titulo, segmentos, linhas, base, pergunta="Pergunta: x?"):
    """Escreve uma tabela no formato do relatório e devolve a próxima linha livre."""
    ncol = 2 + len(segmentos)
    ws.cell(linha, 1, titulo)
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=max(ncol, 2))
    ws.cell(linha + 1, 2, "Total")
    for i, s in enumerate(segmentos):
        ws.cell(linha + 1, 3 + i, s)
    r = linha + 2
    for rot, vals in linhas:
        ws.cell(r, 1, rot)
        for j, v in enumerate(vals):
            ws.cell(r, 2 + j, v)
        r += 1
    ws.cell(r, 1, "Base")
    for j, v in enumerate(base):
        ws.cell(r, 2 + j, v)
    ws.cell(r + 1, 1, pergunta)
    return r + 3


def _ws():
    return openpyxl.Workbook().active


def test_segmento_maior_que_total_agrupa_por_segmento():
    ws = _ws()
    r = _tabela(ws, 1, "Avaliação A", ["Facebook", "Mulher"], [("Bom", [50, 50, 50])], [301, 353, 150])
    _tabela(ws, r, "Avaliação B", ["Facebook", "Mulher"], [("Bom", [50, 50, 50])], [301, "353", "25*"])
    av = gerar_avisos_consistencia(ws)
    assert len(av["segmento_maior_que_total"]) == 2
    assert all(a["segmento"] == "Facebook" for a in av["segmento_maior_que_total"])
    msgs = textos_avisos(av)
    assert len(msgs) == 1 and "em 2 tabelas" in msgs[0]


def test_base_pequena_com_asterisco_nao_quebra():
    ws = _ws()
    _tabela(ws, 1, "T", ["Jovem"], [("Sim", [10, 10])], [604, "25*"])
    assert gerar_avisos_consistencia(ws)["segmento_maior_que_total"] == []


def test_mesmo_titulo_com_bases_diferentes():
    ws = _ws()
    r = _tabela(ws, 1, "Avaliação do Prefeito (Estimulada e Única)", [], [("Bom", [50])], [604])
    r = _tabela(ws, r, "Avaliação do Prefeito (Estimulada e Única)", ["Homem"], [("Bom", [50, 50])], [580, 300])
    _tabela(ws, r, "Outra pergunta", [], [("Sim", [50])], [604])
    tit = gerar_avisos_consistencia(ws)["titulos_bases_diferentes"]
    assert len(tit) == 1
    assert [v for _, v in tit[0]["ocorrencias"]] == [604, 580]


def test_mesmo_titulo_mesma_base_nao_avisa():
    ws = _ws()
    r = _tabela(ws, 1, "P", [], [("Bom", [50])], [604])
    _tabela(ws, r, "P", ["Homem"], [("Bom", [50, 50])], [604, 300])
    assert gerar_avisos_consistencia(ws)["titulos_bases_diferentes"] == []


def test_candidato_some_entre_cenarios_e_grafia_diferente():
    ws = _ws()
    linhas1 = [("Ana (PT)", [10]), ("Alexda Confecção (AVANTE)", [5]), ("Bruno (PL)", [5]),
               ("Outros", [3]), ("Não sabe", [70])]
    linhas2 = [("Ana (PT)", [12]), ("Alex da Confecção (AVANTE)", [6]), ("Não irá votar", [2]),
               ("Não sabe", [80])]
    r = _tabela(ws, 1, "Intenção de voto para Deputado - Cenário 1 (Estimulada e Única)", [], linhas1, [604])
    _tabela(ws, r, "Intenção de voto para Deputado - Cenário 2 (Estimulada e Única)", [], linhas2, [604])
    cen = gerar_avisos_consistencia(ws)["candidatos_cenarios"]
    assert len(cen) == 1
    assert [f["candidato"] for f in cen[0]["faltas"]] == ["Bruno (PL)"]
    assert len(cen[0]["grafias"]) == 1
    msgs = "\n".join(textos_avisos(gerar_avisos_consistencia(ws)))
    assert "Só no Cenário 1: Bruno (PL)" in msgs
    assert "escrito diferente" in msgs
    assert "Não irá votar" not in msgs


def test_relatorio_limpo_sem_avisos():
    ws = _ws()
    r = _tabela(ws, 1, "Voto - Cenário 1", ["Homem"], [("Ana", [10, 10]), ("Bruno", [5, 5])], [604, 300])
    _tabela(ws, r, "Voto - Cenário 2", ["Homem"], [("Ana", [12, 12]), ("Bruno", [6, 6])], [604, 300])
    assert textos_avisos(gerar_avisos_consistencia(ws)) == []
