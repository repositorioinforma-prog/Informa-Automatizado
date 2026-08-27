from pathlib import Path
import sqlite3
import tempfile
import unittest
from contextlib import closing

import pandas as pd
from openpyxl import load_workbook

from core.banco_amostras import (
    CAMINHO_BANCO_PADRAO,
    buscar_territorios,
    carregar_base_amostragem,
    carregar_base_motor_amostragem,
    consultar_cobertura,
    verificar_integridade,
)
from core.amostra_math import calcular_amostra
from core.amostra_exportador import exportar_amostra


class BancoAmostrasTests(unittest.TestCase):
    def test_banco_existe_e_esta_integro(self):
        self.assertTrue(CAMINHO_BANCO_PADRAO.exists())
        resultado = verificar_integridade()
        self.assertEqual(resultado["integridade"], "ok")
        self.assertEqual(resultado["chaves_estrangeiras_invalidas"], 0)

    def test_quantidades_demograficas(self):
        with closing(sqlite3.connect(CAMINHO_BANCO_PADRAO)) as conexao:
            contagens = dict(
                conexao.execute(
                    """SELECT nivel, COUNT(*) FROM territorios
                       WHERE nivel IN ('municipio','distrito','subdistrito','bairro')
                       GROUP BY nivel"""
                ).fetchall()
            )
        self.assertEqual(contagens["municipio"], 5570)
        self.assertEqual(contagens["distrito"], 10698)
        self.assertEqual(contagens["subdistrito"], 643)
        self.assertEqual(contagens["bairro"], 17576)

    def test_municipio_tem_populacao_16_mais_municipal(self):
        base = carregar_base_amostragem(
            "municipio", chaves_territoriais=["municipio:3550308"]
        )
        self.assertEqual(len(base), 1)
        self.assertEqual(base.iloc[0]["nome"], "São Paulo")
        self.assertGreater(int(base.iloc[0]["populacao_16_mais_municipal"]), 0)
        self.assertIn("municipal_sidra", base.iloc[0]["metodo_populacao_16_mais_referencia"])

    def test_codigo_de_bairro_duplicado_foi_preservado(self):
        bairros = buscar_territorios(
            "bairro", texto="Sao To", codigo_municipio="3205200", limite=20
        )
        repetidos = bairros[bairros["codigo"] == "3205200069"]
        self.assertEqual(len(repetidos), 2)
        self.assertEqual(set(repetidos["codigo_duplicado"]), {1})
        self.assertEqual(repetidos["chave_territorial"].nunique(), 2)

    def test_distritos_fecham_com_municipios(self):
        with closing(sqlite3.connect(CAMINHO_BANCO_PADRAO)) as conexao:
            divergencias = conexao.execute(
                """SELECT COUNT(*) FROM cobertura_hierarquia
                   WHERE nivel_filho = 'distrito' AND fecha_exatamente <> 1"""
            ).fetchone()[0]
        self.assertEqual(divergencias, 0)

    def test_cobertura_de_bairros_e_explicita(self):
        cobertura = consultar_cobertura(
            nivel_pai="municipio", codigo_pai="3304557", nivel_filho="bairro"
        )
        self.assertEqual(len(cobertura), 1)
        self.assertEqual(int(cobertura.iloc[0]["fecha_exatamente"]), 1)
        self.assertAlmostEqual(float(cobertura.iloc[0]["proporcao_cobertura"]), 1.0)

    def test_base_sqlite_e_convertida_para_o_motor(self):
        base = carregar_base_motor_amostragem("municipio")
        self.assertEqual(len(base), 5570)
        for coluna in [
            "cod_municipio",
            "municipio",
            "masc_16_19",
            "fem_50mais",
            "populacao_16_mais_referencia",
            "renda_disponivel",
        ]:
            self.assertIn(coluna, base.columns)

    def test_motor_calcula_usando_base_sqlite_injetada(self):
        base = carregar_base_motor_amostragem("municipio")
        resultado = calcular_amostra(
            uf="SP",
            amostra_total=1000,
            nivel_quebra="regiao_intermediaria",
            base_populacional="16_mais",
            df_base=base,
        )
        self.assertEqual(int(resultado["amostra"].sum()), 1000)
        self.assertGreater(int(resultado["populacao_base"].sum()), 0)

    def test_renda_ausente_fica_em_branco_e_nao_em_zero(self):
        base = carregar_base_motor_amostragem("municipio")
        resultado = calcular_amostra(
            codigos_municipios=["1504752"],
            amostra_total=20,
            nivel_quebra="municipio",
            base_populacional="16_mais",
            df_base=base,
        )
        self.assertEqual(int(resultado.iloc[0]["renda_disponivel"]), 0)
        self.assertTrue(pd.isna(resultado.iloc[0]["cota_renda_ate_2sm"]))

        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            exportar_amostra(
                resultado,
                caminho_saida=tmp.name,
                exibicao_cotas="totais",
            )
            wb = load_workbook(tmp.name, data_only=False)
            ws = wb["Amostra Completa"]
            cabecalhos = {
                ws.cell(row=2, column=coluna).value: coluna
                for coluna in range(1, ws.max_column + 1)
                if ws.cell(row=2, column=coluna).value
            }
            coluna_renda = cabecalhos["Até 2sm"]
            self.assertIsNone(ws.cell(row=3, column=coluna_renda).value)


if __name__ == "__main__":
    unittest.main()
