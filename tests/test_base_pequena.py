"""Código 08 — Base pequena: marca com '*' as bases abaixo de 30 (maiores que zero)."""
import unittest

from openpyxl import Workbook

from core.relatorio_automatizado_math import aplicar_codigo_08


class TestBasePequena(unittest.TestCase):
    def test_abaixo_de_30_inclusive_menores_que_20(self):
        ws = Workbook().active
        ws.append(["Candidato A", 50, 40, 30, 20, 10])
        ws.append(["Base", 450, 35, 29, 18, 5])
        ws.append(["Base", 0, 30, 31, 100, 0])  # zero e 30+ não marcam
        self.assertEqual(aplicar_codigo_08(ws), 1)
        self.assertEqual([c.value for c in ws[2]], ["Base", 450, 35, "29*", "18*", "5*"])
        self.assertEqual(ws.cell(row=4, column=1).value, "*Base pequena para análise estatística")
        self.assertEqual([c.value for c in ws[3]], ["Base", 0, 30, 31, 100, 0])


if __name__ == "__main__":
    unittest.main()
