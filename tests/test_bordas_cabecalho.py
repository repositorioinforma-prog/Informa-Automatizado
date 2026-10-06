"""Código 17 — bordas espessas no contorno do cabeçalho das tabelas segmentadas."""
import io
import unittest

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from core.relatorio_automatizado_math import aplicar_codigo_02, aplicar_codigo_17

FINA = Side(style="thin")
MEDIA = Side(style="medium")


def _tabela_marica(ws, topo=11):
    """Réplica da tabela 'Participação na Expo Maricá' (cabeçalho em 2 linhas)."""
    t, b = topo, topo + 1
    ws.cell(row=t - 1, column=1, value="Participação na Expo Maricá (Estimulada e Única)")
    ws.merge_cells(start_row=t - 1, start_column=1, end_row=t - 1, end_column=8)
    ws.cell(row=t, column=2, value="Total")
    ws.cell(row=t, column=3, value="Tempo de moradia em Maricá")
    subcats = ["Mora em Maricá desde que nasceu", "Mora há mais de 20 anos", "Mora entre 20 e 10 anos",
               "Mora entre 10 e 5 anos", "Mora há menos de 5 anos"]
    for i, s in enumerate(subcats):
        ws.cell(row=b, column=3 + i, value=s)
    dados = [("Não participou", 62.1, 49.8, 60.7, 65.9, 71.2, 83.1),
             ("Participou", 37.9, 50.2, 39.3, 34.1, 28.8, 16.9),
             ("Base reduzida", 464, 137, 124, 90, 74, 39)]
    for i, linha in enumerate(dados):
        for j, v in enumerate(linha):
            ws.cell(row=b + 1 + i, column=1 + j, value=v)
    # bordas finas/médias como vêm do SPSS (depois do código 02)
    for r in range(t, b + 4):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
    for c in range(1, 9):
        ws.cell(row=t, column=c).border = Border(left=FINA, right=FINA, top=MEDIA, bottom=FINA)
    ws.merge_cells(start_row=t, start_column=1, end_row=b, end_column=1)
    ws.merge_cells(start_row=t, start_column=2, end_row=b, end_column=2)
    ws.merge_cells(start_row=t, start_column=3, end_row=t, end_column=8)
    ws.cell(row=b + 4, column=1, value="Pergunta: Você já participou da Expo Maricá?")
    return t, b


def _recarregar(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active


def _estilo(ws, r, c, lado):
    s = getattr(ws.cell(row=r, column=c).border, lado)
    return s.style if s is not None else None


class TestCodigo17(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        ws = self.wb.active
        self.t, self.b = _tabela_marica(ws)
        self.n = aplicar_codigo_17(ws)
        # duas recargas: o assistente reabre o arquivo a cada etapa seguinte
        self.ws = _recarregar(_recarregar(self.wb).parent)

    def test_um_cabecalho(self):
        self.assertEqual(self.n, 1)

    def test_linha_de_cima_e_de_baixo_espessas_de_a_ate_h(self):
        for c in range(1, 9):
            self.assertEqual(_estilo(self.ws, self.t, c, "top"), "thick", f"topo col {c}")
            self.assertEqual(_estilo(self.ws, self.b, c, "bottom"), "thick", f"base col {c}")
            # primeira linha de dados também (borda compartilhada)
            self.assertEqual(_estilo(self.ws, self.b + 1, c, "top"), "thick", f"dados col {c}")

    def test_laterais_e_coluna_total(self):
        for r in (self.t, self.b):
            self.assertEqual(_estilo(self.ws, r, 1, "left"), "thick")
            self.assertEqual(_estilo(self.ws, r, 8, "right"), "thick")
            self.assertEqual(_estilo(self.ws, r, 1, "right"), "thick")   # A | Total
            self.assertEqual(_estilo(self.ws, r, 2, "left"), "thick")
            self.assertEqual(_estilo(self.ws, r, 2, "right"), "thick")   # Total | segmento
            self.assertEqual(_estilo(self.ws, r, 3, "left"), "thick")

    def test_divisorias_internas_intocadas(self):
        # entre subcategorias (linha das subcategorias) e linha sob o grupo
        for c in range(3, 8):
            self.assertEqual(_estilo(self.ws, self.b, c, "right"), "thin", f"col {c}")
        self.assertEqual(_estilo(self.ws, self.t, 3, "bottom"), "thin")

    def test_dados_intocados(self):
        self.assertEqual(_estilo(self.ws, self.b + 2, 4, "left"), "thin")
        self.assertEqual(_estilo(self.ws, self.b + 3, 4, "bottom"), "thin")

    def test_titulo_intocado(self):
        self.assertIsNone(_estilo(self.ws, self.t - 1, 4, "top"))


class TestVariasTabelas(unittest.TestCase):
    def test_duas_tabelas_e_tabela_sem_total_mesclado(self):
        wb = Workbook()
        ws = wb.active
        _tabela_marica(ws, topo=3)
        _tabela_marica(ws, topo=20)
        ws["B40"] = "Total"   # não mesclado -> ignora
        ws["C40"] = 10
        self.assertEqual(aplicar_codigo_17(ws), 2)
        ws2 = _recarregar(wb)
        self.assertEqual(_estilo(ws2, 20, 5, "top"), "thick")
        self.assertIsNone(_estilo(ws2, 40, 2, "top"))



class TestSequenciaDoAssistente(unittest.TestCase):
    """Etapa do assistente: 17 e logo depois o Muda Bordas (02)."""

    def test_contorno_fica_medio_e_sobrevive_a_recargas(self):
        wb = Workbook()
        ws = wb.active
        t, b = _tabela_marica(ws)
        aplicar_codigo_17(ws)
        ws = _recarregar(wb)
        aplicar_codigo_02(ws)
        ws = _recarregar(_recarregar(ws.parent).parent)
        for c in range(1, 9):
            self.assertEqual(_estilo(ws, t, c, "top"), "medium", f"topo col {c}")
            self.assertEqual(_estilo(ws, b, c, "bottom"), "medium", f"base col {c}")
        for r in (t, b):
            self.assertEqual(_estilo(ws, r, 1, "left"), "medium")
            self.assertEqual(_estilo(ws, r, 8, "right"), "medium")
            self.assertEqual(_estilo(ws, r, 2, "left"), "medium")
            self.assertEqual(_estilo(ws, r, 2, "right"), "medium")
        self.assertEqual(_estilo(ws, b, 4, "right"), "thin")  # divisória interna
        # nenhuma borda espessa sobrando na planilha
        for linha in ws.iter_rows():
            for cel in linha:
                for lado in ("left", "right", "top", "bottom"):
                    self.assertNotEqual(_estilo(ws, cel.row, cel.column, lado), "thick")


if __name__ == "__main__":
    unittest.main()
