import tempfile
import unittest

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.amostra_alocacao import aplicar_alocacao_manual
from core.amostra_camadas import distribuir_amostra_em_camadas
from core.amostra_exportador import exportar_amostra
from core.amostra_math import calcular_amostra
from tests.test_amostra_camadas import base_sintetica


def _linha_com_valor(ws, coluna, valor):
    for linha in range(1, ws.max_row + 1):
        if ws.cell(row=linha, column=coluna).value == valor:
            return linha
    raise AssertionError(f"Valor {valor!r} não encontrado na coluna {coluna}.")


def _cabecalhos_visiveis_linha_2(ws):
    retorno = []
    for coluna in range(1, ws.max_column + 1):
        if ws.column_dimensions[get_column_letter(coluna)].hidden:
            continue
        valor = ws.cell(row=2, column=coluna).value
        if valor is not None:
            retorno.append(valor)
    return retorno


class TestExportadorAuditoria(unittest.TestCase):
    def setUp(self):
        proporcional = calcular_amostra(
            codigos_municipios=["001", "002", "003"],
            amostra_total=20,
            nivel_quebra="municipio",
            base_populacional="total",
            df_base=base_sintetica(),
        )
        self.resultado = aplicar_alocacao_manual(
            proporcional,
            nivel="municipio",
            alocacoes_manuais={"001": 10, "002": 0, "003": 10},
            amostra_total=20,
        )
        self.info = {
            "nome_projeto": "Auditoria de teste",
            "amostra_total": 20,
            "metodo_distribuicao": "manual",
            "fonte_dados_label": "Base sintética",
            "modo": "classico",
            "base_populacional": "total",
            "nivel_quebra_label": "Município",
            "nivel_saida_solicitado_label": "Região Intermediária",
            "nivel_alocacao_manual_label": "Município",
            "recorte_descricao": "Municípios 001, 002 e 003",
        }

    def _exportar(self, formato):
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.close()
        campanhas = {
            "Campanha Norte": self.resultado.iloc[:2].copy(),
            "Campanha Sul": self.resultado.iloc[2:].copy(),
        }
        exportar_amostra(
            self.resultado,
            dict_campanhas=campanhas,
            caminho_saida=tmp.name,
            informacoes=self.info,
            exibicao_cotas=formato,
        )
        return tmp.name

    def test_informacoes_e_a_primeira_aba_e_contem_memoria_de_calculo(self):
        caminho = self._exportar("ambos")
        wb = load_workbook(caminho, data_only=False)
        self.assertEqual(wb.sheetnames[:2], ["Informações", "Amostra Completa"])
        self.assertIn("Campanha Norte", wb.sheetnames)
        self.assertIn("Campanha Sul", wb.sheetnames)

        ws = wb["Informações"]
        linha_nivel_inicial = _linha_com_valor(ws, 1, "Nível inicialmente configurado")
        linha_unidade_metas = _linha_com_valor(ws, 1, "Unidade das metas manuais")
        self.assertEqual(ws.cell(row=linha_nivel_inicial, column=3).value, "Região Intermediária")
        self.assertEqual(ws.cell(row=linha_unidade_metas, column=3).value, "Município")
        linha_header = _linha_com_valor(ws, 6, "Código da unidade")
        headers = [ws.cell(row=linha_header, column=col).value for col in range(1, 20)]
        for esperado in [
            "População total",
            "População-base",
            "% real do recorte",
            "% real no pai",
            "Amostra proporcional do pai",
            "Amostra ideal",
            "Sugestão proporcional",
            "Amostra aplicada",
            "% da amostra",
            "Diferença vs proporcional",
        ]:
            self.assertIn(esperado, headers)

        primeira_linha = linha_header + 1
        self.assertTrue(str(ws.cell(row=primeira_linha, column=9).value).startswith("=IFERROR("))
        self.assertTrue(str(ws.cell(row=primeira_linha, column=10).value).startswith("=I"))
        self.assertTrue(str(ws.cell(row=primeira_linha, column=11).value).startswith("=$C$"))
        self.assertTrue(str(ws.cell(row=primeira_linha, column=12).value).startswith("=J"))
        self.assertTrue(str(ws.cell(row=primeira_linha, column=15).value).startswith("=IFERROR("))
        self.assertEqual(ws.cell(row=primeira_linha, column=14).value, 10)
        self.assertEqual(ws.cell(row=primeira_linha, column=13).value, 10)

        linha_total = _linha_com_valor(ws, 1, "TOTAL")
        self.assertTrue(str(ws.cell(row=linha_total, column=12).value).startswith("=SUM("))
        self.assertTrue(str(ws.cell(row=linha_total, column=14).value).startswith("=SUM("))

    def test_informacoes_por_camadas_reconstroi_ideal_dentro_do_pai(self):
        calculo = distribuir_amostra_em_camadas(
            base_sintetica(),
            amostra_total=101,
            camadas=["uf", "regiao_intermediaria", "municipio"],
            base_populacional="total",
        )
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            exportar_amostra(
                calculo.resultado,
                caminho_saida=tmp.name,
                resumos_camadas=calculo.resumos,
                informacoes={
                    "nome_projeto": "Camadas",
                    "amostra_total": 101,
                    "metodo_distribuicao": "proporcional",
                    "modo": "Por camadas (Beta)",
                    "base_populacional": "total",
                    "nivel_quebra_label": "Município",
                    "camadas": ["uf", "regiao_intermediaria", "municipio"],
                },
                exibicao_cotas="ambos",
            )
            wb = load_workbook(tmp.name, data_only=False)
            ws = wb["Informações"]
            linha_header = _linha_com_valor(ws, 6, "Código da unidade")
            linha = linha_header + 1
            formula_pct_pai = str(ws.cell(row=linha, column=10).value)
            formula_amostra_pai = str(ws.cell(row=linha, column=11).value)
            formula_ideal = str(ws.cell(row=linha, column=12).value)
            self.assertIn("SUMIFS($H$", formula_pct_pai)
            self.assertIn("$A$", formula_pct_pai)
            self.assertIn("$B$", formula_pct_pai)
            self.assertIn("SUMIFS($M$", formula_amostra_pai)
            self.assertEqual(formula_ideal, f"=J{linha}*K{linha}")

    def test_totais_exibe_contagens_e_mantem_total_e_percentual_da_amostra(self):
        caminho = self._exportar("totais")
        wb = load_workbook(caminho, data_only=False)
        ws = wb["Amostra Completa"]
        headers = _cabecalhos_visiveis_linha_2(ws)
        self.assertIn("Masculino", headers)
        self.assertNotIn("Masculino (%)", headers)
        self.assertIn("Total", headers)
        self.assertIn("% da amostra", headers)

        coluna_masc = headers.index("Masculino") + 1
        # Há colunas mescladas antes de Sexo; localizamos pela linha real para evitar
        # depender de deslocamentos fixos.
        coluna_masc = next(
            col for col in range(1, ws.max_column + 1)
            if ws.cell(row=2, column=col).value == "Masculino"
        )
        self.assertEqual(ws.cell(row=3, column=coluna_masc).value, 5)

    def test_percentuais_exibe_percentuais_e_esconde_contagens_auxiliares(self):
        caminho = self._exportar("percentuais")
        wb = load_workbook(caminho, data_only=False)
        ws = wb["Amostra Completa"]
        coluna_masc = next(
            col for col in range(1, ws.max_column + 1)
            if ws.cell(row=2, column=col).value == "Masculino"
        )
        valor = ws.cell(row=3, column=coluna_masc).value
        self.assertTrue(str(valor).startswith("=IF("))
        self.assertEqual(ws.cell(row=3, column=coluna_masc).number_format, "0.00%")

        ocultas = [
            col for col in range(1, ws.max_column + 1)
            if ws.column_dimensions[get_column_letter(col)].hidden
        ]
        self.assertGreater(len(ocultas), 0)
        self.assertIn("% da amostra", _cabecalhos_visiveis_linha_2(ws))

    def test_ambos_exibe_contagens_e_percentuais_lado_a_lado(self):
        caminho = self._exportar("ambos")
        wb = load_workbook(caminho, data_only=False)
        ws = wb["Amostra Completa"]
        headers = _cabecalhos_visiveis_linha_2(ws)
        self.assertIn("Masculino (n)", headers)
        self.assertIn("Masculino (%)", headers)
        self.assertIn("% da amostra", headers)

        col_n = next(
            col for col in range(1, ws.max_column + 1)
            if ws.cell(row=2, column=col).value == "Masculino (n)"
        )
        col_pct = next(
            col for col in range(1, ws.max_column + 1)
            if ws.cell(row=2, column=col).value == "Masculino (%)"
        )
        self.assertEqual(ws.cell(row=3, column=col_n).value, 5)
        self.assertTrue(str(ws.cell(row=3, column=col_pct).value).startswith("=IF("))


if __name__ == "__main__":
    unittest.main()
