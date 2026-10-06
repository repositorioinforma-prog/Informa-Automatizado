"""Mapa de Calor (portado do "Relatória Analítica").

A paridade com a QtBridge original foi conferida na migração (mesmas
chamadas para a tela, mesmos avisos e PPTX idêntico); estes testes fixam as
regras principais para não regredirem.
"""
import base64
import io
import json
import math
import os
import sys
import types
import unittest

import openpyxl
from PIL import Image
from pptx import Presentation

# streamlit mínimo (os testes não abrem a interface)
if "streamlit" not in sys.modules:
    _st = types.ModuleType("streamlit")
    _st.session_state = {}
    _comp = types.ModuleType("streamlit.components")
    _v1 = types.ModuleType("streamlit.components.v1")
    _v1.declare_component = lambda *a, **k: (lambda **kw: None)
    _comp.v1 = _v1
    _st.components = _comp
    sys.modules.update({"streamlit": _st, "streamlit.components": _comp, "streamlit.components.v1": _v1})

from analises.mapa_calor import PonteMapaCalor, VERSAO_TELA  # noqa: E402
from core import mapa_calor_math as mk  # noqa: E402


def _relatorio_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A2"] = "Avaliação do Governo"
    ws["A4"] = "Bairros"
    for j, h in enumerate(["", "Total", "Centro", "Sul", "Norte"], 1):
        ws.cell(5, j, h or None)
    for r, (seg, vals) in enumerate({"Aprovação": [50, 61.25, 40, 52],
                                      "Reprovação": [40, 30, 50, 41]}.items(), 6):
        ws.cell(r, 1, seg)
        for j, v in enumerate(vals, 2):
            ws.cell(r, j, v)
    ws.cell(8, 1, "Base")
    for j, v in enumerate([900, 300, 300, 300], 2):
        ws.cell(8, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _args(js):
    """window.__x(a, b, ...) -> [a, b, ...]"""
    return json.loads("[" + js[js.index("(") + 1:js.rindex(")")] + "]")


class TestMapaCalor(unittest.TestCase):
    def setUp(self):
        self.p = PonteMapaCalor()
        self.p.pickExcel(arquivo={"name": "Relatorio.xlsx", "dados": _relatorio_xlsx()})
        self.caso = list(self.p.cases_by_key)[0]

    def _js(self, nome):
        return [e["codigo"] for e in self.p.eventos if e["tipo"] == "js" and e["codigo"].startswith(nome)]

    def test_leitura_do_relatorio(self):
        payload = json.loads(_args(self._js("window.__setCases")[-1])[0])
        self.assertEqual(len(payload["cases"]), 1)
        self.assertIn("Bairros", payload["cases"][0])
        self.assertEqual(payload["baseByCase"][self.caso], "900")
        self.assertEqual(sorted(payload["segmentsByCase"][self.caso]), ["Aprovação", "Reprovação"])

    def test_estatisticas(self):
        self.p.computeStats(self.caso, "Aprovação", "")
        a = _args(self._js("window.__setStats")[-1])
        media = (61.25 + 40 + 52) / 3  # o Total não entra
        erro = 1.96 * math.sqrt(media * (100 - media) / 900)
        self.assertEqual(a[1], "51,1")
        self.assertEqual(a[7], mk.round_half_up(media, 1))
        self.assertEqual(a[6], mk.round_half_up(media - erro, 1))
        self.p.computeStats(self.caso, "Aprovação", "abc")  # base manual inválida -> usa a do Excel
        self.assertEqual(_args(self._js("window.__setStats")[-1])[6], a[6])

    def test_pintura_automatica(self):
        self.p.applyAutoPaint(self.caso, "Aprovação", "", True)
        payload = json.loads(_args(self._js("window.__applyAutoPaint")[-1])[0])
        cls = {i["name"]: i["cls"] for i in payload["items"]}
        self.assertEqual(cls, {"Centro": "high", "Sul": "low", "Norte": "mid"})
        self.assertTrue(payload["noOverwriteManual"])

    def test_fila_pptx_e_png(self):
        png = io.BytesIO()
        Image.new("RGB", (400, 300), (10, 100, 200)).save(png, "PNG")
        url = "data:image/png;base64," + base64.b64encode(png.getvalue()).decode()
        meta = {"caseKey": self.caso, "segment": "Aprovação",
                "palette": {"low": "#87cefa", "mid": "#1e90ff", "high": "#121958"},
                "stats": {"min": 47.3, "mean": 51.1, "max": 54.8},
                "legend": [{"name": "Centro", "value": 61.3, "color": "#121958"}]}
        self.p.addToPptQueue(json.dumps(meta), url)
        self.p.addToPptQueue(json.dumps({"caseKey": ""}), url)  # meta incompleta: não entra
        self.assertEqual(len(self.p.ppt_queue), 1)
        self.p.exportPptx()
        dl = [e for e in self.p.eventos if e["tipo"] == "download"][-1]
        self.assertEqual(dl["nome"], "mapas_calor.pptx")
        prs = Presentation(io.BytesIO(base64.b64decode(dl["b64"])))
        self.assertEqual(len(prs.slides), 1)
        self.p.savePng(url)
        self.assertEqual(self.p.ultimo_download["nome"], "mapa_calor.png")
        self.p.clearPptQueue()
        self.p.eventos = []
        self.p.exportPptx()  # fila vazia: só avisa
        self.assertEqual([e["tipo"] for e in self.p.eventos], ["aviso"])

    def test_fila_repetido_substituir_e_estado(self):
        png = io.BytesIO()
        Image.new("RGB", (400, 300), (10, 100, 200)).save(png, "PNG")
        url = "data:image/png;base64," + base64.b64encode(png.getvalue()).decode()
        meta = {"caseKey": self.caso, "segment": "Aprovação",
                "palette": {"low": "#87cefa", "mid": "#1e90ff", "high": "#121958"},
                "stats": {"min": 47.3, "mean": 51.1, "max": 54.8}, "legend": []}
        self.p.addToPptQueue(json.dumps(meta), url)
        self.p.addToPptQueue(json.dumps({**meta, "stats": {"min": 1, "mean": 2, "max": 3}}), url)  # repetido
        self.assertEqual(self.p.ppt_queue[0]["stats"]["min"], 47.3)
        self.p.addToPptQueue(json.dumps({**meta, "substituir": True, "stats": {"min": 1, "mean": 2, "max": 3}}), url)
        self.assertEqual(len(self.p.ppt_queue), 1)
        self.assertEqual(self.p.ppt_queue[0]["stats"]["min"], 1.0)
        self.assertTrue(self.p.ppt_queue[0]["miniatura"].startswith("data:image/jpeg"))
        self.assertTrue(self.p.pptx_da_fila().startswith(b"PK"))
        self.p.removerDaFila(f"{self.caso}|Aprovação")
        self.assertEqual(self.p.ppt_queue, [])
        # estado do mapa volta quando a página reabre
        self.p.saveKmls('[{"name": "Centro.kml", "text": "<kml/>"}]')
        self.p.saveMapState('{"caseKey": "x"}')
        self.p.eventos = []
        self.p.estado_inicial()
        ultimo = [e["codigo"] for e in self.p.eventos if e["tipo"] == "js"][-1]
        self.assertTrue(ultimo.startswith("window.__restoreMapState("))
        self.assertIn("Centro.kml", ultimo)

    def test_bases_por_regiao(self):
        payload = json.loads(_args(self._js("window.__setCases")[-1])[0])
        self.assertEqual(payload["baseRegionsByCase"][self.caso], {"Centro": 300.0, "Sul": 300.0, "Norte": 300.0})
        self.assertEqual(payload["regionsByCase"][self.caso], ["Centro", "Sul", "Norte"])

    def test_versao_da_tela(self):
        pasta = os.path.join(os.path.dirname(os.path.dirname(__file__)), "componentes", "relatoria_mapa_calor")
        html = open(os.path.join(pasta, "index.html"), encoding="utf-8").read()
        self.assertIn(f'<meta name="versao-tela" content="{VERSAO_TELA}" />', html)
        ponte_js = os.path.join(os.path.dirname(pasta), "ponte_qt.js")
        self.assertEqual(open(os.path.join(pasta, "ponte_qt.js"), "rb").read(), open(ponte_js, "rb").read())


if __name__ == "__main__":
    unittest.main()
