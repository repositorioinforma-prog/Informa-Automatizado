import unittest

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from core.total_automatico_math import (
    aplicar_codigo_09_total,
    aplicar_codigo_11_total,
    aplicar_codigo_12_total,
)


class TestAlturasTotalAutomatico(unittest.TestCase):
    def setUp(self):
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.column_dimensions["A"].width = 21
        self.ws.column_dimensions["C"].width = 15

    def test_label_uma_linha_reduz_para_15(self):
        c = self.ws["C1"]
        c.value = "Masculino"
        c.font = Font(name="DIN", size=10)
        c.alignment = Alignment(wrap_text=True)
        self.ws.row_dimensions[1].height = 30

        aplicar_codigo_09_total(self.ws)

        self.assertEqual(self.ws.row_dimensions[1].height, 15)

    def test_label_duas_linhas_manuais_fica_30(self):
        c = self.ws["C1"]
        c.value = "Ensino\nMédio"
        c.font = Font(name="DIN", size=10)
        c.alignment = Alignment(wrap_text=True)
        self.ws.row_dimensions[1].height = 45

        aplicar_codigo_09_total(self.ws)

        self.assertEqual(self.ws.row_dimensions[1].height, 30)

    def test_label_tres_linhas_manuais_fica_45(self):
        c = self.ws["C1"]
        c.value = "Ensino\nFunda-\nmental"
        c.font = Font(name="DIN", size=10)
        c.alignment = Alignment(wrap_text=True)

        aplicar_codigo_09_total(self.ws)

        self.assertEqual(self.ws.row_dimensions[1].height, 45)

    def test_titulo_curto_pode_ficar_em_15(self):
        c = self.ws["A1"]
        c.value = "Titulo: Intenção de voto"
        c.font = Font(name="DIN", size=10)
        c.alignment = Alignment(wrap_text=True)
        self.ws.merge_cells("A1:H1")
        self.ws.row_dimensions[1].height = 30

        aplicar_codigo_11_total(self.ws)

        self.assertEqual(c.value, "Intenção de voto")
        self.assertEqual(self.ws.row_dimensions[1].height, 15)
        self.assertTrue(c.font.bold)

    def test_pergunta_com_quebra_manual_usa_multiplos_de_15(self):
        c = self.ws["A1"]
        c.value = "Pergunta: Primeira linha\nSegunda linha"
        c.font = Font(name="DIN", size=10)
        c.alignment = Alignment(wrap_text=True)
        self.ws.merge_cells("A1:H1")

        aplicar_codigo_12_total(self.ws)

        self.assertEqual(self.ws.row_dimensions[1].height, 30)


if __name__ == "__main__":
    unittest.main()
