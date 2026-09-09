import io
import unittest

import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.pagebreak import Break

from core.total_automatico_math import adicionar_linhas_finais_em_branco


class TestLinhasFinaisTotalAutomatico(unittest.TestCase):
    def _criar_planilha(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Relatorio"
        ws["A1"] = "Titulo"
        ws["M10"] = "fim visual"
        ws["A10"] = "Pergunta: teste"
        ws.print_area = "A1:M10"
        ws.row_breaks.append(Break(id=10))
        return wb, ws

    def test_adiciona_duas_linhas_e_move_limite_final(self):
        wb, ws = self._criar_planilha()

        ultima = adicionar_linhas_finais_em_branco(ws)

        self.assertEqual(ultima, 10)
        self.assertEqual(ws.row_dimensions[11].height, 15)
        self.assertEqual(ws.row_dimensions[12].height, 15)
        self.assertEqual(ws["A11"].font.name, "DIN")
        self.assertEqual(ws["A12"].font.name, "DIN")
        self.assertIn("H$12", str(ws.print_area))
        self.assertNotIn(10, [int(b.id) for b in ws.row_breaks.brk if b.id is not None])
        self.assertIn(12, [int(b.id) for b in ws.row_breaks.brk if b.id is not None])

    def test_persiste_apos_salvar_e_reabrir(self):
        wb, ws = self._criar_planilha()
        adicionar_linhas_finais_em_branco(ws)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        wb2 = openpyxl.load_workbook(buf)
        ws2 = wb2.active

        self.assertGreaterEqual(ws2.max_row, 12)
        self.assertEqual(ws2.row_dimensions[11].height, 15)
        self.assertEqual(ws2.row_dimensions[12].height, 15)
        self.assertIn("H$12", str(ws2.print_area))
        self.assertIn(12, [int(b.id) for b in ws2.row_breaks.brk if b.id is not None])

    def test_reaplicar_nao_cria_novo_respiro(self):
        _, ws = self._criar_planilha()
        adicionar_linhas_finais_em_branco(ws)
        adicionar_linhas_finais_em_branco(ws)

        self.assertIn("H$12", str(ws.print_area))
        self.assertEqual(
            [int(b.id) for b in ws.row_breaks.brk if b.id is not None and int(b.id) >= 10],
            [12],
        )


if __name__ == "__main__":
    unittest.main()
